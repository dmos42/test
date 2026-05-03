import sys
import os
import json
import tempfile
import shutil
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use("Qt5Agg")
import matplotlib.pyplot as plt
from matplotlib.backends.backend_qt5agg import FigureCanvasQTAgg as FigureCanvas
from matplotlib.backends.backend_qt5agg import NavigationToolbar2QT as NavigationToolbar
from matplotlib.figure import Figure
from scipy import stats as scipy_stats

from PyQt5.QtWidgets import (
    QApplication, QMainWindow, QTabWidget, QWidget, QVBoxLayout, QHBoxLayout,
    QLabel, QLineEdit, QPushButton, QTextEdit, QFileDialog, QMessageBox,
    QGroupBox, QFormLayout, QCheckBox, QComboBox, QSpinBox, QDoubleSpinBox,
    QSplitter, QMenuBar, QMenu, QAction, QDialog, QDialogButtonBox, QStatusBar,
    QTableWidget, QTableWidgetItem, QHeaderView, QAbstractItemView,
    QFrame, QToolBar, QSizePolicy, QListWidget, QListWidgetItem, QStackedWidget, QStyle,
)
from PyQt5.QtCore import Qt, QMimeData, pyqtSignal
from PyQt5.QtGui import QFont, QKeySequence, QPixmap, QIcon

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from stats_engine.capability import CapabilityAnalysis
from stats_engine.normality import NormalityTests
from stats_engine.outliers import OutlierDetection
from stats_engine.correlation import CorrelationMatrix
from stats_engine.probability_plots import ProbabilityPlot
from stats_engine.msa import GageRRStudy
from stats_engine.export import ReportExporter



# ===== MINIQUAL CORE COPIÉ DE core.py =====
from pathlib import Path
from datetime import datetime
import json, math
import numpy as np, pandas as pd
from scipy import stats
import matplotlib.pyplot as plt
SUPPORTED_DISTRIBUTIONS=['normal','lognormal','weibull','gamma','exponential']
LABEL={'normal':'normale','lognormal':'lognormale','weibull':'Weibull','gamma':'gamma','exponential':'exponentielle'}
def read_table(path, sheet_name=0):
    p=Path(path); s=p.suffix.lower()
    if s=='.csv': df=pd.read_csv(p,sep=None,engine='python',encoding='utf-8-sig')
    elif s in ['.xlsx','.xlsm']: df=pd.read_excel(p,sheet_name=sheet_name,engine='openpyxl')
    elif s=='.xls': df=pd.read_excel(p,sheet_name=sheet_name)
    else: raise ValueError('Format non supporté: '+s)
    df.columns=[str(c).lstrip('\ufeff').strip() for c in df.columns]; return df
def numeric_series(df,col):
    if col not in df.columns: raise KeyError(f'Colonne introuvable: {col}. Colonnes: {list(df.columns)}')
    s=pd.to_numeric(df[col],errors='coerce').dropna()
    if s.empty: raise ValueError('Aucune valeur numérique exploitable')
    return s
def descriptive(s):
    s=pd.to_numeric(s,errors='coerce').dropna(); return {'n':len(s),'moyenne':float(s.mean()),'mediane':float(s.median()),'ecart_type_echantillon':float(s.std(ddof=1)),'minimum':float(s.min()),'q1':float(s.quantile(.25)),'q3':float(s.quantile(.75)),'maximum':float(s.max())}
def _pos(x,d):
    if d in ['lognormal','weibull','gamma']: return x[x>0]
    if d=='exponential': return x[x>=0]
    return x
def _dist(d): return {'normal':stats.norm,'lognormal':stats.lognorm,'weibull':stats.weibull_min,'gamma':stats.gamma,'exponential':stats.expon}[d]
def _name(d): return {'normal':'norm','lognormal':'lognorm','weibull':'weibull_min','gamma':'gamma','exponential':'expon'}[d]
def fit_distribution(series,d):
    x=pd.Series(series).dropna().astype(float).to_numpy(); xp=_pos(x,d)
    if len(xp)<3: return {'loi':d,'p_value':np.nan,'statistique':np.nan,'paramètres':'Données incompatibles','params':None}
    params=stats.norm.fit(xp) if d=='normal' else _dist(d).fit(xp,floc=0)
    D,p=stats.kstest(xp,_name(d),args=params)
    return {'loi':d,'p_value':float(p),'statistique':float(D),'paramètres':str(tuple(round(float(v),6) for v in params)),'params':params}
def distribution_pvalues(s): return sorted([{k:v for k,v in fit_distribution(s,d).items() if k!='params'} for d in SUPPORTED_DISTRIBUTIONS], key=lambda r:(-1 if r['p_value']!=r['p_value'] else -r['p_value']))
def chi_square_gof(s, distribution='normal', bins='auto'):
    x=pd.Series(s).dropna().astype(float).to_numpy(); n=len(x)
    if n<10: return {'Test':'Khi²','Loi testée':distribution,'Statistique':np.nan,'ddl':np.nan,'p-value':np.nan,'Lecture':'Effectif insuffisant (<10)','Classes':'n/a'}
    k=max(4,min(10,int(np.sqrt(n)))) if str(bins)=='auto' else int(bins)
    fit=fit_distribution(x,distribution); p=fit.get('params')
    if p is None: return {'Test':'Khi²','Loi testée':distribution,'Statistique':np.nan,'ddl':np.nan,'p-value':np.nan,'Lecture':'Paramètres non estimables','Classes':'n/a'}
    edges=_dist(distribution).ppf(np.linspace(0,1,k+1),*p); edges[0],edges[-1]=-np.inf,np.inf
    obs,_=np.histogram(x,bins=edges); exp=np.ones(k)*n/k; chi=float(((obs-exp)**2/exp).sum()); ddl=max(1,k-1-len(p)); pv=float(stats.chi2.sf(chi,ddl))
    return {'Test':'Khi²','Loi testée':distribution,'Statistique':chi,'ddl':ddl,'p-value':pv,'Lecture':'Compatible avec la loi testée' if pv>0.05 else 'Écart possible avec la loi testée','Classes':k}
def boxcox_transform(s,lsl=None,usl=None,target=None):
    x=pd.Series(s).dropna().astype(float); shift=0.0
    if x.min()<=0: shift=float(abs(x.min())+1e-6)
    y,lam=stats.boxcox(x+shift)
    def tr(v):
        if v is None: return None
        vv=float(v)+shift; return float(stats.boxcox([vv],lmbda=lam)[0]) if vv>0 else None
    return pd.Series(y,index=x.index),tr(lsl),tr(usl),tr(target),{'Box-Cox lambda':float(lam),'Décalage appliqué':shift}
def validate_capability(df,col,lsl=None,usl=None,target=None,subgroup=None,distribution='normal'):
    rows=[]; add=lambda n,m: rows.append({'Niveau':n,'Message':m})
    if col not in df.columns: add('ERREUR',f'Colonne mesure introuvable: {col}'); return rows
    s=pd.to_numeric(df[col],errors='coerce'); n=s.notna().sum(); bad=len(df)-n
    add('OK',f'Colonne mesure trouvée: {col}'); add('OK',f'{int(n)} valeurs numériques exploitables sur {len(df)} lignes')
    if bad: add('AVERTISSEMENT',f'{int(bad)} valeurs vides ou non numériques seront ignorées')
    if n<10: add('ERREUR','Moins de 10 valeurs numériques: analyse non robuste')
    elif n<30: add('AVERTISSEMENT','Moins de 30 valeurs: interprétation prudente')
    if lsl is None and usl is None: add('ERREUR','Aucune limite de spécification renseignée')
    if lsl is not None and usl is not None and usl<=lsl: add('ERREUR','USL doit être strictement supérieure à LSL')
    if target is not None and lsl is not None and usl is not None and not(lsl<=target<=usl): add('AVERTISSEMENT','La cible est hors intervalle [LSL;USL]')
    if subgroup and subgroup not in df.columns: add('AVERTISSEMENT',f'Sous-groupe introuvable: {subgroup}. Calcul sans sous-groupe')
    vals=s.dropna()
    if distribution in ['lognormal','weibull','gamma'] and (vals<=0).any(): add('AVERTISSEMENT',f'Loi {distribution}: des valeurs <=0 sont incompatibles avec cette loi')
    if distribution=='exponential' and (vals<0).any(): add('AVERTISSEMENT','Loi exponentielle: des valeurs <0 sont incompatibles')
    return rows
def status(rows): return 'BLOQUÉ' if any(r['Niveau']=='ERREUR' for r in rows) else ('OK AVEC AVERTISSEMENT' if any(r['Niveau']=='AVERTISSEMENT' for r in rows) else 'OK')
def capability(s,lsl=None,usl=None,target=None,subgroup=None):
    x=pd.Series(s).dropna().astype(float); m=float(x.mean()); so=float(x.std(ddof=1)); mr=x.diff().abs().dropna(); sw=float(mr.mean()/1.128) if len(mr) and mr.mean()>0 else so
    if lsl is None and usl is None: raise ValueError('Renseigner LSL ou USL')
    if lsl is not None and usl is not None and target is None: target=(lsl+usl)/2
    cu=(usl-m)/(3*sw) if usl is not None and sw>0 else None; cl=(m-lsl)/(3*sw) if lsl is not None and sw>0 else None
    pu=(usl-m)/(3*so) if usl is not None and so>0 else None; pl=(m-lsl)/(3*so) if lsl is not None and so>0 else None
    cp=(usl-lsl)/(6*sw) if lsl is not None and usl is not None and sw>0 else None; pp=(usl-lsl)/(6*so) if lsl is not None and usl is not None and so>0 else None
    valid=lambda v: v is not None and not np.isnan(v)
    mask=pd.Series(False,index=x.index)
    if lsl is not None: mask|=x<lsl
    if usl is not None: mask|=x>usl
    return {'n':len(x),'mean':m,'stdev_within':sw,'stdev_overall':so,'lsl':lsl,'usl':usl,'target':target,'cp':cp,'cpk':min([v for v in [cu,cl] if valid(v)]) if any(valid(v) for v in [cu,cl]) else None,'cpk_upper':cu,'cpk_lower':cl,'pp':pp,'ppk':min([v for v in [pu,pl] if valid(v)]) if any(valid(v) for v in [pu,pl]) else None,'ppk_upper':pu,'ppk_lower':pl,'ppm_below_lsl':float(stats.norm.cdf((lsl-m)/so)*1e6) if lsl is not None and so>0 else None,'ppm_above_usl':float((1-stats.norm.cdf((usl-m)/so))*1e6) if usl is not None and so>0 else None,'observed_nc_count':int(mask.sum()),'observed_nc_percent':float(mask.mean()*100),'spec_mode':'Bilatéral LSL+USL' if lsl is not None and usl is not None else ('Unilatéral supérieur USL' if usl is not None else 'Unilatéral inférieur LSL')}
def outlier_tests(s,alpha=.05,z_threshold=3.0):
    x=pd.Series(s).dropna().astype(float); q1,q3=x.quantile(.25),x.quantile(.75); iqr=q3-q1; lo,hi=q1-1.5*iqr,q3+1.5*iqr; oi=x[(x<lo)|(x>hi)]
    z=(x-x.mean())/x.std(ddof=1) if x.std(ddof=1)>0 else x*0; oz=x[z.abs()>z_threshold]
    return {'table':[{'Méthode':'IQR','Seuil / statistique':f'[{lo:.6g}; {hi:.6g}]','Valeurs détectées':', '.join(map(str,oi.tolist()[:12])) if len(oi) else 'Aucune','Nombre':len(oi)},{'Méthode':f'Z-score absolu > {z_threshold:g}','Seuil / statistique':f'|Z| > {z_threshold:g}','Valeurs détectées':', '.join(map(str,oz.tolist()[:12])) if len(oz) else 'Aucune','Nombre':len(oz)}], 'values_to_exclude':list(set(map(float,oi)).union(set(map(float,oz))))}
def normality_tests(s,alpha=.05):
    x=pd.Series(s).dropna().astype(float).to_numpy(); sh,shp=stats.shapiro(x if len(x)<=5000 else pd.Series(x).sample(5000,random_state=42)); ks,ksp=stats.kstest(x,'norm',args=(np.mean(x),np.std(x,ddof=1))); ad=stats.anderson(x,'norm'); crit=float(ad.critical_values[2]); ok=[shp>alpha,ksp>alpha,float(ad.statistic)<crit]
    return {'table':[{'Test':'Shapiro-Wilk','Statistique':float(sh),'p-value / seuil':float(shp),'Lecture':'Compatible avec une loi normale' if ok[0] else 'Écart possible à la normalité'},{'Test':'Kolmogorov-Smirnov','Statistique':float(ks),'p-value / seuil':float(ksp),'Lecture':'Compatible avec une loi normale' if ok[1] else 'Écart possible à la normalité'},{'Test':'Anderson-Darling','Statistique':float(ad.statistic),'p-value / seuil':f'Critique 5% : {crit:.6g}','Lecture':'Compatible avec une loi normale' if ok[2] else 'Écart possible à la normalité'}],'favorable_count':int(sum(ok)),'global_ok':sum(ok)>=2}
def nonnormal_capability(s,lsl=None,usl=None):
    x=pd.Series(s).dropna().astype(float); lo=x.quantile(.00135); hi=x.quantile(.99865); med=x.median(); out={'méthode':'Percentiles empiriques 0,135% / 99,865%','q0_135':float(lo),'q99_865':float(hi),'mediane':float(med)}
    if lsl is not None and usl is not None and hi>lo: out['CNp']=float((usl-lsl)/(hi-lo)); out['CNpk']=float(min((usl-med)/(hi-med),(med-lsl)/(med-lo))) if hi!=med and med!=lo else None
    return out
def dashboard(s,res,out,normality=None,distribution='normal',accept=1.33,excellent=1.67):
    data=pd.Series(s).dropna().astype(float).to_numpy(); m=np.mean(data); sd=np.std(data,ddof=1); fig,ax=plt.subplots(2,2,figsize=(15,12)); fig.suptitle('Analyse de capabilité du procédé',fontsize=18,fontweight='bold')
    a=ax[0,0]; a.hist(data,bins=min(15,max(5,int(np.sqrt(len(data))))),density=True,alpha=.75,edgecolor='black'); xs=[data.min(),data.max(),m-4*sd,m+4*sd]+[v for v in [res.get('lsl'),res.get('usl'),res.get('target')] if v is not None]; xx=np.linspace(min(xs),max(xs),300); fit=fit_distribution(data,distribution); p=fit.get('params')
    if p is not None: a.plot(xx,_dist(distribution).pdf(xx,*p),'r-',label=f"Courbe {LABEL.get(distribution,distribution)} estimée")
    for key,c,lab,ls in [('lsl','green','LSL','--'),('usl','red','USL','--'),('target','purple','Cible',':')]:
        if res.get(key) is not None: a.axvline(res[key],color=c,ls=ls,label=f'{lab}={res[key]:.3g}')
    a.axvline(m,color='orange',label=f'Moyenne={m:.3g}'); a.legend(fontsize=9); a.set_title('Distribution')
    ax[0,1].boxplot(data); ax[0,1].set_title('Boxplot')
    qfit=fit_distribution(data,distribution); p=qfit.get('params')
    if p is not None:
        probs=(np.arange(1,len(data)+1)-.5)/len(data); q=_dist(distribution).ppf(probs,*p); ordered=np.sort(_pos(data,distribution)); ax[1,0].scatter(q,ordered); lo=min(q.min(),ordered.min()); hi=max(q.max(),ordered.max()); ax[1,0].plot([lo,hi],[lo,hi],'r-')
    ax[1,0].set_title(f"Q-Q plot — {LABEL.get(distribution,distribution)}")
    labels=[]; vals=[]
    for lab,key in [('Cp','cp'),('Cpk inf.','cpk_lower'),('Cpk sup.','cpk_upper'),('Cpk global','cpk')]:
        if res.get(key) is not None: labels.append(lab); vals.append(res[key])
    ax[1,1].bar(labels,vals); ax[1,1].axhline(accept,color='orange',ls='--'); ax[1,1].axhline(excellent,color='green',ls='--'); ax[1,1].set_title('Indices')
    fig.tight_layout(rect=[0,0,1,.96]); fig.savefig(out,dpi=180,bbox_inches='tight'); plt.close(fig)
def write_excel(path,sheets):
    with pd.ExcelWriter(path,engine='openpyxl') as w:
        for n,d in sheets.items(): (pd.DataFrame([d]) if isinstance(d,dict) else pd.DataFrame(d)).to_excel(w,sheet_name=str(n)[:31],index=False)
def write_docx(title,sections,path,signature=None):
    from docx import Document
    from docx.shared import Inches
    doc=Document(); doc.add_heading(title,0); doc.add_paragraph('MiniQual Python v7.5.2 — '+datetime.now().strftime('%d/%m/%Y %H:%M'))
    for name,content in sections:
        doc.add_heading(name,1)
        if isinstance(content,dict):
            t=doc.add_table(rows=1,cols=2); t.style='Table Grid'; t.rows[0].cells[0].text='Indicateur'; t.rows[0].cells[1].text='Valeur'
            for k,v in content.items(): c=t.add_row().cells; c[0].text=str(k); c[1].text='' if v is None else str(v)
        elif isinstance(content,list) and content:
            cols=list(content[0].keys()); t=doc.add_table(rows=1,cols=len(cols)); t.style='Table Grid'
            for i,c0 in enumerate(cols): t.rows[0].cells[i].text=str(c0)
            for row in content:
                c=t.add_row().cells
                for i,c0 in enumerate(cols): c[i].text=str(row.get(c0,''))
        elif str(content).endswith('.png') and Path(content).exists(): doc.add_picture(str(content),width=Inches(6.7))
        else: doc.add_paragraph(str(content))
    doc.save(path)
def save_project(path,settings):
    p=Path(path); p.parent.mkdir(parents=True,exist_ok=True); p.write_text(json.dumps({k:(str(v) if callable(v) else v) for k,v in settings.items() if k!='func'},ensure_ascii=False,indent=2),encoding='utf-8')
def load_project(path): return json.loads(Path(path).read_text(encoding='utf-8'))

# ===== FIN MINIQUAL CORE =====

NUM_COLS = 26
NUM_ROWS = 500


def col_letter(n):
    result = ""
    while n >= 0:
        result = chr(65 + (n % 26)) + result
        n = n // 26 - 1
    return result


class DataSheet(QTableWidget):
    def __init__(self, parent=None):
        super().__init__(NUM_ROWS, NUM_COLS, parent)
        self.setHorizontalHeaderLabels([col_letter(i) for i in range(NUM_COLS)])
        self.setVerticalHeaderLabels([str(i + 1) for i in range(NUM_ROWS)])

        header = self.horizontalHeader()
        # Colonnes redimensionnables manuellement et plus larges par défaut.
        header.setSectionResizeMode(QHeaderView.Interactive)
        header.setDefaultSectionSize(95)
        header.setMinimumSectionSize(55)
        header.setDefaultAlignment(Qt.AlignCenter)

        vheader = self.verticalHeader()
        vheader.setSectionResizeMode(QHeaderView.Interactive)
        vheader.setDefaultSectionSize(26)
        vheader.setMinimumSectionSize(22)
        vheader.setDefaultAlignment(Qt.AlignCenter)

        self.setContextMenuPolicy(Qt.CustomContextMenu)
        self.customContextMenuRequested.connect(self._show_context_menu)

        self.setAlternatingRowColors(False)
        self.setStyleSheet("""
            QTableWidget {
                gridline-color: #c0c0c0;
                font-family: 'Courier New', monospace;
                font-size: 12px;
                background-color: white;
            }
            QHeaderView::section {
                background-color: #e8e8e8;
                color: #1b1b1b;
                border: 1px solid #9a9a9a;
                padding: 2px;
                font-weight: bold;
            }
            QTableCornerButton::section {
                background-color: #e8e8e8;
                border: 1px solid #9a9a9a;
            }
            QTableWidget::item {
                padding: 2px 6px;
                color: #1a1a1a;
                background-color: white;
            }
            QTableWidget::item:selected {
                background-color: #3399ff;
                color: white;
            }
        """)
        self._undo_stack = []
        self._redo_stack = []
        self._max_undo = 30
        self.setSelectionBehavior(QAbstractItemView.SelectItems)
        self.setSelectionMode(QAbstractItemView.ExtendedSelection)

    def _show_context_menu(self, pos):
        menu = QMenu(self)
        
        cut_action = menu.addAction(" Couper")
        cut_action.triggered.connect(self._cut)
        
        copy_action = menu.addAction(" Copier")
        copy_action.triggered.connect(self._copy)
        
        paste_action = menu.addAction(" Coller")
        paste_action.triggered.connect(self._paste)
        
        menu.addSeparator()
        
        clear_action = menu.addAction(" Effacer cellules")
        clear_action.triggered.connect(self._clear_cells)
        
        fill_action = menu.addAction(" Remplir avec séquence...")
        fill_action.triggered.connect(self._fill_sequence)
        
        menu.addSeparator()
        
        insert_row_action = menu.addAction(" Insérer ligne")
        insert_row_action.triggered.connect(self._insert_row)
        
        delete_row_action = menu.addAction(" Supprimer ligne")
        delete_row_action.triggered.connect(self._delete_row)

        menu.addSeparator()

        insert_col_action = menu.addAction(" Insérer colonne")
        insert_col_action.triggered.connect(self._insert_col)

        delete_col_action = menu.addAction(" Supprimer colonne")
        delete_col_action.triggered.connect(self._delete_col)

        clear_col_action = menu.addAction(" Effacer colonne")
        clear_col_action.triggered.connect(self._clear_column)
        
        menu.addSeparator()
        
        stats_action = menu.addAction(" Statistiques rapides...")
        stats_action.triggered.connect(self._quick_stats)
        
        menu.exec_(self.mapToGlobal(pos))

    def _serialize_cells(self):
        data = {}
        for row in range(self.rowCount()):
            for col in range(self.columnCount()):
                item = self.item(row, col)
                if item and item.text():
                    data[(row, col)] = item.text()
        return data

    def _restore_cells(self, data):
        self.clearContents()
        for (row, col), value in data.items():
            if row < self.rowCount() and col < self.columnCount():
                item = QTableWidgetItem(value)
                item.setTextAlignment(Qt.AlignRight | Qt.AlignVCenter)
                self.setItem(row, col, item)

    def _push_undo(self):
        self._undo_stack.append(self._serialize_cells())
        if len(self._undo_stack) > self._max_undo:
            self._undo_stack.pop(0)
        self._redo_stack.clear()

    def _undo(self):
        if not self._undo_stack:
            return
        self._redo_stack.append(self._serialize_cells())
        self._restore_cells(self._undo_stack.pop())

    def _redo(self):
        if not self._redo_stack:
            return
        self._undo_stack.append(self._serialize_cells())
        self._restore_cells(self._redo_stack.pop())

    def _clear_cells(self):
        self._push_undo()
        selected = self.selectedIndexes()
        for idx in selected:
            self.setItem(idx.row(), idx.column(), None)

    def _fill_sequence(self):
        self._push_undo()
        selected = self.selectedIndexes()
        if not selected:
            return
        rows = sorted(set(idx.row() for idx in selected))
        cols = sorted(set(idx.column() for idx in selected))
        if not rows or not cols:
            return
        
        from PyQt5.QtWidgets import QInputDialog
        start, ok = QInputDialog.getDouble(self, "Remplir séquence", "Valeur de départ :", 1.0, -1e9, 1e9, 6)
        if not ok:
            return
        step, ok2 = QInputDialog.getDouble(self, "Remplir séquence", "Incrément :", 1.0, -1e9, 1e9, 6)
        if not ok2:
            return
        
        val = start
        for row in rows:
            for col in cols:
                item = QTableWidgetItem(f"{val:.6f}")
                item.setTextAlignment(Qt.AlignRight | Qt.AlignVCenter)
                self.setItem(row, col, item)
                val += step

    def _insert_row(self):
        self._push_undo()
        current = self.currentIndex()
        if not current.isValid():
            return
        row = current.row()
        self.insertRow(row)
        new_label = [str(i + 1) for i in range(self.rowCount())]
        self.setVerticalHeaderLabels(new_label)

    def _delete_row(self):
        self._push_undo()
        current = self.currentIndex()
        if not current.isValid():
            return
        row = current.row()
        self.removeRow(row)
        new_label = [str(i + 1) for i in range(self.rowCount())]
        self.setVerticalHeaderLabels(new_label)

    def _insert_col(self):
        self._push_undo()
        current = self.currentIndex()
        col = current.column() if current.isValid() else self.columnCount()
        self.insertColumn(col)
        self.setHorizontalHeaderLabels([col_letter(i) for i in range(self.columnCount())])

    def _delete_col(self):
        current = self.currentIndex()
        if not current.isValid():
            return
        self._push_undo()
        self.removeColumn(current.column())
        self.setHorizontalHeaderLabels([col_letter(i) for i in range(self.columnCount())])

    def _clear_column(self):
        current = self.currentIndex()
        if not current.isValid():
            return
        self._push_undo()
        col = current.column()
        for row in range(self.rowCount()):
            self.setItem(row, col, None)

    def _quick_stats(self):
        selected = self.selectedIndexes()
        if not selected:
            return
        rows = sorted(set(idx.row() for idx in selected))
        cols = sorted(set(idx.column() for idx in selected))
        values = []
        for row in rows:
            for col in cols:
                item = self.item(row, col)
                if item and item.text().strip():
                    try:
                        values.append(float(item.text().strip()))
                    except ValueError:
                        pass
        if not values:
            QMessageBox.information(self, "Statistiques rapides", "Aucune valeur numérique dans la sélection.")
            return
        arr = np.array(values)
        msg = (f"Sélection : {len(arr)} valeurs\n\n"
               f"Moyenne = {np.mean(arr):.6f}\n"
               f"Médiane = {np.median(arr):.6f}\n"
               f"Écart-type = {np.std(arr, ddof=1):.6f}\n"
               f"Min = {np.min(arr):.6f}\n"
               f"Max = {np.max(arr):.6f}\n"
               f"Somme = {np.sum(arr):.6f}")
        QMessageBox.information(self, "Statistiques rapides", msg)

    def keyPressEvent(self, event):
        if event.matches(QKeySequence.Copy):
            self._copy()
        elif event.matches(QKeySequence.Paste):
            self._paste()
        elif event.matches(QKeySequence.Cut):
            self._cut()
        elif event.matches(QKeySequence.Undo):
            self._undo()
        elif event.matches(QKeySequence.Redo):
            self._redo()
        else:
            super().keyPressEvent(event)

    def _copy(self):
        selected = self.selectedIndexes()
        if not selected:
            return

        rows = sorted(set(idx.row() for idx in selected))
        cols = sorted(set(idx.column() for idx in selected))

        lines = []
        for row in rows:
            row_data = []
            for col in cols:
                idx = self.model().index(row, col)
                val = self.model().data(idx, Qt.DisplayRole)
                row_data.append(val if val else "")
            lines.append("\t".join(row_data))

        QApplication.clipboard().setText("\n".join(lines))

    def _paste(self):
        self._push_undo()
        text = QApplication.clipboard().text()
        if not text:
            return

        current = self.currentIndex()
        if not current.isValid():
            return

        start_row = current.row()
        start_col = current.column()

        clipboard = QApplication.clipboard()
        mime = clipboard.mimeData()
        if mime.hasText():
            text = mime.text()
        else:
            return

        rows = text.split("\n")
        if rows and rows[-1] == "":
            rows = rows[:-1]

        for i, row_text in enumerate(rows):
            cells = row_text.split("\t")
            for j, cell_text in enumerate(cells):
                row = start_row + i
                col = start_col + j
                if row < self.rowCount() and col < self.columnCount():
                    item = QTableWidgetItem(cell_text.strip())
                    item.setTextAlignment(Qt.AlignRight | Qt.AlignVCenter)
                    self.setItem(row, col, item)

    def _cut(self):
        self._push_undo()
        self._copy()
        selected = self.selectedIndexes()
        for idx in selected:
            self.setItem(idx.row(), idx.column(), None)

    def get_column_data(self, col_idx):
        values = []
        if col_idx < 0 or col_idx >= self.columnCount():
            return None

        for row in range(self.rowCount()):
            item = self.item(row, col_idx)
            if item and item.text().strip():
                txt = item.text().strip().replace(",", ".")
                try:
                    values.append(float(txt))
                except ValueError:
                    pass
        return np.array(values, dtype=float) if values else None

    def get_selected_column_data(self):
        selected = self.selectionModel().selectedColumns()
        if not selected:
            selected = self.selectionModel().selectedIndexes()
            if not selected:
                cols = set()
                for row in range(self.rowCount()):
                    for col in range(self.columnCount()):
                        item = self.item(row, col)
                        if item and item.text().strip():
                            cols.add(col)
                if cols:
                    col_idx = sorted(cols)[0]
                    return self.get_column_data(col_idx), col_letter(col_idx)
                return None, None

        if hasattr(selected, '__iter__'):
            col_idx = list(selected)[0].column() if hasattr(list(selected)[0], 'column') else 0
        else:
            col_idx = selected.column() if hasattr(selected, 'column') else 0

        return self.get_column_data(col_idx), col_letter(col_idx)

    def load_from_dataframe(self, df):
        self.clearContents()

        max_rows = min(len(df), self.rowCount())
        max_cols = min(len(df.columns), self.columnCount())

        self.setHorizontalHeaderLabels([col_letter(i) for i in range(self.columnCount())])

        for col_idx in range(max_cols):
            for row_idx in range(max_rows):
                val = df.iloc[row_idx, col_idx]
                if pd.isna(val):
                    continue
                item = QTableWidgetItem(str(val))
                item.setTextAlignment(Qt.AlignRight | Qt.AlignVCenter)
                self.setItem(row_idx, col_idx, item)

    def load_from_array(self, data, start_col=0):
        max_rows = min(len(data), NUM_ROWS)
        for row_idx in range(max_rows):
            item = QTableWidgetItem(f"{data[row_idx]:.6f}")
            item.setTextAlignment(Qt.AlignRight | Qt.AlignVCenter)
            self.setItem(row_idx, start_col, item)


class GenerateDataDialog(QDialog):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Générer des données")
        self.resize(400, 250)

        layout = QFormLayout(self)

        self.dist_combo = QComboBox()
        self.dist_combo.addItems(["Normal", "Uniforme", "Exponentielle"])
        layout.addRow("Distribution :", self.dist_combo)

        self.n_spin = QSpinBox()
        self.n_spin.setRange(10, NUM_ROWS)
        self.n_spin.setValue(100)
        layout.addRow("Nombre de valeurs :", self.n_spin)

        self.p1_spin = QDoubleSpinBox()
        self.p1_spin.setRange(-1e9, 1e9)
        self.p1_spin.setValue(50.0)
        layout.addRow("Moyenne (ou min) :", self.p1_spin)

        self.p2_spin = QDoubleSpinBox()
        self.p2_spin.setRange(-1e9, 1e9)
        self.p2_spin.setSingleStep(0.5)
        self.p2_spin.setValue(5.0)
        layout.addRow("Écart-type (ou max) :", self.p2_spin)

        self.seed_spin = QSpinBox()
        self.seed_spin.setRange(0, 99999)
        self.seed_spin.setValue(42)
        layout.addRow("Seed (optionnel) :", self.seed_spin)

        btn_box = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        btn_box.accepted.connect(self.accept)
        btn_box.rejected.connect(self.reject)
        layout.addWidget(btn_box)

    def generate(self):
        n = self.n_spin.value()
        np.random.seed(self.seed_spin.value())
        dist = self.dist_combo.currentText()
        p1 = self.p1_spin.value()
        p2 = self.p2_spin.value()

        if dist == "Normal":
            return np.random.normal(p1, p2, n)
        elif dist == "Uniforme":
            return np.random.uniform(p1, p2, n)
        elif dist == "Exponentielle":
            return np.random.exponential(p2, n) + p1


class ColumnSelectDialog(QDialog):
    def __init__(self, parent=None, columns=None):
        super().__init__(parent)
        self.setWindowTitle("Sélectionner la colonne")
        self.resize(300, 200)

        layout = QVBoxLayout(self)

        layout.addWidget(QLabel("Colonne de données à analyser :"))

        self.col_combo = QComboBox()
        if columns:
            self.col_combo.addItems(columns)
        layout.addWidget(self.col_combo)

        btn_box = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        btn_box.accepted.connect(self.accept)
        btn_box.rejected.connect(self.reject)
        layout.addWidget(btn_box)

    def get_column(self):
        return self.col_combo.currentText()


class MplCanvas(FigureCanvas):
    def __init__(self, parent=None, width=7, height=8, dpi=100):
        self.fig = Figure(figsize=(width, height), dpi=dpi)
        super().__init__(self.fig)
        self.setParent(parent)


class SafeNavigationToolbar(NavigationToolbar):
    def __init__(self, canvas, parent=None):
        super().__init__(canvas, parent)
        self._add_copy_action_on_left()

    def _add_copy_action_on_left(self):
        """Ajoute l'action Copier près des icônes de navigation, à gauche de la barre."""
        icon = QIcon.fromTheme("edit-copy")
        if icon.isNull():
            # Icône différente de l'icône disquette/enregistrer : deux feuilles / vue détaillée selon le style Qt.
            icon = QApplication.style().standardIcon(QStyle.SP_FileDialogDetailedView)
        copy_action = QAction(icon, "Copier", self)
        copy_action.setToolTip("Copier le graphique dans le presse-papiers")
        copy_action.triggered.connect(self.copy_figure_to_clipboard)

        actions = self.actions()
        before_action = None
        # Matplotlib place généralement un séparateur après Home/Back/Forward.
        # On insère Copier juste avant ce premier séparateur pour le garder à gauche avec les autres icônes.
        for action in actions:
            if action.isSeparator():
                before_action = action
                break
        if before_action is not None:
            self.insertAction(before_action, copy_action)
        elif actions:
            self.insertAction(actions[0], copy_action)
        else:
            self.addAction(copy_action)

    def copy_figure_to_clipboard(self):
        try:
            if self.canvas is None or self.canvas.figure is None or not self.canvas.figure.axes:
                QMessageBox.information(self, "Copier graphique", "Aucun graphique disponible à copier.")
                return
            tmp = os.path.join(tempfile.gettempdir(), "statpro_graphique_clipboard.png")
            self.canvas.figure.savefig(tmp, dpi=160, bbox_inches="tight")
            pixmap = QPixmap(tmp)
            if pixmap.isNull():
                QMessageBox.warning(self, "Copier graphique", "Impossible de préparer l'image du graphique.")
                return
            QApplication.clipboard().setPixmap(pixmap)
            parent = self.parent()
            while parent is not None and not hasattr(parent, "status_bar"):
                parent = parent.parent()
            if parent is not None and hasattr(parent, "status_bar"):
                parent.status_bar.showMessage("Graphique copié dans le presse-papiers")
        except Exception as e:
            QMessageBox.critical(self, "Erreur", f"Impossible de copier le graphique :\n{e}")

    def save_figure(self, *args, **kwargs):
        try:
            file_filters = (
                "PNG Image (*.png);;"
                "PDF Document (*.pdf);;"
                "SVG Image (*.svg);;"
                "JPEG Image (*.jpg *.jpeg)"
            )
            filepath, _ = QFileDialog.getSaveFileName(
                self,
                "Enregistrer le graphique",
                "graphique.png",
                file_filters
            )
            if filepath:
                self.canvas.figure.savefig(filepath)
                QMessageBox.information(self, "Succès", f"Graphique enregistré :\n{filepath}")
        except Exception as e:
            QMessageBox.critical(self, "Erreur", f"Impossible de sauvegarder le graphique :\n{e}")


class VerticalNavTabs(QWidget):
    """Menu vertical à gauche compatible avec les usages principaux de QTabWidget.

    Les éléments sont regroupés par catégories avec des séparateurs visuels.
    L'ordre d'affichage du menu peut donc être différent de l'ordre de création des pages,
    mais chaque ligne pointe vers le bon index du QStackedWidget.
    """

    currentChanged = pyqtSignal(int)

    CATEGORY_ORDER = [
        "Données",
        "Distributions",
        "Qualité procédé",
        "Tests statistiques",
        "Graphiques",
        "Rapport",
    ]

    CATEGORY_MAP = {
        "Données": "Données",
        "Identification distribution": "Distributions",
        "Normalité": "Distributions",
        "Graphiques probabilité": "Distributions",
        "Capabilité": "Qualité procédé",
        "Cartes de contrôle": "Qualité procédé",
        "MSA / Gage R&R": "Qualité procédé",
        "Valeurs aberrantes": "Tests statistiques",
        "Test t": "Tests statistiques",
        "ANOVA": "Tests statistiques",
        "Corrélation": "Tests statistiques",
        "Régression": "Tests statistiques",
        "Boxplots": "Graphiques",
        "Rapport d'analyse": "Rapport",
        "MiniQual": "Rapport",
    }

    def __init__(self, parent=None):
        super().__init__(parent)
        self._titles = []
        self._widgets = []
        self._nav_rows_for_stack = {}
        self._building_nav = False

        layout = QHBoxLayout(self)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.setSpacing(6)

        self.nav = QListWidget()
        self.nav.setObjectName("verticalNav")
        self.set_theme("Clair")
        self.nav.setFixedWidth(215)
        self.nav.setSpacing(1)
        self.nav.setAlternatingRowColors(False)
        self.nav.currentRowChanged.connect(self._on_nav_changed)

        self.stack = QStackedWidget()
        self.stack.setObjectName("analysisStack")

        layout.addWidget(self.nav)
        layout.addWidget(self.stack, 1)

    def set_theme(self, theme):
        header_common = "padding: 7px 8px 3px 8px; margin-top: 8px; font-weight: bold;"
        if theme == "Sombre":
            self.nav.setStyleSheet(f"""
                QListWidget#verticalNav {{ background-color: #1f2327; color: #f0f0f0; border: 1px solid #4b5560; padding: 5px; outline: none; }}
                QListWidget#verticalNav::item {{ padding: 8px 10px; border-radius: 3px; margin: 1px; color: #f0f0f0; }}
                QListWidget#verticalNav::item:selected {{ background-color: #3d444b; color: #ffffff; border: 1px solid #7aa7d9; font-weight: bold; }}
                QListWidget#verticalNav::item:hover {{ background-color: #4f5b66; color: #ffffff; }}
                QListWidget#verticalNav::item:disabled {{ color: #9fb2c3; background: transparent; {header_common} }}
            """)
        elif theme == "Minitab-like":
            self.nav.setStyleSheet(f"""
                QListWidget#verticalNav {{ background-color: #edf5ff; color: #102a43; border: 1px solid #9bb7d4; padding: 5px; outline: none; }}
                QListWidget#verticalNav::item {{ padding: 8px 10px; border-radius: 3px; margin: 1px; color: #102a43; }}
                QListWidget#verticalNav::item:selected {{ background-color: #ffffff; color: #0b2239; border: 1px solid #4d90d9; font-weight: bold; }}
                QListWidget#verticalNav::item:hover {{ background-color: #d9eafc; color: #102a43; }}
                QListWidget#verticalNav::item:disabled {{ color: #42627d; background: transparent; {header_common} }}
            """)
        else:
            self.nav.setStyleSheet(f"""
                QListWidget#verticalNav {{ background-color: #f7f9fc; color: #1b1b1b; border: 1px solid #c5d4e3; padding: 5px; outline: none; }}
                QListWidget#verticalNav::item {{ padding: 8px 10px; border-radius: 3px; margin: 1px; color: #1b1b1b; }}
                QListWidget#verticalNav::item:selected {{ background-color: #d9eafc; color: #102a43; border: 1px solid #7aa7d9; font-weight: bold; }}
                QListWidget#verticalNav::item:hover {{ background-color: #eef6ff; color: #102a43; }}
                QListWidget#verticalNav::item:disabled {{ color: #5f6f7f; background: transparent; {header_common} }}
            """)

    def _category_for_title(self, title):
        return self.CATEGORY_MAP.get(title, "Autres")

    def _add_category_header(self, category):
        item = QListWidgetItem(f"── {category}")
        item.setFlags(Qt.NoItemFlags)
        item.setData(Qt.UserRole, None)
        font = item.font()
        font.setBold(True)
        font.setPointSize(max(8, font.pointSize() - 1))
        item.setFont(font)
        self.nav.addItem(item)

    def _rebuild_nav(self, selected_stack_index=None):
        self._building_nav = True
        self.nav.clear()
        self._nav_rows_for_stack = {}

        categories = list(self.CATEGORY_ORDER)
        existing_categories = {self._category_for_title(title) for title in self._titles}
        for category in sorted(existing_categories):
            if category not in categories:
                categories.append(category)

        for category in categories:
            stack_indices = [i for i, title in enumerate(self._titles) if self._category_for_title(title) == category]
            if not stack_indices:
                continue
            self._add_category_header(category)
            for stack_index in stack_indices:
                title = self._titles[stack_index]
                item = QListWidgetItem(title)
                item.setToolTip(title)
                item.setData(Qt.UserRole, stack_index)
                self.nav.addItem(item)
                self._nav_rows_for_stack[stack_index] = self.nav.count() - 1

        if selected_stack_index is None:
            selected_stack_index = self.stack.currentIndex()
        if selected_stack_index < 0 and self._widgets:
            selected_stack_index = 0
        row = self._nav_rows_for_stack.get(selected_stack_index)
        if row is not None:
            self.nav.setCurrentRow(row)
        self._building_nav = False

    def addTab(self, widget, title):
        index = len(self._widgets)
        self._widgets.append(widget)
        self._titles.append(title)
        self.stack.addWidget(widget)
        self._rebuild_nav(0 if index == 0 else self.stack.currentIndex())
        return index

    def _on_nav_changed(self, row):
        if self._building_nav:
            return
        item = self.nav.item(row)
        if item is None:
            return
        stack_index = item.data(Qt.UserRole)
        if stack_index is None:
            current = self.stack.currentIndex()
            if current in self._nav_rows_for_stack:
                self.nav.setCurrentRow(self._nav_rows_for_stack[current])
            return
        self.stack.setCurrentIndex(stack_index)
        self.currentChanged.emit(stack_index)

    def tabText(self, index):
        return self._titles[index] if 0 <= index < len(self._titles) else ""

    def count(self):
        return len(self._widgets)

    def currentIndex(self):
        return self.stack.currentIndex()

    def currentWidget(self):
        return self.stack.currentWidget()

    def setCurrentIndex(self, index):
        if 0 <= index < len(self._widgets):
            row = self._nav_rows_for_stack.get(index)
            if row is not None:
                self.nav.setCurrentRow(row)
            self.stack.setCurrentIndex(index)

    def setCurrentWidget(self, widget):
        idx = self.stack.indexOf(widget)
        if idx >= 0:
            self.setCurrentIndex(idx)

class StatisticalApp(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("StatPro - Analyse Statistique")
        self.resize(1400, 900)

        self._create_menu()
        self._create_toolbar()

        self.current_project_path = None
        self.current_theme = "Clair"
        self.analysis_report_sections = []
        
        self.cap_canvas = None
        self.norm_canvas = None
        self.out_canvas = None
        self.dist_canvas = None
        self.corr_canvas = None
        self.reg_canvas = None
        self.tt_canvas = None
        self.anova_canvas = None
        self.boxplot_canvas = None
        self.cc_canvas = None
        self.prob_canvas = None
        self.msa_canvas = None
        
        self._create_central_widget()

        self.status_bar = QStatusBar()
        self.setStatusBar(self.status_bar)
        self.status_bar.showMessage("Prêt - Saisissez des données dans le tableur")

    def _create_menu(self):
        menubar = self.menuBar()

        file_menu = menubar.addMenu("Fichier")
        file_menu.addAction("Nouveau projet", self._new_project)
        file_menu.addAction("Ouvrir projet...", self._open_project)
        file_menu.addAction("Enregistrer projet", self._save_project)
        file_menu.addAction("Enregistrer projet sous...", self._save_project_as)
        file_menu.addSeparator()
        file_menu.addAction("Importer CSV", self._import_csv)
        file_menu.addAction("Importer Excel", self._import_excel)
        file_menu.addSeparator()
        file_menu.addAction("Exporter résultats", self._export_results)
        file_menu.addAction("Exporter rapport complet PDF/DOCX...", self._export_full_report)
        file_menu.addAction("Exporter Excel multi-feuilles...", self._export_excel_workbook)
        file_menu.addSeparator()
        file_menu.addAction("Quitter", self.close)

        edit_menu = menubar.addMenu("Édition")
        edit_menu.addAction("Générer données exemple", self._generate_sample_data)
        edit_menu.addAction("Actualiser rapport d'analyse", self._refresh_report_panel)

        view_menu = menubar.addMenu("Affichage")
        view_menu.addAction("Thème clair", lambda: self._apply_theme("Clair"))
        view_menu.addAction("Thème sombre", lambda: self._apply_theme("Sombre"))
        view_menu.addAction("Thème Minitab-like", lambda: self._apply_theme("Minitab-like"))

        help_menu = menubar.addMenu("Aide")
        help_menu.addAction("À propos", self._show_about)

    def _create_toolbar(self):
        toolbar = QToolBar("Actions rapides")
        toolbar.setObjectName("mainToolbar")
        toolbar.setMovable(False)
        self.addToolBar(toolbar)

        new_action = toolbar.addAction(" Nouveau")
        new_action.triggered.connect(self._new_project)

        open_action = toolbar.addAction(" Ouvrir")
        open_action.triggered.connect(self._open_project)

        save_action = toolbar.addAction(" Enregistrer")
        save_action.triggered.connect(self._save_project)

        toolbar.addSeparator()

        report_action = toolbar.addAction(" Rapport")
        report_action.triggered.connect(self._refresh_report_panel)

    def _create_central_widget(self):
        central = QWidget()
        self.setCentralWidget(central)
        main_layout = QVBoxLayout(central)
        main_layout.setContentsMargins(5, 5, 5, 5)

        self.tabs = VerticalNavTabs()
        main_layout.addWidget(self.tabs)

        self._create_data_tab()
        self._create_distribution_tab()
        self._create_outliers_tab()
        self._create_normality_tab()
        self._create_capability_tab()
        self._create_correlation_tab()
        self._create_regression_tab()
        self._create_ttest_tab()
        self._create_anova_tab()
        self._create_boxplot_tab()
        self._create_control_charts_tab()
        self._create_probplot_tab()
        self._create_msa_tab()
        self._create_miniqual_tab()
        self._create_report_tab()

        self.tabs.currentChanged.connect(self._on_tab_changed)
        self._refresh_all_combos()

    def _on_tab_changed(self, index):
        self._refresh_all_combos()
        if self.tabs.tabText(index) == "ANOVA":
            self._update_anova_combos()
        elif self.tabs.tabText(index) == "Boxplots":
            self._update_boxplot_combos()
        elif self.tabs.tabText(index) == "Corrélation" and hasattr(self, "corr_combo_layout"):
            self._update_correlation_combos()

    def _refresh_all_combos(self):
        combos = []

        for attr in dir(self):
            obj = getattr(self, attr)
            if isinstance(obj, QComboBox):
                if attr.endswith("_col_combo") or attr in (
                    "reg_y_combo", "reg_x_combo", "tt_col1_combo", "tt_col2_combo",
                    "msa_part_combo", "msa_op_combo", "msa_meas_combo", "cc_col_combo",
                    "prob_col_combo", "corr_col_combo",
                ):
                    combos.append(obj)

        unique = []
        seen = set()
        for combo in combos:
            if id(combo) not in seen:
                seen.add(id(combo))
                unique.append(combo)

        cols = []
        if hasattr(self, "sheet"):
            for i in range(self.sheet.columnCount()):
                data = self.sheet.get_column_data(i)
                if data is not None and len(data) > 0:
                    cols.append(col_letter(i))

        if not cols:
            cols = ["(aucune donnée)"]

        for combo in unique:
            current = combo.currentText()
            combo.clear()
            combo.addItems(cols)
            if current in cols:
                combo.setCurrentText(current)

        if hasattr(self, "anova_col_combos") and hasattr(self, "anova_combo_layout"):
            self._update_anova_combos()
        if hasattr(self, "boxplot_col_combos") and hasattr(self, "boxplot_combo_layout"):
            self._update_boxplot_combos()
        if hasattr(self, "corr_col_combos") and hasattr(self, "corr_combo_layout"):
            self._update_correlation_combos()

    def _create_data_tab(self):
        tab = QWidget()
        self.tabs.addTab(tab, "Données")
        layout = QVBoxLayout(tab)
        layout.setContentsMargins(2, 2, 2, 2)

        info_label = QLabel(
            "Données — saisissez ou collez vos mesures dans le tableur. "
            "Sélectionnez ensuite les colonnes dans les onglets d'analyse. "
            "Les colonnes sont nommées A, B, C…"
        )
        info_label.setStyleSheet("""
            QLabel {
                background-color: #fff3b0;
                color: #1b1b1b;
                padding: 8px 10px;
                border: 1px solid #b49b00;
                border-radius: 3px;
                font-weight: bold;
                font-size: 12px;
            }
        """)
        info_label.setMinimumHeight(34)
        info_label.setWordWrap(True)
        layout.addWidget(info_label)

        data_splitter = QSplitter(Qt.Vertical)
        data_splitter.setChildrenCollapsible(False)
        layout.addWidget(data_splitter, stretch=1)

        self.sheet = DataSheet()
        self.sheet.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Expanding)
        data_splitter.addWidget(self.sheet)

        stats_group = QGroupBox("Statistiques descriptives - Colonne sélectionnée")
        stats_group.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Preferred)
        stats_layout = QVBoxLayout(stats_group)

        self.stats_text = QTextEdit()
        self.stats_text.setReadOnly(True)
        self.stats_text.setFont(QFont("Courier", 10))
        self.stats_text.setMinimumHeight(90)
        self.stats_text.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Expanding)
        self.stats_text.setLineWrapMode(QTextEdit.NoWrap)
        stats_layout.addWidget(self.stats_text)

        data_splitter.addWidget(stats_group)
        data_splitter.setStretchFactor(0, 5)
        data_splitter.setStretchFactor(1, 1)
        data_splitter.setSizes([650, 180])

        btn_frame = QWidget()
        btn_layout = QHBoxLayout(btn_frame)
        btn_layout.setContentsMargins(0, 0, 0, 0)
        refresh_btn = QPushButton(" Actualiser stats")
        refresh_btn.clicked.connect(self._update_data_display)
        clear_btn = QPushButton(" Effacer feuille")
        clear_btn.clicked.connect(self._clear_sheet)
        gen_btn = QPushButton(" Générer données")
        gen_btn.clicked.connect(self._generate_sample_data)
        btn_layout.addWidget(refresh_btn)
        btn_layout.addWidget(clear_btn)
        btn_layout.addWidget(gen_btn)
        btn_layout.addStretch()
        layout.addWidget(btn_frame)
        self.sheet.itemSelectionChanged.connect(self._on_selection_changed)

    def _on_selection_changed(self):
        self._update_data_display()

    def _get_active_data(self):
        data, col_name = self.sheet.get_selected_column_data()
        return data, col_name

    def _clear_sheet(self):
        if hasattr(self, "sheet"):
            self.sheet._push_undo()
        self.sheet.clearContents()
        self.stats_text.clear()
        self.status_bar.showMessage("Feuille effacée")

    def _update_data_display(self):
        data, col_name = self._get_active_data()
        if data is None or len(data) == 0:
            self.stats_text.setText("Aucune donnée dans la colonne sélectionnée.")
            return

        mean = np.mean(data)
        std = np.std(data, ddof=1)
        stats_info = (
            f"Colonne : {col_name}\n"
            f"N = {len(data)}\n"
            f"Moyenne = {mean:.6f}\n"
            f"Médiane = {np.median(data):.6f}\n"
            f"Écart-type = {std:.6f}\n"
            f"Min = {np.min(data):.6f}\n"
            f"Max = {np.max(data):.6f}\n"
            f"Asymétrie = {scipy_stats.skew(data):.4f}\n"
            f"Aplatissement = {scipy_stats.kurtosis(data):.4f}"
        )
        self.stats_text.setText(stats_info)
        self.status_bar.showMessage(f"Colonne {col_name} sélectionnée : {len(data)} valeurs")

    def _create_distribution_tab(self):
        tab = QWidget()
        self.tabs.addTab(tab, "Identification distribution")
        left, left_layout, right, self.dist_canvas = self._setup_analysis_layout(tab)

        col_row = QHBoxLayout()
        self.dist_col_combo = QComboBox()
        self.dist_col_combo.setMinimumWidth(80)
        col_row.addWidget(QLabel("Colonne :"))
        col_row.addWidget(self.dist_col_combo)
        refresh_btn = QPushButton("↻")
        refresh_btn.setMaximumWidth(30)
        refresh_btn.clicked.connect(self._refresh_all_combos)
        col_row.addWidget(refresh_btn)
        left_layout.addLayout(col_row)

        params_group = QGroupBox("Paramètres")
        params_layout = QFormLayout(params_group)

        self.dist_alpha = QDoubleSpinBox()
        self.dist_alpha.setRange(0.001, 0.5)
        self.dist_alpha.setSingleStep(0.005)
        self.dist_alpha.setValue(0.05)
        self._tip(self.dist_alpha, "Seuil de significativité α. Si p-value > α, la distribution est considérée comme un bon ajustement aux données.")
        params_layout.addRow("Seuil α :", self.dist_alpha)

        dists_group = QGroupBox("Distributions à tester")
        dists_layout = QVBoxLayout(dists_group)

        self.dist_checks = {}
        dist_desc = [
            ("Normal", "Loi gaussienne. Symétrique, définie par moyenne et écart-type. La plus courante."),
            ("Log-Normale", "Le logarithme des données suit une loi normale. Pour données strictement positives et asymétriques."),
            ("Weibull (2P)", "Flexible pour fiabilité/durée de vie. Shape + scale. Inclut exponentielle comme cas particulier."),
            ("Exponentielle", "Modélise les temps entre événements. Cas particulier de Weibull avec shape=1."),
            ("Gamma", "Généralisation de l'exponentielle. Pour temps d'attente de k événements."),
            ("Logistique", "Similaire à normale mais avec queues plus lourdes."),
            ("Gumbel (max)", "Pour valeurs extrêmes (maxima). Utilisée en hydrologie et fiabilité."),
            ("Cauchy", "Queue très lourde, pas de moyenne ni variance définies. Pour données avec outliers fréquents."),
            ("Rayleigh", "Pour magnitude de vecteurs 2D. Cas particulier de Weibull."),
            ("Uniforme", "Toutes les valeurs equally probables entre min et max."),
            ("Student-t", "Similaire à normale mais queue plus lourde. Pour petits échantillons."),
            ("Laplace", "Double exponentielle. Plus de kurtosis que la normale. Pour données avec pics centraux."),
        ]
        for name, desc in dist_desc:
            cb = QCheckBox(name)
            cb.setChecked(True)
            cb.setToolTip(desc)
            self.dist_checks[name] = cb
            dists_layout.addWidget(cb)

        params_layout.addRow(dists_group)
        left_layout.addWidget(params_group)

        calc_btn = QPushButton(" Identifier")
        calc_btn.setMinimumHeight(40)
        calc_btn.setStyleSheet("font-weight: bold; font-size: 14px;")
        calc_btn.clicked.connect(self._run_distribution)
        left_layout.addWidget(calc_btn)

        result_group = QGroupBox("Résultats (classés par p-value)")
        result_layout = QVBoxLayout(result_group)
        self.dist_result_text = QTextEdit()
        self.dist_result_text.setReadOnly(True)
        self.dist_result_text.setFont(QFont("Courier", 10))
        result_layout.addWidget(self.dist_result_text)
        left_layout.addWidget(result_group)

    def _run_distribution(self):
        data, col_name = self._get_combo_data(self.dist_col_combo)
        if data is None:
            QMessageBox.warning(self, "Attention", "Sélectionnez une colonne")
            return

        try:
            alpha = self.dist_alpha.value()

            dist_map = {
                "Normal": ("norm", lambda d: scipy_stats.norm.fit(d)),
                "Log-Normale": ("lognorm", lambda d: scipy_stats.lognorm.fit(d, floc=0)),
                "Weibull (2P)": ("weibull_min", lambda d: scipy_stats.weibull_min.fit(d, floc=0)),
                "Exponentielle": ("expon", lambda d: scipy_stats.expon.fit(d)),
                "Gamma": ("gamma", lambda d: scipy_stats.gamma.fit(d, floc=0)),
                "Logistique": ("logistic", lambda d: scipy_stats.logistic.fit(d)),
                "Gumbel (max)": ("gumbel_r", lambda d: scipy_stats.gumbel_r.fit(d)),
                "Cauchy": ("cauchy", lambda d: scipy_stats.cauchy.fit(d)),
                "Rayleigh": ("rayleigh", lambda d: scipy_stats.rayleigh.fit(d, floc=0)),
                "Uniforme": ("uniform", lambda d: scipy_stats.uniform.fit(d)),
                "Student-t": ("t", lambda d: scipy_stats.t.fit(d)),
                "Laplace": ("laplace", lambda d: scipy_stats.laplace.fit(d)),
            }

            results = []
            for display_name, cb in self.dist_checks.items():
                if cb.isChecked() and display_name in dist_map:
                    dist_name, fit_func = dist_map[display_name]
                    try:
                        params = fit_func(data)
                        loc = params[-2] if len(params) >= 2 else 0
                        scale = params[-1] if len(params) >= 1 else 1
                        shape_params = params[:-2] if len(params) > 2 else ()

                        ks_stat, ks_pvalue = scipy_stats.kstest(
                            data, dist_name, args=params
                        )
                        results.append({
                            "name": display_name,
                            "dist_name": dist_name,
                            "params": params,
                            "ks_stat": ks_stat,
                            "p_value": ks_pvalue,
                            "passes": ks_pvalue > alpha,
                        })
                    except Exception:
                        results.append({
                            "name": display_name,
                            "dist_name": dist_name,
                            "params": None,
                            "ks_stat": None,
                            "p_value": None,
                            "passes": False,
                            "error": True,
                        })

            results.sort(key=lambda r: r["p_value"] if r["p_value"] is not None else -1, reverse=True)

            lines = ["=" * 60, "IDENTIFICATION DE DISTRIBUTION", "=" * 60]
            lines.append(f"\nColonne : {col_name}")
            lines.append(f"N = {len(data)}")
            lines.append(f"Moyenne = {np.mean(data):.6f}")
            lines.append(f"Écart-type = {np.std(data, ddof=1):.6f}")
            lines.append(f"Seuil α = {alpha}")
            lines.append("\n" + "-" * 60)
            lines.append(f"{'Rang':<6} {'Distribution':<18} {'KS Stat':<12} {'p-value':<12} {'Ajusté?':<10}")
            lines.append("-" * 60)

            best_dist = results[0] if results else None
            for i, r in enumerate(results):
                if r.get("error"):
                    lines.append(f"{i+1:<6} {r['name']:<18} {'ERREUR':<12} {'-':<12} {'NON':<10}")
                else:
                    normal_str = "OUI" if r["passes"] else "NON"
                    lines.append(f"{i+1:<6} {r['name']:<18} {r['ks_stat']:<12.6f} {r['p_value']:<12.6f} {normal_str:<10}")

            lines.append("-" * 60)
            if best_dist and not best_dist.get("error"):
                lines.append(f"\nMeilleure distribution : {best_dist['name']}")
                lines.append(f"  p-value = {best_dist['p_value']:.6f}")
                lines.append(f"  Statistique KS = {best_dist['ks_stat']:.6f}")
                if best_dist["params"] is not None:
                    lines.append(f"  Paramètres : {best_dist['params']}")
            lines.append("=" * 60)

            self.dist_result_text.setText("\n".join(lines))
            self._plot_distribution(data, results)
            self.status_bar.showMessage("Identification de distribution terminée")

        except Exception as e:
            QMessageBox.critical(self, "Erreur", f"Erreur lors de l'identification :\n{e}")

    def _plot_distribution(self, data, results):
        self.dist_canvas.fig.clear()

        valid_results = [r for r in results if r["p_value"] is not None and not r.get("error")]
        n = len(valid_results)
        if n == 0:
            ax = self.dist_canvas.fig.add_subplot(111)
            ax.text(0.5, 0.5, "Aucune distribution valide", ha="center", va="center", fontsize=14)
            self.dist_canvas.draw()
            return

        cols = 3
        rows = (n + cols - 1) // cols

        self.dist_canvas.fig.set_size_inches(14, rows * 5)

        gs = self.dist_canvas.fig.add_gridspec(
            rows * 2, cols,
            left=0.06, right=0.96, top=0.96, bottom=0.04,
            wspace=0.30, hspace=0.50,
        )

        data_min = np.min(data)
        data_max = np.max(data)

        for i, r in enumerate(valid_results):
            row = i // cols
            col = i % cols

            ax_hist = self.dist_canvas.fig.add_subplot(gs[row * 2, col])
            ax_prob = self.dist_canvas.fig.add_subplot(gs[row * 2 + 1, col])

            bins = max(10, int(np.sqrt(len(data))))
            ax_hist.hist(data, bins=bins, density=True, alpha=0.6, color="lightblue", edgecolor="black")

            if r["params"] is not None:
                x = np.linspace(data_min, data_max, 300)
                dist_obj = getattr(scipy_stats, r["dist_name"])
                pdf_vals = dist_obj.pdf(x, *r["params"])
                ax_hist.plot(x, pdf_vals, "r-", linewidth=2)

            ax_hist.set_title(f"{r['name']}  (KS={r['ks_stat']:.4f}, p={r['p_value']:.4f})")
            ax_hist.grid(True, alpha=0.3)

            try:
                if r["dist_name"] == "lognorm":
                    valid = data[data > 0]
                    scipy_stats.probplot(np.log(valid), dist="norm", plot=ax_prob)
                else:
                    scipy_stats.probplot(data, dist=r["dist_name"], plot=ax_prob)
            except Exception:
                ax_prob.text(0.5, 0.5, "ProbPlot indisponible", ha="center", va="center", fontsize=8)
            ax_prob.grid(True, alpha=0.3)

        # Ne pas appeler tight_layout() ici : la figure utilise déjà add_gridspec()
        # avec des marges explicites (left/right/top/bottom). tight_layout peut alors
        # générer un UserWarning avec certains axes. On ajuste simplement les marges.
        self.dist_canvas.fig.subplots_adjust(
            left=0.06, right=0.96, top=0.96, bottom=0.04,
            wspace=0.30, hspace=0.50,
        )
        self.dist_canvas.draw()

    def _import_csv(self):
        filepath, _ = QFileDialog.getOpenFileName(self, "Importer CSV", "", "CSV files (*.csv);;All files (*)")
        if not filepath:
            return
        try:
            df = pd.read_csv(filepath)
            numeric_cols = df.select_dtypes(include=[np.number])
            if len(numeric_cols.columns) == 0:
                QMessageBox.critical(self, "Erreur", "Aucune colonne numérique trouvée")
                return
            self.sheet.load_from_dataframe(numeric_cols)
            self.status_bar.showMessage(f"Données importées : {filepath}")
            self._update_data_display()
        except Exception as e:
            QMessageBox.critical(self, "Erreur", f"Impossible de lire le fichier :\n{e}")

    def _import_excel(self):
        filepath, _ = QFileDialog.getOpenFileName(self, "Importer Excel", "", "Excel files (*.xlsx *.xls);;All files (*)")
        if not filepath:
            return
        try:
            df = pd.read_excel(filepath)
            numeric_cols = df.select_dtypes(include=[np.number])
            if len(numeric_cols.columns) == 0:
                QMessageBox.critical(self, "Erreur", "Aucune colonne numérique trouvée")
                return
            self.sheet.load_from_dataframe(numeric_cols)
            self.status_bar.showMessage(f"Données importées : {filepath}")
            self._update_data_display()
        except Exception as e:
            QMessageBox.critical(self, "Erreur", f"Impossible de lire le fichier :\n{e}")

    def _generate_sample_data(self):
        dialog = GenerateDataDialog(self)
        if dialog.exec_() == QDialog.Accepted:
            try:
                data = dialog.generate()
                col_letter_selected = "A"
                selected = self.sheet.selectionModel().selectedColumns()
                if selected:
                    col_idx = list(selected)[0].column()
                    col_letter_selected = col_letter(col_idx)

                self.sheet.load_from_array(data, start_col=ord(col_letter_selected[0]) - 65)
                dist = dialog.dist_combo.currentText()
                n = dialog.n_spin.value()
                self.status_bar.showMessage(f"Données générées en {col_letter_selected} : {dist} ({n} valeurs)")
                self._update_data_display()
            except Exception as e:
                QMessageBox.critical(self, "Erreur", str(e))

    def _get_data_from_sheet(self):
        data, col_name = self._get_active_data()
        if data is None or len(data) == 0:
            QMessageBox.warning(self, "Attention", "Aucune donnée.\nSaisissez des données dans le tableur ou sélectionnez une colonne.")
            return None, None
        return data, col_name

    def _setup_analysis_layout(self, tab):
        layout = QHBoxLayout(tab)
        layout.setContentsMargins(5, 5, 5, 5)

        splitter = QSplitter(Qt.Horizontal)
        splitter.setChildrenCollapsible(False)

        left = QWidget()
        left.setMinimumWidth(300)
        left.setMaximumWidth(500)
        left_layout = QVBoxLayout(left)
        left_layout.setContentsMargins(0, 0, 0, 0)

        right = QWidget()
        right_layout = QVBoxLayout(right)
        right_layout.setContentsMargins(0, 0, 0, 0)

        canvas = MplCanvas(right, width=7, height=8)
        right_layout.addWidget(canvas)
        toolbar = SafeNavigationToolbar(canvas, right)
        toolbar.setObjectName("plotToolbar")
        toolbar.setIconSize(toolbar.iconSize())
        right_layout.addWidget(toolbar)


        splitter.addWidget(left)
        splitter.addWidget(right)
        splitter.setSizes([350, 700])
        layout.addWidget(splitter)

        return left, left_layout, right, canvas

    def _create_capability_tab(self):
        tab = QWidget()
        self.tabs.addTab(tab, "Capabilité")
        left, left_layout, right, self.cap_canvas = self._setup_analysis_layout(tab)

        col_row = QHBoxLayout()
        self.cap_col_combo = QComboBox()
        self.cap_col_combo.setMinimumWidth(80)
        col_row.addWidget(QLabel("Colonne :"))
        col_row.addWidget(self.cap_col_combo)
        refresh_btn = QPushButton("↻")
        refresh_btn.setMaximumWidth(30)
        refresh_btn.clicked.connect(self._refresh_all_combos)
        col_row.addWidget(refresh_btn)
        left_layout.addLayout(col_row)

        params_group = QGroupBox("Paramètres")
        params_layout = QFormLayout(params_group)

        self.cap_lsl = QDoubleSpinBox()
        self.cap_lsl.setRange(-1e9, 1e9)
        self.cap_lsl.setSpecialValueText("Aucune")
        self._tip(self.cap_lsl, "Limite inférieure de spécification (LSL). Valeur minimale acceptable du processus. Si aucune, l'analyse est unilatérale supérieure.")
        params_layout.addRow("LSL (limite inférieure) :", self.cap_lsl)

        self.cap_usl = QDoubleSpinBox()
        self.cap_usl.setRange(-1e9, 1e9)
        self.cap_usl.setSpecialValueText("Aucune")
        self._tip(self.cap_usl, "Limite supérieure de spécification (USL). Valeur maximale acceptable du processus. Si aucune, l'analyse est unilatérale inférieure.")
        params_layout.addRow("USL (limite supérieure) :", self.cap_usl)

        self.cap_target = QDoubleSpinBox()
        self.cap_target.setRange(-1e9, 1e9)
        self.cap_target.setSpecialValueText("Aucune")
        self._tip(self.cap_target, "Valeur cible idéale du processus. Utilisée pour calculer Cpm, qui pénalise les écarts à la cible même dans les limites.")
        params_layout.addRow("Cible (Target) :", self.cap_target)

        self.cap_subgroup = QSpinBox()
        self.cap_subgroup.setRange(1, 100)
        self.cap_subgroup.setValue(1)
        self._tip(self.cap_subgroup, "Taille du sous-groupe pour l'estimation σ intra. 1 = individus (MR). >1 = sous-groupes rationnels.")
        params_layout.addRow("Taille sous-groupe :", self.cap_subgroup)

        self.cap_method = QComboBox()
        self.cap_method.addItems(["pooled", "rbar", "sbar"])
        self.cap_method.setCurrentText("pooled")
        self._tip(self.cap_method, "Méthode d'estimation de σ intra : pooled (écart-type groupé), rbar (R̄/d2), sbar (S̄/c4). Pooled est recommandé pour sous-groupes inégaux.")
        params_layout.addRow("Méthode estimation σ :", self.cap_method)

        self.cap_dist = QComboBox()
        self.cap_dist.addItems([
            "Normal", "Log-Normale", "Weibull (2P)", "Exponentielle",
            "Gamma", "Logistique", "Gumbel (max)", "Cauchy",
            "Rayleigh", "Uniforme", "Student-t", "Laplace",
        ])
        self._tip(self.cap_dist, "Distribution supposée des données. Normal = indices classiques. Autres = indices basés sur les percentiles (Pp/Ppk).")
        params_layout.addRow("Distribution :", self.cap_dist)

        left_layout.addWidget(params_group)

        calc_btn = QPushButton(" Calculer")
        calc_btn.setMinimumHeight(40)
        calc_btn.setStyleSheet("font-weight: bold; font-size: 14px;")
        calc_btn.clicked.connect(self._run_capability)
        left_layout.addWidget(calc_btn)

        result_group = QGroupBox("Résultats")
        result_layout = QVBoxLayout(result_group)
        self.cap_result_text = QTextEdit()
        self.cap_result_text.setReadOnly(True)
        self.cap_result_text.setFont(QFont("Courier", 10))
        result_layout.addWidget(self.cap_result_text)
        left_layout.addWidget(result_group)

    def _run_capability(self):
        data, col_name = self._get_combo_data(self.cap_col_combo)
        if data is None:
            QMessageBox.warning(self, "Attention", "Sélectionnez une colonne")
            return

        lsl = None if self.cap_lsl.value() == self.cap_lsl.minimum() else self.cap_lsl.value()
        usl = None if self.cap_usl.value() == self.cap_usl.minimum() else self.cap_usl.value()
        target = None if self.cap_target.value() == self.cap_target.minimum() else self.cap_target.value()

        if lsl is None and usl is None:
            QMessageBox.warning(self, "Attention", "Spécifiez au moins une limite (LSL ou USL)")
            return

        try:
            subgroup = self.cap_subgroup.value()
            method = self.cap_method.currentText()
            distribution = self.cap_dist.currentText()
            is_normal = (distribution == "Normal")

            analysis = CapabilityAnalysis(data, usl=usl, lsl=lsl, target=target,
                                         subgroup_size=subgroup, estimation_method=method)
            results = analysis.get_results()

            non_normal_info = None
            if not is_normal:
                dist_map = {
                    "Log-Normale": ("lognorm", lambda d: scipy_stats.lognorm.fit(d, floc=0)),
                    "Weibull (2P)": ("weibull_min", lambda d: scipy_stats.weibull_min.fit(d, floc=0)),
                    "Exponentielle": ("expon", lambda d: scipy_stats.expon.fit(d)),
                    "Gamma": ("gamma", lambda d: scipy_stats.gamma.fit(d, floc=0)),
                    "Logistique": ("logistic", lambda d: scipy_stats.logistic.fit(d)),
                    "Gumbel (max)": ("gumbel_r", lambda d: scipy_stats.gumbel_r.fit(d)),
                    "Cauchy": ("cauchy", lambda d: scipy_stats.cauchy.fit(d)),
                    "Rayleigh": ("rayleigh", lambda d: scipy_stats.rayleigh.fit(d, floc=0)),
                    "Uniforme": ("uniform", lambda d: scipy_stats.uniform.fit(d)),
                    "Student-t": ("t", lambda d: scipy_stats.t.fit(d)),
                    "Laplace": ("laplace", lambda d: scipy_stats.laplace.fit(d)),
                }
                dist_info = dist_map.get(distribution)
                if dist_info:
                    dist_name, fit_func = dist_info
                    params = fit_func(data)
                    dist_obj = getattr(scipy_stats, dist_name)
                    p00135 = dist_obj.ppf(0.00135, *params)
                    p50 = dist_obj.ppf(0.5, *params)
                    p99865 = dist_obj.ppf(0.99865, *params)

                if usl is not None and lsl is not None:
                    pp_nn = (usl - lsl) / (p99865 - p00135)
                    ppu_nn = (usl - p50) / (p99865 - p50)
                    ppl_nn = (p50 - lsl) / (p50 - p00135)
                    ppk_nn = min(ppu_nn, ppl_nn)
                elif usl is not None:
                    pp_nn = None
                    ppl_nn = None
                    ppu_nn = (usl - p50) / (p99865 - p50)
                    ppk_nn = ppu_nn
                else:
                    pp_nn = None
                    ppu_nn = None
                    ppl_nn = (p50 - lsl) / (p50 - p00135)
                    ppk_nn = ppl_nn

                non_normal_info = {
                    "distribution": distribution,
                    "dist_name": dist_name,
                    "params": params,
                    "p00135": p00135,
                    "p50": p50,
                    "p99865": p99865,
                    "pp": pp_nn,
                    "ppu": ppu_nn,
                    "ppl": ppl_nn,
                    "ppk": ppk_nn,
                }
                results["pp_non_normal"] = pp_nn
                results["ppu_non_normal"] = ppu_nn
                results["ppl_non_normal"] = ppl_nn
                results["ppk_non_normal"] = ppk_nn

            lines = ["=" * 60, "ANALYSE DE CAPABILITÉ", "=" * 60]
            lines.append(f"\nColonne : {col_name}")
            lines.append(f"Distribution : {distribution}")
            lines.append(f"N = {results['n']}")
            lines.append(f"Moyenne = {results['mean']:.6f}")
            lines.append(f"Écart-type (intra) = {results['std_within']:.6f}")
            lines.append(f"Écart-type (global) = {results['std_overall']:.6f}")
            lines.append(f"\nLimites :")
            if lsl is not None:
                lines.append(f"  LSL = {lsl:.6f}")
            if usl is not None:
                lines.append(f"  USL = {usl:.6f}")
            if target is not None:
                lines.append(f"  Cible = {target:.6f}")

            if non_normal_info:
                lines.append(f"\nPercentiles ({distribution}) :")
                lines.append(f"  P(0.135%) = {non_normal_info['p00135']:.6f}")
                lines.append(f"  P(50%)    = {non_normal_info['p50']:.6f}")
                lines.append(f"  P(99.865%)= {non_normal_info['p99865']:.6f}")

            lines.append(f"\nIndice intra (dans) :")
            if "cp" in results:
                lines.append(f"  Cp  = {results['cp']:.2f}")
            if "cpk" in results:
                lines.append(f"  Cpk = {results['cpk']:.2f}")
            if "cpl" in results:
                lines.append(f"  Cpl = {results['cpl']:.2f}")
            if "cpu" in results:
                lines.append(f"  Cpu = {results['cpu']:.2f}")

            lines.append(f"\nIndice global :")
            if non_normal_info:
                if non_normal_info["pp"] is not None:
                    lines.append(f"  Pp  = {non_normal_info['pp']:.2f} (non-normal)")
                if non_normal_info["ppk"] is not None:
                    lines.append(f"  Ppk = {non_normal_info['ppk']:.2f} (non-normal)")
                if non_normal_info["ppl"] is not None:
                    lines.append(f"  Ppl = {non_normal_info['ppl']:.2f} (non-normal)")
                if non_normal_info["ppu"] is not None:
                    lines.append(f"  Ppu = {non_normal_info['ppu']:.2f} (non-normal)")
            else:
                if "pp" in results:
                    lines.append(f"  Pp  = {results['pp']:.2f}")
                if "ppk" in results:
                    lines.append(f"  Ppk = {results['ppk']:.2f}")
                if "ppl" in results:
                    lines.append(f"  Ppl = {results['ppl']:.2f}")
                if "ppu" in results:
                    lines.append(f"  Ppu = {results['ppu']:.2f}")

            if "cpm" in results:
                lines.append(f"\nCpm = {results['cpm']:.2f}")
            lines.append("=" * 60)
            self.cap_result_text.setText("\n".join(lines))

            self._plot_capability(analysis, results, distribution, data=data)
            self.status_bar.showMessage("Analyse de capabilité terminée")

        except Exception as e:
            QMessageBox.critical(self, "Erreur", f"Erreur lors du calcul :\n{e}")

    def _plot_capability(self, analysis, results, distribution="Normal", data=None):
        if self.cap_canvas is None or self.cap_canvas.fig is None:
            return
        self.cap_canvas.fig.clear()

        # Utiliser explicitement la colonne sélectionnée dans l'onglet Capabilité.
        if data is None:
            data, _ = self._get_combo_data(self.cap_col_combo)
        if data is None or len(data) == 0:
            return

        data = np.asarray(data, dtype=float)
        data = data[~np.isnan(data)]
        if len(data) == 0:
            return

        mean = results.get("mean", float(np.mean(data)))
        sigma_within = results.get("std_within", float(np.std(data, ddof=1)) if len(data) > 1 else 0)
        sigma_overall = results.get("std_overall", float(np.std(data, ddof=1)) if len(data) > 1 else 0)
        lsl = results.get("lsl")
        usl = results.get("usl")
        target = results.get("target")
        if target is None and hasattr(self, "cap_target"):
            target = None if self.cap_target.value() == self.cap_target.minimum() else self.cap_target.value()

        self.cap_canvas.fig.set_size_inches(14, 9)
        gs = self.cap_canvas.fig.add_gridspec(
            2, 2,
            left=0.06, right=0.96, top=0.94, bottom=0.08,
            wspace=0.25, hspace=0.30
        )
        ax1 = self.cap_canvas.fig.add_subplot(gs[0, 0])
        ax2 = self.cap_canvas.fig.add_subplot(gs[0, 1])
        ax3 = self.cap_canvas.fig.add_subplot(gs[1, 0])
        ax4 = self.cap_canvas.fig.add_subplot(gs[1, 1])

        # === 1) Distribution — reprise du style MiniQual ===
        dist_map = {
            "Normal": ("norm", scipy_stats.norm, "normale", lambda d: scipy_stats.norm.fit(d)),
            "Log-Normale": ("lognorm", scipy_stats.lognorm, "lognormale", lambda d: scipy_stats.lognorm.fit(d[d > 0], floc=0)),
            "Weibull (2P)": ("weibull_min", scipy_stats.weibull_min, "Weibull", lambda d: scipy_stats.weibull_min.fit(d[d > 0], floc=0)),
            # Pour une exponentielle décalée, ne pas forcer loc=0.
            "Exponentielle": ("expon", scipy_stats.expon, "exponentielle", lambda d: scipy_stats.expon.fit(d[d >= 0])),
            "Gamma": ("gamma", scipy_stats.gamma, "gamma", lambda d: scipy_stats.gamma.fit(d[d > 0], floc=0)),
            "Logistique": ("logistic", scipy_stats.logistic, "logistique", lambda d: scipy_stats.logistic.fit(d)),
            "Gumbel (max)": ("gumbel_r", scipy_stats.gumbel_r, "Gumbel", lambda d: scipy_stats.gumbel_r.fit(d)),
            "Cauchy": ("cauchy", scipy_stats.cauchy, "Cauchy", lambda d: scipy_stats.cauchy.fit(d)),
            "Rayleigh": ("rayleigh", scipy_stats.rayleigh, "Rayleigh", lambda d: scipy_stats.rayleigh.fit(d[d > 0], floc=0)),
            "Uniforme": ("uniform", scipy_stats.uniform, "uniforme", lambda d: scipy_stats.uniform.fit(d)),
            "Student-t": ("t", scipy_stats.t, "Student-t", lambda d: scipy_stats.t.fit(d)),
            "Laplace": ("laplace", scipy_stats.laplace, "Laplace", lambda d: scipy_stats.laplace.fit(d)),
        }
        dist_name, dist_obj, dist_label, fit_func = dist_map.get(distribution, dist_map["Normal"])

        bins = min(15, max(5, int(np.sqrt(len(data)))))
        ax1.hist(data, bins=bins, density=True, alpha=0.75, edgecolor="black")

        data_std = float(np.std(data, ddof=1)) if len(data) > 1 else 0.0
        xs = [float(np.min(data)), float(np.max(data))]
        if data_std > 0:
            xs += [float(mean - 4 * data_std), float(mean + 4 * data_std)]
        for spec in (lsl, usl, target):
            if spec is not None:
                xs.append(float(spec))
        x_min, x_max = min(xs), max(xs)
        if x_min == x_max:
            x_min -= 1
            x_max += 1
        xx = np.linspace(x_min, x_max, 300)

        try:
            fit_params = fit_func(data)
            yy = dist_obj.pdf(xx, *fit_params)
            if np.all(np.isfinite(yy)):
                ax1.plot(xx, yy, "r-", linewidth=2, label=f"Courbe {dist_label} estimée")
        except Exception:
            # Si la loi sélectionnée ne peut pas être ajustée, on conserve l'histogramme et les repères.
            pass

        if lsl is not None:
            ax1.axvline(lsl, color="green", linestyle="--", linewidth=2, label=f"LSL={lsl:.3g}")
        if usl is not None:
            ax1.axvline(usl, color="red", linestyle="--", linewidth=2, label=f"USL={usl:.3g}")
        if target is not None:
            ax1.axvline(target, color="purple", linestyle=":", linewidth=2, label=f"Cible={target:.3g}")
        ax1.axvline(mean, color="orange", linewidth=1.8, label=f"Moyenne={mean:.3g}")
        ax1.set_title("Distribution")
        ax1.legend(fontsize=8)
        ax1.grid(True, alpha=0.25)

        # === 2) Graphique de probabilité — conservé ===
        try:
            if distribution == "Log-Normale":
                valid = data[data > 0]
                if len(valid) > 0:
                    scipy_stats.probplot(np.log(valid), dist="norm", plot=ax2)
                else:
                    ax2.text(0.5, 0.5, "Données invalides", ha="center", va="center")
            else:
                scipy_stats.probplot(data, dist=dist_name, plot=ax2)
        except Exception:
            ax2.text(0.5, 0.5, "ProbPlot indisponible", ha="center", va="center")
        ax2.set_title("Graphique de probabilité")
        ax2.grid(True, alpha=0.3)

        # === 3) Run chart — conservé ===
        ax3.plot(range(len(data)), data, "b-", linewidth=0.8, marker=".", markersize=3)
        ax3.axhline(mean, color="green", linestyle="-", linewidth=1.5)
        if lsl is not None:
            ax3.axhline(lsl, color="red", linestyle="--", linewidth=1)
        if usl is not None:
            ax3.axhline(usl, color="red", linestyle="--", linewidth=1)
        ax3.set_title("Run chart")
        ax3.set_xlabel("Observation")
        ax3.grid(True, alpha=0.25)

        # === 4) Indices Cp/Cpk — remplacement des résidus par la vue MiniQual ===
        labels = []
        values = []
        for label, key in [
            ("Cp", "cp"),
            ("Cpk inf.", "cpl"),
            ("Cpk sup.", "cpu"),
            ("Cpk global", "cpk"),
        ]:
            val = results.get(key)
            if val is not None and np.isfinite(val):
                labels.append(label)
                values.append(float(val))

        if values:
            ax4.bar(labels, values)
            ax4.axhline(1.33, color="orange", linestyle="--", linewidth=1.5, label="Accept. 1.33")
            ax4.axhline(1.67, color="green", linestyle="--", linewidth=1.5, label="Excellent 1.67")
            y_max = max(values + [1.67]) * 1.20
            ax4.set_ylim(0, y_max if y_max > 0 else 1)
            for i, val in enumerate(values):
                ax4.text(i, val, f"{val:.2f}", ha="center", va="bottom", fontsize=9)
            ax4.legend(fontsize=8)
        else:
            ax4.text(0.5, 0.5, "Indices Cp/Cpk indisponibles", ha="center", va="center")
        ax4.set_title("Indices")
        ax4.grid(True, axis="y", alpha=0.25)

        try:
            self.cap_canvas.fig.tight_layout()
        except Exception:
            pass
        self.cap_canvas.draw()

    def _create_normality_tab(self):
        tab = QWidget()
        self.tabs.addTab(tab, "Normalité")
        left, left_layout, right, self.norm_canvas = self._setup_analysis_layout(tab)

        col_row = QHBoxLayout()
        self.norm_col_combo = QComboBox()
        self.norm_col_combo.setMinimumWidth(80)
        col_row.addWidget(QLabel("Colonne :"))
        col_row.addWidget(self.norm_col_combo)
        refresh_btn = QPushButton("↻")
        refresh_btn.setMaximumWidth(30)
        refresh_btn.clicked.connect(self._refresh_all_combos)
        col_row.addWidget(refresh_btn)
        left_layout.addLayout(col_row)

        params_group = QGroupBox("Paramètres")
        params_layout = QFormLayout(params_group)

        self.norm_alpha = QDoubleSpinBox()
        self.norm_alpha.setRange(0.001, 0.5)
        self.norm_alpha.setSingleStep(0.005)
        self.norm_alpha.setValue(0.05)
        self._tip(self.norm_alpha, "Seuil de significativité α. Si p-value < α, on rejette l'hypothèse de normalité. 0.05 = risque 5%.")
        params_layout.addRow("Seuil α :", self.norm_alpha)

        tests_group = QGroupBox("Tests à effectuer")
        tests_layout = QVBoxLayout(tests_group)

        self.norm_tests = {}
        test_names = [
            ("Shapiro-Wilk", "Test le plus puissant pour N < 5000. Hypothèse H0: les données suivent une loi normale."),
            ("Anderson-Darling", "Plus sensible aux écarts dans les queues de distribution que Shapiro-Wilk."),
            ("D'Agostino-Pearson", "Combine l'asymétrie (skewness) et l'aplatissement (kurtosis). Bon pour grands échantillons."),
            ("Jarque-Bera", "Test asymptotique basé sur skewness et kurtosis. Adapté aux grands échantillons (N > 50)."),
            ("Ryan-Joiner", "Similaire à Shapiro-Wilk. Mesure la corrélation entre les données et les quantiles normaux."),
        ]
        for name, tooltip in test_names:
            cb = QCheckBox(name)
            cb.setChecked(True)
            cb.setToolTip(tooltip)
            self.norm_tests[name] = cb
            tests_layout.addWidget(cb)

        khi_row = QHBoxLayout()
        cb_chi = QCheckBox("Khi-deux (χ²)")
        cb_chi.setChecked(True)
        cb_chi.setToolTip("Test d'ajustement basé sur la comparaison des fréquences observées vs attendues dans des classes.")
        self.norm_tests["Khi-deux (χ²)"] = cb_chi
        khi_row.addWidget(cb_chi)
        khi_row.addWidget(QLabel("ddl :"))
        self.norm_chi2_df = QSpinBox()
        self.norm_chi2_df.setRange(1, 100)
        self.norm_chi2_df.setValue(10)
        self.norm_chi2_df.setMaximumWidth(60)
        self._tip(self.norm_chi2_df, "Degrés de liberté pour le test du Khi-deux. Par défaut = nombre de classes - 1 - paramètres estimés.")
        khi_row.addWidget(self.norm_chi2_df)
        khi_row.addStretch()
        tests_layout.addLayout(khi_row)

        params_layout.addRow(tests_group)
        left_layout.addWidget(params_group)

        run_btn = QPushButton(" Exécuter tests")
        run_btn.setMinimumHeight(40)
        run_btn.setStyleSheet("font-weight: bold; font-size: 14px;")
        run_btn.clicked.connect(self._run_normality)
        left_layout.addWidget(run_btn)

        result_group = QGroupBox("Résultats")
        result_layout = QVBoxLayout(result_group)
        self.norm_result_text = QTextEdit()
        self.norm_result_text.setReadOnly(True)
        self.norm_result_text.setFont(QFont("Courier", 10))
        result_layout.addWidget(self.norm_result_text)
        left_layout.addWidget(result_group)

    def _run_normality(self):
        data, col_name = self._get_combo_data(self.norm_col_combo)
        if data is None:
            QMessageBox.warning(self, "Attention", "Sélectionnez une colonne")
            return

        try:
            alpha = self.norm_alpha.value()
            chi2_df = self.norm_chi2_df.value()
            tests = NormalityTests(data, alpha=alpha, chi2_df=chi2_df)
            full_results = tests.get_results()

            summary_lines = ["=" * 60, "TESTS DE NORMALITÉ", "=" * 60]
            summary_lines.append(f"\nColonne : {col_name}")
            summary_lines.append(f"N = {full_results['n']}")
            summary_lines.append(f"Moyenne = {full_results['mean']:.6f}")
            summary_lines.append(f"Écart-type = {full_results['std']:.6f}")
            summary_lines.append(f"Asymétrie = {full_results['skewness']:.4f}")
            summary_lines.append(f"Aplatissement = {full_results['kurtosis']:.4f}")
            summary_lines.append(f"\nSeuil α = {full_results['alpha']}")
            summary_lines.append("-" * 60)

            for name, cb in self.norm_tests.items():
                if cb.isChecked() and name in full_results["tests"]:
                    result = full_results["tests"][name]
                    summary_lines.append(f"\n{name} :")
                    summary_lines.append(f"  Statistique = {result['statistic']:.6f}")
                    if result.get("p_value") is not None:
                        summary_lines.append(f"  p-value = {result['p_value']:.6f}")
                        normal_str = "OUI" if result["normal"] else "NON"
                        summary_lines.append(f"  Normal = {normal_str}")
                    if result.get("df") is not None:
                        summary_lines.append(f"  Degrés de liberté = {result['df']}")
                    if "critical_value" in result and result["critical_value"] is not None:
                        summary_lines.append(f"  Valeur critique (α={alpha}) = {result['critical_value']:.6f}")
                    if "critical_95" in result:
                        summary_lines.append(f"  Critique (95%) = {result['critical_95']:.6f}")
                        normal_str = "OUI" if result["normal"] else "NON"
                        summary_lines.append(f"  Normal = {normal_str}")

            summary_lines.append("=" * 60)
            self.norm_result_text.setText("\n".join(summary_lines))

            self._plot_normality(data)
            self.status_bar.showMessage("Tests de normalité terminés")

        except Exception as e:
            QMessageBox.critical(self, "Erreur", f"Erreur lors des tests :\n{e}")

    def _plot_normality(self, data):
        self.norm_canvas.fig.clear()

        ax1 = self.norm_canvas.fig.add_subplot(221)
        ax2 = self.norm_canvas.fig.add_subplot(222)
        ax3 = self.norm_canvas.fig.add_subplot(223)
        ax4 = self.norm_canvas.fig.add_subplot(224)

        bins = max(10, int(np.sqrt(len(data))))
        ax1.hist(data, bins=bins, density=True, alpha=0.6, color="lightblue", edgecolor="black")
        x = np.linspace(np.min(data), np.max(data), 100)
        ax1.plot(x, scipy_stats.norm.pdf(x, np.mean(data), np.std(data, ddof=1)), "r-", linewidth=2)
        ax1.set_title("Histogramme avec courbe normale")

        scipy_stats.probplot(data, dist="norm", plot=ax2)
        ax2.set_title("QQ Plot")
        ax2.grid(True, alpha=0.3)

        ax3.boxplot(data, vert=True)
        ax3.set_title("Boîte à moustaches")

        ax4.hist(data, bins=bins, density=True, alpha=0.6, color="lightblue", edgecolor="black")
        from scipy.stats import gaussian_kde
        kde = gaussian_kde(data)
        x_kde = np.linspace(np.min(data), np.max(data), 200)
        ax4.plot(x_kde, kde(x_kde), "g-", linewidth=2, label="KDE")
        ax4.plot(x_kde, scipy_stats.norm.pdf(x_kde, np.mean(data), np.std(data, ddof=1)), "r--", linewidth=2, label="Normal")
        ax4.legend(fontsize=8)
        ax4.set_title("Densité KDE vs Normale")

        try:
            self.norm_canvas.fig.tight_layout()
        except Exception:
            pass
        self.norm_canvas.draw()

    def _create_outliers_tab(self):
        tab = QWidget()
        self.tabs.addTab(tab, "Valeurs aberrantes")
        left, left_layout, right, self.out_canvas = self._setup_analysis_layout(tab)

        col_row = QHBoxLayout()
        self.out_col_combo = QComboBox()
        self.out_col_combo.setMinimumWidth(80)
        col_row.addWidget(QLabel("Colonne :"))
        col_row.addWidget(self.out_col_combo)
        refresh_btn = QPushButton("↻")
        refresh_btn.setMaximumWidth(30)
        refresh_btn.clicked.connect(self._refresh_all_combos)
        col_row.addWidget(refresh_btn)
        left_layout.addLayout(col_row)

        params_group = QGroupBox("Paramètres")
        params_layout = QFormLayout(params_group)

        self.out_alpha = QDoubleSpinBox()
        self.out_alpha.setRange(0.001, 0.5)
        self.out_alpha.setSingleStep(0.005)
        self.out_alpha.setValue(0.05)
        self._tip(self.out_alpha, "Seuil de détection. Plus petit = moins de valeurs détectées comme aberrantes.")
        params_layout.addRow("Seuil α :", self.out_alpha)

        tests_group = QGroupBox("Méthodes")
        tests_layout = QVBoxLayout(tests_group)

        self.out_tests = {}
        test_desc = [
            ("IQR", "Méthode de l'intervalle interquartile. Valeurs < Q1-1.5×IQR ou > Q3+1.5×IQR. Robuste, ne suppose pas de normalité."),
            ("Z-score", "Écart à la moyenne en unités d'écart-type. Seuil typique : |Z| > 3. Suppose la normalité."),
            ("Z-score modifié", "Utilise la médiane et le MAD (Median Absolute Deviation). Plus robuste aux outliers que Z-score classique."),
            ("Grubbs", "Test statistique pour détecter un outlier à la fois. Basé sur la distribution normale. Pour N ≥ 7."),
            ("Dixon", "Test Q de Dixon. Pour petits échantillons (N < 30). Compare l'écart entre valeurs extrêmes."),
            ("Chauvenet", "Critère probabiliste. Rejette une valeur si la probabilité de l'observer est < 1/(2N)."),
            ("Tukey (extrêmes)", "Variante stricte de IQR avec facteur 3× au lieu de 1.5×. Détecte seulement les outliers extrêmes."),
        ]
        for name, desc in test_desc:
            cb = QCheckBox(name)
            cb.setChecked(True)
            cb.setToolTip(desc)
            self.out_tests[name] = cb
            tests_layout.addWidget(cb)

        params_layout.addRow(tests_group)
        left_layout.addWidget(params_group)

        detect_btn = QPushButton(" Détecter")
        detect_btn.setMinimumHeight(40)
        detect_btn.setStyleSheet("font-weight: bold; font-size: 14px;")
        detect_btn.clicked.connect(self._run_outliers)
        left_layout.addWidget(detect_btn)

        result_group = QGroupBox("Résultats")
        result_layout = QVBoxLayout(result_group)
        self.out_result_text = QTextEdit()
        self.out_result_text.setReadOnly(True)
        self.out_result_text.setFont(QFont("Courier", 10))
        result_layout.addWidget(self.out_result_text)
        left_layout.addWidget(result_group)

    def _run_outliers(self):
        data, col_name = self._get_combo_data(self.out_col_combo)
        if data is None:
            QMessageBox.warning(self, "Attention", "Sélectionnez une colonne")
            return

        try:
            alpha = self.out_alpha.value()
            detector = OutlierDetection(data, alpha=alpha)
            full_results = detector.get_results()

            summary_lines = ["=" * 60, "DÉTECTION DE VALEURS ABERRANTES", "=" * 60]
            summary_lines.append(f"\nColonne : {col_name}")
            summary_lines.append(f"N = {full_results['n']}")
            summary_lines.append(f"Moyenne = {full_results['mean']:.6f}")
            summary_lines.append(f"Médiane = {full_results['median']:.6f}")
            summary_lines.append(f"Écart-type = {full_results['std']:.6f}")
            summary_lines.append(f"Q1 = {full_results['q1']:.6f}")
            summary_lines.append(f"Q3 = {full_results['q3']:.6f}")
            summary_lines.append(f"IQR = {full_results['iqr']:.6f}")
            summary_lines.append(f"\nSeuil α = {full_results['alpha']}")
            summary_lines.append("-" * 60)

            method_map = {
                "IQR": "IQR (1.5×IQR)",
                "Z-score": "Z-score",
                "Z-score modifié": "Z-score modifié (MAD)",
                "Grubbs": "Grubbs",
                "Dixon": "Dixon",
                "Chauvenet": "Chauvenet",
                "Tukey (extrêmes)": "Tukey (3×IQR, extrêmes)",
            }

            for display_name, cb in self.out_tests.items():
                if cb.isChecked():
                    test_name = method_map.get(display_name, display_name)
                    if test_name in full_results["tests"]:
                        result = full_results["tests"][test_name]
                        summary_lines.append(f"\n{test_name} :")
                        summary_lines.append(f"  Valeurs détectées : {result['outlier_count']}")
                        if result["outlier_count"] > 0:
                            for val, idx in zip(result["outlier_values"], result["outlier_indices"]):
                                summary_lines.append(f"    [{idx}] Valeur = {val:.6f}")
                        else:
                            summary_lines.append(f"    Aucune valeur aberrante détectée")

            summary_lines.append("=" * 60)
            self.out_result_text.setText("\n".join(summary_lines))

            self._plot_outliers(full_results, detector, data)
            self.status_bar.showMessage("Détection de valeurs aberrantes terminée")

        except Exception as e:
            QMessageBox.critical(self, "Erreur", f"Erreur lors de la détection :\n{e}")

    def _plot_outliers(self, full_results, detector, data):
        self.out_canvas.fig.clear()

        ax1 = self.out_canvas.fig.add_subplot(311)
        ax2 = self.out_canvas.fig.add_subplot(312)
        ax3 = self.out_canvas.fig.add_subplot(313)

        all_outlier_indices = detector.get_outlier_indices()

        colors = ["red" if i in all_outlier_indices else "blue" for i in range(len(data))]

        ax1.scatter(range(len(data)), data, c=colors, alpha=0.6, s=10)
        ax1.axhline(np.mean(data), color="green", linestyle="--", label=f"Moy={np.mean(data):.2f}")
        ax1.axhline(np.mean(data) + 3 * np.std(data, ddof=1), color="red", linestyle=":", alpha=0.5)
        ax1.axhline(np.mean(data) - 3 * np.std(data, ddof=1), color="red", linestyle=":", alpha=0.5)
        ax1.legend(fontsize=7)
        ax1.set_title("Graphique des données (valeurs aberrantes en rouge)")
        ax1.set_ylabel("Valeur")

        bp = ax2.boxplot(data, vert=True, patch_artist=True, showfliers=True)
        bp["boxes"][0].set_facecolor("lightblue")
        ax2.set_title("Boîte à moustaches avec valeurs aberrantes")

        deviations = np.abs(data - np.median(data))
        ax3.scatter(range(len(data)), deviations, alpha=0.6, color="blue", s=10)
        median = np.median(data)
        mad = np.median(np.abs(data - median))
        if mad > 0:
            threshold = 3.5 / 0.6745 * mad
            ax3.axhline(threshold, color="red", linestyle="--", label="Seuil (3.5 MAD)")
            ax3.legend(fontsize=7)
        ax3.set_title("Écarts à la médiane")
        ax3.set_xlabel("Index")
        ax3.set_ylabel("Écart absolu")

        try:
            self.out_canvas.fig.tight_layout()
        except Exception:
            pass
        self.out_canvas.draw()

    def _tip(self, widget, text):
        widget.setToolTip(text)

    def _copy_text(self, widget):
        QApplication.clipboard().setText(widget.toPlainText())

    def _add_copy_button(self, layout, text_widget):
        btn = QPushButton(" Copier")
        btn.clicked.connect(lambda: self._copy_text(text_widget))
        layout.addWidget(btn)

    def _refresh_combos(self, *combos):
        cols = [col_letter(i) for i in range(NUM_COLS)]
        has_data = False
        for i in range(NUM_COLS):
            if self.sheet.get_column_data(i) is not None:
                has_data = True
                break
        if not has_data:
            cols = ["(aucune donnée)"]
        for combo in combos:
            current = combo.currentText()
            combo.clear()
            combo.addItems(cols)
            if current in cols:
                combo.setCurrentText(current)

    def _get_combo_data(self, combo):
        letter = combo.currentText()
        if not letter or letter == "(aucune donnée)":
            return None, None

        try:
            idx = 0
            for char in letter:
                if not char.isalpha():
                    return None, None
                idx = idx * 26 + (ord(char.upper()) - ord("A") + 1)
            idx -= 1

            if idx < 0 or idx >= self.sheet.columnCount():
                return None, None

            return self.sheet.get_column_data(idx), letter
        except Exception:
            return None, None

    def _get_all_numeric_columns(self):
        cols = {}
        for c in range(self.sheet.columnCount()):
            data = self.sheet.get_column_data(c)
            if data is not None and len(data) >= 2:
                cols[col_letter(c)] = data
        return cols

    # ===== CORRELATION =====
    def _create_correlation_tab(self):
        tab = QWidget()
        self.tabs.addTab(tab, "Corrélation")
        left, left_layout, right, self.corr_canvas = self._setup_analysis_layout(tab)

        params_group = QGroupBox("Paramètres")
        params_layout = QFormLayout(params_group)
        self.corr_method = QComboBox()
        self.corr_method.addItems(["Pearson", "Spearman", "Kendall"])
        self._tip(self.corr_method, "Pearson: linéaire, Spearman: monotone, Kendall: rangs concordants")
        params_layout.addRow("Méthode :", self.corr_method)
        self.corr_min_corr = QDoubleSpinBox()
        self.corr_min_corr.setRange(0, 1)
        self.corr_min_corr.setSingleStep(0.05)
        self.corr_min_corr.setValue(0.7)
        self._tip(self.corr_min_corr, "Seuil minimum pour mettre en évidence les corrélations fortes dans la matrice.")
        params_layout.addRow("Seuil corrélation forte :", self.corr_min_corr)
        self.corr_ngroups = QSpinBox()
        self.corr_ngroups.setRange(2, NUM_COLS)
        self.corr_ngroups.setValue(2)
        self._tip(self.corr_ngroups, "Nombre de colonnes à inclure dans la matrice de corrélation.")
        params_layout.addRow("Nombre de colonnes :", self.corr_ngroups)
        left_layout.addWidget(params_group)

        select_group = QGroupBox("Colonnes à corréler")
        select_layout = QVBoxLayout(select_group)
        self.corr_combo_container = QWidget()
        self.corr_combo_layout = QVBoxLayout(self.corr_combo_container)
        self.corr_combo_layout.setContentsMargins(0, 0, 0, 0)
        select_layout.addWidget(self.corr_combo_container)
        refresh_btn = QPushButton("↻ Actualiser les colonnes")
        refresh_btn.clicked.connect(self._refresh_all_combos)
        select_layout.addWidget(refresh_btn)
        left_layout.addWidget(select_group)
        self.corr_col_combos = []
        self._update_correlation_combos()
        self.corr_ngroups.valueChanged.connect(self._update_correlation_combos)

        calc_btn = QPushButton(" Calculer")
        calc_btn.setMinimumHeight(40)
        calc_btn.clicked.connect(self._run_correlation)
        left_layout.addWidget(calc_btn)
        result_group = QGroupBox("Résultats")
        result_layout = QVBoxLayout(result_group)
        self.corr_result_text = QTextEdit()
        self.corr_result_text.setReadOnly(True)
        self.corr_result_text.setFont(QFont("Courier", 10))
        result_layout.addWidget(self.corr_result_text)
        self._add_copy_button(result_layout, self.corr_result_text)
        left_layout.addWidget(result_group)

    def _update_correlation_combos(self):
        current_selections = []
        if hasattr(self, "corr_col_combos") and self.corr_col_combos:
            current_selections = [combo.currentText() for combo in self.corr_col_combos]
        if not hasattr(self, "corr_combo_layout"):
            return
        while self.corr_combo_layout.count():
            item = self.corr_combo_layout.takeAt(0)
            if item.widget():
                item.widget().deleteLater()
            elif item.layout():
                while item.layout().count():
                    child = item.layout().takeAt(0)
                    if child.widget():
                        child.widget().deleteLater()
        self.corr_col_combos = []
        cols = []
        if hasattr(self, "sheet"):
            for i in range(self.sheet.columnCount()):
                data = self.sheet.get_column_data(i)
                if data is not None and len(data) > 0:
                    cols.append(col_letter(i))
        if not cols:
            cols = ["(aucune donnée)"]
        n = self.corr_ngroups.value() if hasattr(self, "corr_ngroups") else 2
        for i in range(n):
            row = QHBoxLayout()
            label = QLabel(f"Colonne {i + 1} :")
            label.setMinimumWidth(80)
            row.addWidget(label)
            combo = QComboBox()
            combo.setMinimumWidth(100)
            combo.addItems(cols)
            if i < len(current_selections) and current_selections[i] in cols:
                combo.setCurrentText(current_selections[i])
            elif i < len(cols) and cols[0] != "(aucune donnée)":
                combo.setCurrentText(cols[i % len(cols)])
            row.addWidget(combo)
            row.addStretch()
            self.corr_combo_layout.addLayout(row)
            self.corr_col_combos.append(combo)

    def _run_correlation(self):
        if not hasattr(self, "corr_col_combos") or not self.corr_col_combos:
            QMessageBox.warning(self, "Attention", "Sélectionnez au moins 2 colonnes")
            return
        selected_labels = []
        selected_data = []
        for combo in self.corr_col_combos:
            data, label = self._get_combo_data(combo)
            if data is not None and label is not None:
                selected_labels.append(label)
                selected_data.append(data)
        if len(selected_data) < 2:
            QMessageBox.warning(self, "Attention", "Au moins 2 colonnes avec données sont requises")
            return
        if len(set(selected_labels)) != len(selected_labels):
            QMessageBox.warning(self, "Attention", "Sélectionnez des colonnes différentes pour la corrélation")
            return
        try:
            method = self.corr_method.currentText()
            min_corr = self.corr_min_corr.value()
            min_len = min(len(col) for col in selected_data)
            if min_len < 2:
                QMessageBox.warning(self, "Attention", "Chaque colonne doit contenir au moins 2 valeurs numériques")
                return
            data = np.column_stack([col[:min_len] for col in selected_data])
            analysis = CorrelationMatrix(data, labels=selected_labels, method=method.lower())
            self.corr_result_text.setText(analysis.get_summary())
            self._plot_correlation(data, selected_labels, method, min_corr)
            self.status_bar.showMessage(f"Corrélation terminée : {len(selected_labels)} colonnes, N={min_len}")
        except Exception as e:
            QMessageBox.critical(self, "Erreur", str(e))

    def _plot_correlation(self, data, labels, method, min_corr=0.7):
        self.corr_canvas.fig.clear()
        ax = self.corr_canvas.fig.add_subplot(111)
        df = pd.DataFrame(data, columns=labels)
        method_lower = method.lower()
        if method_lower not in ("pearson", "spearman", "kendall"):
            method_lower = "pearson"
        corr = df.corr(method=method_lower).values
        im = ax.imshow(corr, cmap="RdBu_r", vmin=-1, vmax=1)
        ax.set_xticks(range(len(labels)))
        ax.set_yticks(range(len(labels)))
        ax.set_xticklabels(labels, rotation=45, ha="right")
        ax.set_yticklabels(labels)
        for i in range(len(labels)):
            for j in range(len(labels)):
                value = corr[i, j]
                color = "white" if abs(value) > min_corr else "black"
                weight = "bold" if abs(value) > min_corr else "normal"
                ax.text(j, i, f"{value:.2f}", ha="center", va="center", color=color, fontsize=9, fontweight=weight)
        self.corr_canvas.fig.colorbar(im, ax=ax, label="Corrélation")
        ax.set_title(f"Matrice de corrélation ({method})")
        self.corr_canvas.fig.tight_layout()
        self.corr_canvas.draw()

    def _create_regression_tab(self):
        tab = QWidget()
        self.tabs.addTab(tab, "Régression")
        left, left_layout, right, self.reg_canvas = self._setup_analysis_layout(tab)
        col_row = QHBoxLayout()
        self.reg_y_combo = QComboBox()
        self.reg_y_combo.setMinimumWidth(80)
        col_row.addWidget(QLabel("Y (réponse) :"))
        col_row.addWidget(self.reg_y_combo)
        self.reg_x_combo = QComboBox()
        self.reg_x_combo.setMinimumWidth(80)
        col_row.addWidget(QLabel("X (prédicteur) :"))
        col_row.addWidget(self.reg_x_combo)
        refresh_btn = QPushButton("↻")
        refresh_btn.setMaximumWidth(30)
        refresh_btn.clicked.connect(self._refresh_all_combos)
        col_row.addWidget(refresh_btn)
        left_layout.addLayout(col_row)
        self._refresh_all_combos()
        params_group = QGroupBox("Options")
        params_layout = QFormLayout(params_group)
        self.reg_show_ci = QCheckBox("Intervalle de confiance")
        self.reg_show_ci.setChecked(True)
        params_layout.addRow(self.reg_show_ci)
        self.reg_confidence = QDoubleSpinBox()
        self.reg_confidence.setRange(80, 99.9)
        self.reg_confidence.setValue(95)
        self.reg_confidence.setSuffix(" %")
        self._tip(self.reg_confidence, "Niveau de confiance pour les intervalles")
        params_layout.addRow("Confiance :", self.reg_confidence)
        left_layout.addWidget(params_group)
        calc_btn = QPushButton(" Calculer")
        calc_btn.setMinimumHeight(40)
        calc_btn.clicked.connect(self._run_regression)
        left_layout.addWidget(calc_btn)
        result_group = QGroupBox("Résultats")
        result_layout = QVBoxLayout(result_group)
        self.reg_result_text = QTextEdit()
        self.reg_result_text.setReadOnly(True)
        self.reg_result_text.setFont(QFont("Courier", 10))
        result_layout.addWidget(self.reg_result_text)
        self._add_copy_button(result_layout, self.reg_result_text)
        left_layout.addWidget(result_group)

    def _run_regression(self):
        y_data, y_name = self._get_combo_data(self.reg_y_combo)
        x_data, x_name = self._get_combo_data(self.reg_x_combo)
        if y_data is None or x_data is None:
            QMessageBox.warning(self, "Attention", "Sélectionnez Y et X")
            return
        try:
            min_len = min(len(y_data), len(x_data))
            y_data = y_data[:min_len]
            x_data = x_data[:min_len]

            if min_len < 3:
                QMessageBox.warning(self, "Attention", "Au moins 3 points sont nécessaires pour la régression.")
                return
            if len(np.unique(x_data)) < 2:
                QMessageBox.warning(self, "Attention", "La colonne X doit contenir au moins deux valeurs différentes.")
                return

            slope, intercept, r_value, p_value, std_err = scipy_stats.linregress(x_data, y_data)
            confidence = self.reg_confidence.value() / 100.0
            lines = ["=" * 60, "RÉGRESSION LINÉAIRE", "=" * 60,
                     f"\nY = {y_name}, X = {x_name}", f"N = {min_len}",
                     f"\nÉquation : Y = {intercept:.6f} + {slope:.6f} * X",
                     f"\nR² = {r_value**2:.6f}", f"R = {r_value:.6f}",
                     f"p-value = {p_value:.6f}", f"Erreur standard = {std_err:.6f}",
                     f"\nIntervalle confiance : {confidence*100:.0f}%"]
            self.reg_result_text.setText("\n".join(lines))
            self._plot_regression(x_data, y_data, slope, intercept)
        except Exception as e:
            QMessageBox.critical(self, "Erreur", str(e))

    def _plot_regression(self, x, y, slope, intercept):
        self.reg_canvas.fig.clear()
        ax = self.reg_canvas.fig.add_subplot(211)
        ax.scatter(x, y, alpha=0.6, color="blue", label="Données")
        x_line = np.linspace(np.min(x), np.max(x), 100)
        ax.plot(x_line, intercept + slope * x_line, "r-", linewidth=2, label=f"Y = {intercept:.2f} + {slope:.2f}X")
        if self.reg_show_ci.isChecked():
            n = len(x)
            mean_x = np.mean(x)
            ss_x = np.sum((x - mean_x)**2)
            y_pred = intercept + slope * x_line
            mse = np.sum((y - (intercept + slope * x))**2) / (n - 2)
            t_val = scipy_stats.t.ppf(0.975, n - 2)
            se_pred = np.sqrt(mse * (1/n + (x_line - mean_x)**2 / ss_x))
            ci_upper = y_pred + t_val * se_pred
            ci_lower = y_pred - t_val * se_pred
            ax.fill_between(x_line, ci_lower, ci_upper, alpha=0.2, color="red", label="IC 95%")
        ax.legend()
        ax.set_title("Régression linéaire")
        ax2 = self.reg_canvas.fig.add_subplot(212)
        residuals = y - (intercept + slope * x)
        ax2.scatter(x, residuals, alpha=0.6, color="green")
        ax2.axhline(0, color="red", linestyle="--")
        ax2.set_title("Résidus")
        ax2.set_xlabel("X")
        ax2.set_ylabel("Résidu")
        self.reg_canvas.fig.tight_layout()
        self.reg_canvas.draw()

    # ===== T-TEST =====
    def _create_ttest_tab(self):
        tab = QWidget()
        self.tabs.addTab(tab, "Test t")
        left, left_layout, right, self.tt_canvas = self._setup_analysis_layout(tab)
        col_row = QHBoxLayout()
        self.tt_col1_combo = QComboBox()
        self.tt_col1_combo.setMinimumWidth(80)
        col_row.addWidget(QLabel("Colonne 1 :"))
        col_row.addWidget(self.tt_col1_combo)
        self.tt_col2_combo = QComboBox()
        self.tt_col2_combo.setMinimumWidth(80)
        col_row.addWidget(QLabel("Colonne 2 :"))
        col_row.addWidget(self.tt_col2_combo)
        refresh_btn = QPushButton("↻")
        refresh_btn.setMaximumWidth(30)
        refresh_btn.clicked.connect(self._refresh_all_combos)
        col_row.addWidget(refresh_btn)
        left_layout.addLayout(col_row)
        params_group = QGroupBox("Paramètres")
        params_layout = QFormLayout(params_group)
        self.tt_type = QComboBox()
        self.tt_type.addItems(["Deux échantillons", "Un échantillon"])
        self._tip(self.tt_type, "Un échantillon: compare la moyenne à une valeur. Deux échantillons: compare deux moyennes.")
        params_layout.addRow("Type :", self.tt_type)
        self.tt_mu = QDoubleSpinBox()
        self.tt_mu.setRange(-1e9, 1e9)
        self.tt_mu.setValue(0)
        self._tip(self.tt_mu, "Valeur hypothétique de la moyenne (H0: μ = valeur)")
        params_layout.addRow("μ hypothétique :", self.tt_mu)
        self.tt_alpha = QDoubleSpinBox()
        self.tt_alpha.setRange(0.001, 0.5)
        self.tt_alpha.setValue(0.05)
        params_layout.addRow("Seuil α :", self.tt_alpha)
        self.tt_equal_var = QCheckBox("Variances égales")
        self.tt_equal_var.setChecked(True)
        params_layout.addRow(self.tt_equal_var)
        left_layout.addWidget(params_group)
        calc_btn = QPushButton(" Calculer")
        calc_btn.setMinimumHeight(40)
        calc_btn.clicked.connect(self._run_ttest)
        left_layout.addWidget(calc_btn)
        result_group = QGroupBox("Résultats")
        result_layout = QVBoxLayout(result_group)
        self.tt_result_text = QTextEdit()
        self.tt_result_text.setReadOnly(True)
        self.tt_result_text.setFont(QFont("Courier", 10))
        result_layout.addWidget(self.tt_result_text)
        self._add_copy_button(result_layout, self.tt_result_text)
        left_layout.addWidget(result_group)

    def _run_ttest(self):
        data1, name1 = self._get_combo_data(self.tt_col1_combo)
        if data1 is None:
            QMessageBox.warning(self, "Attention", "Sélectionnez une colonne")
            return
        try:
            alpha = self.tt_alpha.value()
            if self.tt_type.currentText() == "Un échantillon":
                mu = self.tt_mu.value()
                t_stat, p_value = scipy_stats.ttest_1samp(data1, mu)
                lines = ["=" * 60, "TEST T - UN ÉCHANTILLON", "=" * 60,
                         f"\nColonne : {name1}", f"N = {len(data1)}",
                         f"Moyenne = {np.mean(data1):.6f}", f"μ hypothétique = {mu}",
                         f"\nt = {t_stat:.6f}", f"p-value = {p_value:.6f}",
                         f"Seuil α = {alpha}",
                         f"\nH0 rejetée = {'OUI' if p_value < alpha else 'NON'}"]
            else:
                data2, name2 = self._get_combo_data(self.tt_col2_combo)
                if data2 is None:
                    QMessageBox.warning(self, "Attention", "Sélectionnez une 2ème colonne")
                    return
                equal_var = self.tt_equal_var.isChecked()
                t_stat, p_value = scipy_stats.ttest_ind(data1, data2, equal_var=equal_var)
                lines = ["=" * 60, "TEST T - DEUX ÉCHANTILLONS", "=" * 60,
                         f"\nColonne 1 : {name1} (N={len(data1)}, moy={np.mean(data1):.4f})",
                         f"Colonne 2 : {name2} (N={len(data2)}, moy={np.mean(data2):.4f})",
                         f"\nt = {t_stat:.6f}", f"p-value = {p_value:.6f}",
                         f"Variances égales = {'Oui' if equal_var else 'Non'}",
                         f"Seuil α = {alpha}",
                         f"\nH0 rejetée = {'OUI' if p_value < alpha else 'NON'}"]
            self.tt_result_text.setText("\n".join(lines))
            self._plot_ttest(data1, data2 if self.tt_type.currentText() == "Deux échantillons" else None)
        except Exception as e:
            QMessageBox.critical(self, "Erreur", str(e))

    def _plot_ttest(self, data1, data2):
        self.tt_canvas.fig.clear()
        ax = self.tt_canvas.fig.add_subplot(111)
        if data2 is not None:
            ax.boxplot([data1, data2], labels=[self.tt_col1_combo.currentText(), self.tt_col2_combo.currentText()],
                      patch_artist=True)
        else:
            ax.boxplot([data1], labels=[self.tt_col1_combo.currentText()], patch_artist=True)
            ax.axhline(self.tt_mu.value(), color="red", linestyle="--", label=f"μ={self.tt_mu.value()}")
            ax.legend()
        ax.set_title("Comparaison")
        self.tt_canvas.fig.tight_layout()
        self.tt_canvas.draw()

    # ===== ANOVA =====
    def _create_anova_tab(self):
        tab = QWidget()
        self.tabs.addTab(tab, "ANOVA")
        left, left_layout, right, self.anova_canvas = self._setup_analysis_layout(tab)

        params_group = QGroupBox("Paramètres")
        params_layout = QFormLayout(params_group)
        self.anova_alpha = QDoubleSpinBox()
        self.anova_alpha.setRange(0.001, 0.5)
        self.anova_alpha.setValue(0.05)
        params_layout.addRow("Seuil α :", self.anova_alpha)
        self.anova_ngroups = QSpinBox()
        self.anova_ngroups.setRange(2, NUM_COLS)
        self.anova_ngroups.setValue(3)
        params_layout.addRow("Nombre de colonnes :", self.anova_ngroups)
        left_layout.addWidget(params_group)

        select_group = QGroupBox("Colonnes / groupes à comparer")
        select_layout = QVBoxLayout(select_group)
        self.anova_combo_container = QWidget()
        self.anova_combo_layout = QVBoxLayout(self.anova_combo_container)
        self.anova_combo_layout.setContentsMargins(0, 0, 0, 0)
        select_layout.addWidget(self.anova_combo_container)
        refresh_btn = QPushButton("↻ Actualiser les colonnes")
        refresh_btn.clicked.connect(self._refresh_all_combos)
        select_layout.addWidget(refresh_btn)
        left_layout.addWidget(select_group)

        self.anova_col_combos = []
        self._update_anova_combos()
        self.anova_ngroups.valueChanged.connect(self._update_anova_combos)

        calc_btn = QPushButton(" Calculer ANOVA")
        calc_btn.setMinimumHeight(40)
        calc_btn.setStyleSheet("font-weight: bold; font-size: 14px;")
        calc_btn.clicked.connect(self._run_anova)
        left_layout.addWidget(calc_btn)

        result_group = QGroupBox("Résultats")
        result_layout = QVBoxLayout(result_group)
        self.anova_result_text = QTextEdit()
        self.anova_result_text.setReadOnly(True)
        self.anova_result_text.setFont(QFont("Courier", 10))
        result_layout.addWidget(self.anova_result_text)
        self._add_copy_button(result_layout, self.anova_result_text)
        left_layout.addWidget(result_group)

    def _update_anova_combos(self):
        current_selections = [c.currentText() for c in getattr(self, "anova_col_combos", [])]
        if not hasattr(self, "anova_combo_layout"):
            return
        while self.anova_combo_layout.count():
            item = self.anova_combo_layout.takeAt(0)
            if item.widget():
                item.widget().deleteLater()
            elif item.layout():
                while item.layout().count():
                    child = item.layout().takeAt(0)
                    if child.widget():
                        child.widget().deleteLater()
        self.anova_col_combos = []
        cols = []
        if hasattr(self, "sheet"):
            for i in range(self.sheet.columnCount()):
                data = self.sheet.get_column_data(i)
                if data is not None and len(data) > 0:
                    cols.append(col_letter(i))
        if not cols:
            cols = ["(aucune donnée)"]
        n = self.anova_ngroups.value() if hasattr(self, "anova_ngroups") else 2
        for i in range(n):
            row = QHBoxLayout()
            label = QLabel(f"Groupe {i + 1} :")
            label.setMinimumWidth(80)
            row.addWidget(label)
            combo = QComboBox()
            combo.setMinimumWidth(100)
            combo.addItems(cols)
            if i < len(current_selections) and current_selections[i] in cols:
                combo.setCurrentText(current_selections[i])
            elif i < len(cols) and cols[0] != "(aucune donnée)":
                combo.setCurrentText(cols[i % len(cols)])
            row.addWidget(combo)
            row.addStretch()
            self.anova_combo_layout.addLayout(row)
            self.anova_col_combos.append(combo)

    def _run_anova(self):
        group_data, group_labels, selected_letters = [], [], []
        for i, combo in enumerate(getattr(self, "anova_col_combos", [])):
            data, letter = self._get_combo_data(combo)
            if data is not None and letter is not None:
                group_data.append(data)
                group_labels.append(f"G{i + 1} ({letter})")
                selected_letters.append(letter)
        if len(group_data) < 2:
            QMessageBox.warning(self, "Attention", "Au moins 2 groupes avec données sont requis pour ANOVA")
            return
        if len(set(selected_letters)) != len(selected_letters):
            QMessageBox.warning(self, "Attention", "Sélectionnez des colonnes différentes pour l'ANOVA")
            return
        if any(len(g) < 2 for g in group_data):
            QMessageBox.warning(self, "Attention", "Chaque groupe doit contenir au moins 2 valeurs numériques")
            return
        try:
            alpha = self.anova_alpha.value()
            f_stat, p_value = scipy_stats.f_oneway(*group_data)
            grand_mean = np.mean(np.concatenate(group_data))
            ss_between = sum(len(g) * (np.mean(g) - grand_mean) ** 2 for g in group_data)
            ss_within = sum(np.sum((g - np.mean(g)) ** 2) for g in group_data)
            ss_total = ss_between + ss_within
            k = len(group_data)
            n_total = sum(len(g) for g in group_data)
            df_between = k - 1
            df_within = n_total - k
            ms_between = ss_between / df_between
            ms_within = ss_within / df_within if df_within > 0 else np.nan
            eta_sq = ss_between / ss_total if ss_total > 0 else 0
            lines = ["=" * 60, "ANOVA - Analyse de variance", "=" * 60]
            lines += [f"\nNombre de groupes : {k}", f"Colonnes : {', '.join(selected_letters)}", f"N total = {n_total}", f"Seuil α = {alpha}"]
            lines.append("\nSource      |   SS      |  df  |   MS      |    F")
            lines.append("-" * 60)
            lines.append(f"Entre      | {ss_between:10.4f} | {df_between:4d} | {ms_between:10.4f} | {f_stat:.4f}")
            lines.append(f"Intra      | {ss_within:10.4f} | {df_within:4d} | {ms_within:10.4f} |")
            lines.append(f"Total      | {ss_total:10.4f} | {n_total - 1:4d} |           |")
            lines.append(f"\np-value = {p_value:.6f}")
            lines.append(f"η² (eta-squared) = {eta_sq:.4f}")
            lines.append(f"\nH0 rejetée = {'OUI' if p_value < alpha else 'NON'}")
            if p_value < alpha:
                lines.append("\n→ Au moins un groupe diffère significativement")
                lines.append("\nTests post-hoc paire à paire (Welch + correction Bonferroni) :")
                lines.extend(self._anova_posthoc(group_data, group_labels, alpha))
            lines.append("\nStatistiques par groupe :")
            for name, data in zip(group_labels, group_data):
                lines.append(f"  {name}: N={len(data)}, moy={np.mean(data):.4f}, std={np.std(data, ddof=1):.4f}")
            self.anova_result_text.setText("\n".join(lines))
            self._plot_anova(group_data, group_labels)
            self.status_bar.showMessage(f"ANOVA terminée : {k} groupes, N total={n_total}")
            self._refresh_report_panel()
        except Exception as e:
            QMessageBox.critical(self, "Erreur", str(e))

    def _plot_anova(self, group_data, group_labels):
        self.anova_canvas.fig.clear()
        ax = self.anova_canvas.fig.add_subplot(111)
        bp = ax.boxplot(group_data, tick_labels=group_labels, patch_artist=True)
        colors = plt.cm.Set3(np.linspace(0, 1, len(group_data)))
        for patch, color in zip(bp['boxes'], colors):
            patch.set_facecolor(color)
        ax.set_title("ANOVA - Comparaison des groupes")
        ax.set_ylabel("Valeur")
        ax.tick_params(axis='x', rotation=15)
        try:
            self.anova_canvas.fig.tight_layout()
        except Exception:
            pass
        self.anova_canvas.draw()

    # ===== BOXPLOTS =====
    def _create_boxplot_tab(self):
        tab = QWidget()
        self.tabs.addTab(tab, "Boxplots")
        left, left_layout, right, self.boxplot_canvas = self._setup_analysis_layout(tab)
        params_group = QGroupBox("Options")
        params_layout = QFormLayout(params_group)
        self.boxplot_ngroups = QSpinBox()
        self.boxplot_ngroups.setRange(1, NUM_COLS)
        self.boxplot_ngroups.setValue(3)
        params_layout.addRow("Nombre de colonnes :", self.boxplot_ngroups)
        self.boxplot_notched = QCheckBox("Boîtes à encoches")
        self.boxplot_notched.setChecked(False)
        params_layout.addRow(self.boxplot_notched)
        self.boxplot_show_outliers = QCheckBox("Afficher valeurs aberrantes")
        self.boxplot_show_outliers.setChecked(True)
        params_layout.addRow(self.boxplot_show_outliers)
        left_layout.addWidget(params_group)
        select_group = QGroupBox("Colonnes à afficher / contrôler")
        select_layout = QVBoxLayout(select_group)
        self.boxplot_combo_container = QWidget()
        self.boxplot_combo_layout = QVBoxLayout(self.boxplot_combo_container)
        self.boxplot_combo_layout.setContentsMargins(0, 0, 0, 0)
        select_layout.addWidget(self.boxplot_combo_container)
        refresh_btn = QPushButton("↻ Actualiser les colonnes")
        refresh_btn.clicked.connect(self._refresh_all_combos)
        select_layout.addWidget(refresh_btn)
        left_layout.addWidget(select_group)
        self.boxplot_col_combos = []
        self._update_boxplot_combos()
        self.boxplot_ngroups.valueChanged.connect(self._update_boxplot_combos)
        calc_btn = QPushButton(" Tracer")
        calc_btn.setMinimumHeight(40)
        calc_btn.setStyleSheet("font-weight: bold; font-size: 14px;")
        calc_btn.clicked.connect(self._run_boxplot)
        left_layout.addWidget(calc_btn)
        result_group = QGroupBox("Résultats")
        result_layout = QVBoxLayout(result_group)
        self.boxplot_result_text = QTextEdit()
        self.boxplot_result_text.setReadOnly(True)
        self.boxplot_result_text.setFont(QFont("Courier", 10))
        result_layout.addWidget(self.boxplot_result_text)
        self._add_copy_button(result_layout, self.boxplot_result_text)
        left_layout.addWidget(result_group)

    def _update_boxplot_combos(self):
        current_selections = [c.currentText() for c in getattr(self, "boxplot_col_combos", [])]
        if not hasattr(self, "boxplot_combo_layout"):
            return
        while self.boxplot_combo_layout.count():
            item = self.boxplot_combo_layout.takeAt(0)
            if item.widget():
                item.widget().deleteLater()
            elif item.layout():
                while item.layout().count():
                    child = item.layout().takeAt(0)
                    if child.widget():
                        child.widget().deleteLater()
        self.boxplot_col_combos = []
        cols = []
        if hasattr(self, "sheet"):
            for i in range(self.sheet.columnCount()):
                data = self.sheet.get_column_data(i)
                if data is not None and len(data) > 0:
                    cols.append(col_letter(i))
        if not cols:
            cols = ["(aucune donnée)"]
        n = self.boxplot_ngroups.value() if hasattr(self, "boxplot_ngroups") else 1
        for i in range(n):
            row = QHBoxLayout()
            label = QLabel(f"Colonne {i + 1} :")
            label.setMinimumWidth(80)
            row.addWidget(label)
            combo = QComboBox()
            combo.setMinimumWidth(100)
            combo.addItems(cols)
            if i < len(current_selections) and current_selections[i] in cols:
                combo.setCurrentText(current_selections[i])
            elif i < len(cols) and cols[0] != "(aucune donnée)":
                combo.setCurrentText(cols[i % len(cols)])
            row.addWidget(combo)
            row.addStretch()
            self.boxplot_combo_layout.addLayout(row)
            self.boxplot_col_combos.append(combo)

    def _run_boxplot(self):
        groups, selected_letters = {}, []
        for i, combo in enumerate(getattr(self, "boxplot_col_combos", [])):
            data, letter = self._get_combo_data(combo)
            if data is not None and letter is not None:
                groups[f"S{i + 1} ({letter})"] = data
                selected_letters.append(letter)
        if not groups:
            QMessageBox.warning(self, "Attention", "Aucune série avec données")
            return
        if len(set(selected_letters)) != len(selected_letters):
            QMessageBox.warning(self, "Attention", "Sélectionnez des colonnes différentes pour le boxplot")
            return
        try:
            lines = ["=" * 60, "BOXPLOTS", "=" * 60, f"\nColonnes : {', '.join(selected_letters)}"]
            for name, data in groups.items():
                q1, q2, q3 = np.percentile(data, [25, 50, 75])
                iqr = q3 - q1
                lines.append(f"\n{name}: N={len(data)}, Q1={q1:.4f}, Méd={q2:.4f}, Q3={q3:.4f}, IQR={iqr:.4f}")
            self.boxplot_result_text.setText("\n".join(lines))
            self._plot_boxplots(groups)
            self.status_bar.showMessage(f"Boxplot tracé : {len(groups)} colonnes")
            self._refresh_report_panel()
        except Exception as e:
            QMessageBox.critical(self, "Erreur", str(e))

    def _plot_boxplots(self, groups):
        self.boxplot_canvas.fig.clear()
        ax = self.boxplot_canvas.fig.add_subplot(111)
        data = list(groups.values())
        labels = list(groups.keys())
        bp = ax.boxplot(data, tick_labels=labels, patch_artist=True,
                       notch=self.boxplot_notched.isChecked(),
                       showfliers=self.boxplot_show_outliers.isChecked())
        colors = plt.cm.Pastel1(np.linspace(0, 1, len(data)))
        for patch, color in zip(bp['boxes'], colors):
            patch.set_facecolor(color)
        ax.set_title("Boxplots comparatifs")
        ax.set_ylabel("Valeur")
        ax.tick_params(axis='x', rotation=15)
        try:
            self.boxplot_canvas.fig.tight_layout()
        except Exception:
            pass
        self.boxplot_canvas.draw()

    # ===== CONTROL CHARTS =====
    def _create_control_charts_tab(self):
        tab = QWidget()
        self.tabs.addTab(tab, "Cartes de contrôle")
        left, left_layout, right, self.cc_canvas = self._setup_analysis_layout(tab)
        col_row = QHBoxLayout()
        self.cc_col_combo = QComboBox()
        self.cc_col_combo.setMinimumWidth(80)
        col_row.addWidget(QLabel("Colonne :"))
        col_row.addWidget(self.cc_col_combo)
        refresh_btn = QPushButton("↻")
        refresh_btn.setMaximumWidth(30)
        refresh_btn.clicked.connect(self._refresh_all_combos)
        col_row.addWidget(refresh_btn)
        left_layout.addLayout(col_row)
        params_group = QGroupBox("Paramètres")
        params_layout = QFormLayout(params_group)
        self.cc_type = QComboBox()
        self.cc_type.addItems(["X̄-R", "X̄-S", "I-MR", "EWMA", "CUSUM", "P", "U", "C"])
        params_layout.addRow("Type :", self.cc_type)
        self.cc_subgroup = QSpinBox()
        self.cc_subgroup.setRange(2, 25)
        self.cc_subgroup.setValue(5)
        params_layout.addRow("Sous-groupe :", self.cc_subgroup)
        self.cc_sample_size = QSpinBox()
        self.cc_sample_size.setRange(1, 1000000)
        self.cc_sample_size.setValue(100)
        self._tip(self.cc_sample_size, "Carte P : taille d'échantillon constante. Si les données sont déjà des proportions, laisser 1.")
        params_layout.addRow("n échantillon (P) :", self.cc_sample_size)
        self.cc_opportunities = QDoubleSpinBox()
        self.cc_opportunities.setRange(0.001, 1e9)
        self.cc_opportunities.setValue(1.0)
        self._tip(self.cc_opportunities, "Carte U : nombre d'unités/opportunités par observation.")
        params_layout.addRow("Unités/opportunités (U) :", self.cc_opportunities)
        self.cc_lambda = QDoubleSpinBox()
        self.cc_lambda.setRange(0.01, 1.0)
        self.cc_lambda.setSingleStep(0.05)
        self.cc_lambda.setValue(0.2)
        params_layout.addRow("Lambda (EWMA) :", self.cc_lambda)
        self.cc_k = QDoubleSpinBox()
        self.cc_k.setRange(0.1, 2.0)
        self.cc_k.setSingleStep(0.1)
        self.cc_k.setValue(0.5)
        params_layout.addRow("k (CUSUM) :", self.cc_k)
        left_layout.addWidget(params_group)
        calc_btn = QPushButton(" Tracer")
        calc_btn.setMinimumHeight(40)
        calc_btn.clicked.connect(self._run_control_charts)
        left_layout.addWidget(calc_btn)
        result_group = QGroupBox("Résultats")
        result_layout = QVBoxLayout(result_group)
        self.cc_result_text = QTextEdit()
        self.cc_result_text.setReadOnly(True)
        self.cc_result_text.setFont(QFont("Courier", 10))
        result_layout.addWidget(self.cc_result_text)
        self._add_copy_button(result_layout, self.cc_result_text)
        left_layout.addWidget(result_group)


    def _cc_make_subgroups(self, data, subgroup_size):
        """Découpe les données en sous-groupes complets pour cartes Xbar-R / Xbar-S."""
        x = np.asarray(data, dtype=float)
        x = x[~np.isnan(x)]
        if subgroup_size < 2:
            raise ValueError("La taille de sous-groupe doit être au moins 2 pour les cartes X̄-R et X̄-S.")
        n_groups = len(x) // subgroup_size
        if n_groups < 2:
            raise ValueError(
                f"Données insuffisantes : il faut au moins 2 sous-groupes complets de taille {subgroup_size}."
            )
        used = n_groups * subgroup_size
        ignored = len(x) - used
        return x[:used].reshape(n_groups, subgroup_size), ignored

    def _cc_constants(self, subgroup_size):
        """Constantes usuelles pour cartes Xbar-R et Xbar-S, n=2..25."""
        table = {
            2:  {"A2": 1.880, "D3": 0.000, "D4": 3.267, "A3": 2.659, "B3": 0.000, "B4": 3.267},
            3:  {"A2": 1.023, "D3": 0.000, "D4": 2.574, "A3": 1.954, "B3": 0.000, "B4": 2.568},
            4:  {"A2": 0.729, "D3": 0.000, "D4": 2.282, "A3": 1.628, "B3": 0.000, "B4": 2.266},
            5:  {"A2": 0.577, "D3": 0.000, "D4": 2.114, "A3": 1.427, "B3": 0.000, "B4": 2.089},
            6:  {"A2": 0.483, "D3": 0.000, "D4": 2.004, "A3": 1.287, "B3": 0.030, "B4": 1.970},
            7:  {"A2": 0.419, "D3": 0.076, "D4": 1.924, "A3": 1.182, "B3": 0.118, "B4": 1.882},
            8:  {"A2": 0.373, "D3": 0.136, "D4": 1.864, "A3": 1.099, "B3": 0.185, "B4": 1.815},
            9:  {"A2": 0.337, "D3": 0.184, "D4": 1.816, "A3": 1.032, "B3": 0.239, "B4": 1.761},
            10: {"A2": 0.308, "D3": 0.223, "D4": 1.777, "A3": 0.975, "B3": 0.284, "B4": 1.716},
            11: {"A2": 0.285, "D3": 0.256, "D4": 1.744, "A3": 0.927, "B3": 0.321, "B4": 1.679},
            12: {"A2": 0.266, "D3": 0.283, "D4": 1.717, "A3": 0.886, "B3": 0.354, "B4": 1.646},
            13: {"A2": 0.249, "D3": 0.307, "D4": 1.693, "A3": 0.850, "B3": 0.382, "B4": 1.618},
            14: {"A2": 0.235, "D3": 0.328, "D4": 1.672, "A3": 0.817, "B3": 0.406, "B4": 1.594},
            15: {"A2": 0.223, "D3": 0.347, "D4": 1.653, "A3": 0.789, "B3": 0.428, "B4": 1.572},
            16: {"A2": 0.212, "D3": 0.363, "D4": 1.637, "A3": 0.763, "B3": 0.448, "B4": 1.552},
            17: {"A2": 0.203, "D3": 0.378, "D4": 1.622, "A3": 0.739, "B3": 0.466, "B4": 1.534},
            18: {"A2": 0.194, "D3": 0.391, "D4": 1.608, "A3": 0.718, "B3": 0.482, "B4": 1.518},
            19: {"A2": 0.187, "D3": 0.403, "D4": 1.597, "A3": 0.698, "B3": 0.497, "B4": 1.503},
            20: {"A2": 0.180, "D3": 0.415, "D4": 1.585, "A3": 0.680, "B3": 0.510, "B4": 1.490},
            21: {"A2": 0.173, "D3": 0.425, "D4": 1.575, "A3": 0.663, "B3": 0.523, "B4": 1.477},
            22: {"A2": 0.167, "D3": 0.434, "D4": 1.566, "A3": 0.647, "B3": 0.534, "B4": 1.466},
            23: {"A2": 0.162, "D3": 0.443, "D4": 1.557, "A3": 0.633, "B3": 0.545, "B4": 1.455},
            24: {"A2": 0.157, "D3": 0.451, "D4": 1.548, "A3": 0.619, "B3": 0.555, "B4": 1.445},
            25: {"A2": 0.153, "D3": 0.459, "D4": 1.541, "A3": 0.606, "B3": 0.565, "B4": 1.435},
        }
        return table.get(int(subgroup_size), table[25])

    def _cc_xbar_stats(self, data, subgroup_size, chart_kind):
        """Calcule les statistiques et limites pour X̄-R ou X̄-S."""
        groups, ignored = self._cc_make_subgroups(data, subgroup_size)
        xbars = groups.mean(axis=1)
        ranges = groups.max(axis=1) - groups.min(axis=1)
        stds = groups.std(axis=1, ddof=1)
        xbarbar = float(np.mean(xbars))
        const = self._cc_constants(subgroup_size)

        if chart_kind == "XBAR_R":
            rbar = float(np.mean(ranges))
            main = {
                "values": xbars,
                "center": xbarbar,
                "ucl": xbarbar + const["A2"] * rbar,
                "lcl": xbarbar - const["A2"] * rbar,
                "ylabel": "Moyenne sous-groupe",
                "title": "Carte X̄",
            }
            dispersion = {
                "values": ranges,
                "center": rbar,
                "ucl": const["D4"] * rbar,
                "lcl": const["D3"] * rbar,
                "ylabel": "Étendue R",
                "title": "Carte R",
            }
            method = f"X̄-R : limites X̄ = X̄̄ ± A2×R̄ ; limites R = D3/D4×R̄"
        else:
            sbar = float(np.mean(stds))
            main = {
                "values": xbars,
                "center": xbarbar,
                "ucl": xbarbar + const["A3"] * sbar,
                "lcl": xbarbar - const["A3"] * sbar,
                "ylabel": "Moyenne sous-groupe",
                "title": "Carte X̄",
            }
            dispersion = {
                "values": stds,
                "center": sbar,
                "ucl": const["B4"] * sbar,
                "lcl": const["B3"] * sbar,
                "ylabel": "Écart-type S",
                "title": "Carte S",
            }
            method = f"X̄-S : limites X̄ = X̄̄ ± A3×S̄ ; limites S = B3/B4×S̄"

        return {
            "groups": groups,
            "ignored": ignored,
            "subgroup_size": subgroup_size,
            "n_groups": len(groups),
            "constants": const,
            "main": main,
            "dispersion": dispersion,
            "method": method,
        }

    def _cc_count_ooc(self, values, lcl, ucl):
        values = np.asarray(values, dtype=float)
        return int(np.sum((values > ucl) | (values < lcl)))

    def _run_control_charts(self):
        data, col_name = self._get_combo_data(self.cc_col_combo)
        if data is None:
            QMessageBox.warning(self, "Attention", "Sélectionnez une colonne")
            return
        try:
            data = np.asarray(data, dtype=float)
            data = data[~np.isnan(data)]
            if len(data) < 2:
                QMessageBox.warning(self, "Attention", "Au moins 2 valeurs numériques sont nécessaires.")
                return

            cc_type = self.cc_type.currentText()
            subgroup = self.cc_subgroup.value()
            lines = ["=" * 60, f"CARTE DE CONTRÔLE : {cc_type}", "=" * 60,
                     f"\nColonne : {col_name}", f"N = {len(data)}"]

            # Normalisation des libellés : l'interface contient X̄-R et X̄-S.
            if cc_type in ("X̄-R", "Xbar-R", "X-R"):
                stats_xr = self._cc_xbar_stats(data, subgroup, "XBAR_R")
                main = stats_xr["main"]
                disp = stats_xr["dispersion"]
                lines.append(f"\n{stats_xr['method']}")
                lines.append(f"Sous-groupe = {subgroup}")
                lines.append(f"Sous-groupes complets utilisés = {stats_xr['n_groups']}")
                if stats_xr["ignored"]:
                    lines.append(f"Observations ignorées faute de sous-groupe complet = {stats_xr['ignored']}")
                c = stats_xr["constants"]
                lines.append(f"Constantes : A2={c['A2']:.3f}, D3={c['D3']:.3f}, D4={c['D4']:.3f}")
                lines.append(f"\nCarte X̄ : CL={main['center']:.6f}, UCL={main['ucl']:.6f}, LCL={main['lcl']:.6f}")
                lines.append(f"Carte R : CL={disp['center']:.6f}, UCL={disp['ucl']:.6f}, LCL={disp['lcl']:.6f}")
                lines.append(f"\nPoints hors contrôle X̄ : {self._cc_count_ooc(main['values'], main['lcl'], main['ucl'])}/{len(main['values'])}")
                lines.append(f"Points hors contrôle R : {self._cc_count_ooc(disp['values'], disp['lcl'], disp['ucl'])}/{len(disp['values'])}")
                rules = self._detect_control_rules(main["values"], main["center"], (main["ucl"] - main["center"]) / 3 if main["ucl"] != main["center"] else 0)
                lines.append("\nRègles sur la carte X̄ :")
                lines.extend(rules if rules else ["  Aucun signal détecté"])
                self.cc_result_text.setText("\n".join(lines))
                self._plot_control_charts(data, cc_type, subgroup)
                self._refresh_report_panel()
                return

            if cc_type in ("X̄-S", "Xbar-S", "X-S"):
                stats_xs = self._cc_xbar_stats(data, subgroup, "XBAR_S")
                main = stats_xs["main"]
                disp = stats_xs["dispersion"]
                lines.append(f"\n{stats_xs['method']}")
                lines.append(f"Sous-groupe = {subgroup}")
                lines.append(f"Sous-groupes complets utilisés = {stats_xs['n_groups']}")
                if stats_xs["ignored"]:
                    lines.append(f"Observations ignorées faute de sous-groupe complet = {stats_xs['ignored']}")
                c = stats_xs["constants"]
                lines.append(f"Constantes : A3={c['A3']:.3f}, B3={c['B3']:.3f}, B4={c['B4']:.3f}")
                lines.append(f"\nCarte X̄ : CL={main['center']:.6f}, UCL={main['ucl']:.6f}, LCL={main['lcl']:.6f}")
                lines.append(f"Carte S : CL={disp['center']:.6f}, UCL={disp['ucl']:.6f}, LCL={disp['lcl']:.6f}")
                lines.append(f"\nPoints hors contrôle X̄ : {self._cc_count_ooc(main['values'], main['lcl'], main['ucl'])}/{len(main['values'])}")
                lines.append(f"Points hors contrôle S : {self._cc_count_ooc(disp['values'], disp['lcl'], disp['ucl'])}/{len(disp['values'])}")
                rules = self._detect_control_rules(main["values"], main["center"], (main["ucl"] - main["center"]) / 3 if main["ucl"] != main["center"] else 0)
                lines.append("\nRègles sur la carte X̄ :")
                lines.extend(rules if rules else ["  Aucun signal détecté"])
                self.cc_result_text.setText("\n".join(lines))
                self._plot_control_charts(data, cc_type, subgroup)
                self._refresh_report_panel()
                return

            mean_val = np.mean(data)
            std_val = np.std(data, ddof=1) if len(data) > 1 else 0
            plotted = data
            center = mean_val
            sigma = std_val
            ucl = mean_val + 3 * std_val
            lcl = mean_val - 3 * std_val

            if cc_type == "P":
                n = self.cc_sample_size.value()
                plotted = data if np.nanmax(data) <= 1 else data / n
                center = np.mean(plotted)
                sigma = np.sqrt(max(center * (1 - center) / n, 0))
                ucl = min(1, center + 3 * sigma)
                lcl = max(0, center - 3 * sigma)
                lines.append(f"\nCarte P : p̄ = {center:.6f}, n = {n}")
            elif cc_type == "U":
                opp = self.cc_opportunities.value()
                plotted = data / opp
                center = np.mean(plotted)
                sigma = np.sqrt(center / opp) if center >= 0 else np.nan
                ucl = center + 3 * sigma
                lcl = max(0, center - 3 * sigma)
                lines.append(f"\nCarte U : ū = {center:.6f}, unités/opportunités = {opp:.3f}")
            elif cc_type == "C":
                plotted = data
                center = np.mean(plotted)
                sigma = np.sqrt(center) if center >= 0 else np.nan
                ucl = center + 3 * sigma
                lcl = max(0, center - 3 * sigma)
                lines.append(f"\nCarte C : c̄ = {center:.6f}")
            elif cc_type == "EWMA":
                lam = self.cc_lambda.value()
                z = np.zeros(len(data))
                z[0] = mean_val
                for i in range(1, len(data)):
                    z[i] = lam * data[i] + (1 - lam) * z[i-1]
                sigma_z = std_val * np.sqrt(lam / (2 - lam) * (1 - (1 - lam)**(2 * np.arange(1, len(data) + 1))))
                ucl_ewma = mean_val + 3 * sigma_z
                lcl_ewma = mean_val - 3 * sigma_z
                lines.append(f"\nMoyenne = {mean_val:.6f}")
                lines.append(f"Lambda = {lam:.3f}")
                lines.append(f"UCL asymptotique = {mean_val + 3 * std_val * np.sqrt(lam / (2 - lam)):.6f}")
                lines.append(f"LCL asymptotique = {mean_val - 3 * std_val * np.sqrt(lam / (2 - lam)):.6f}")
                out_of_control = int(np.sum((z > ucl_ewma) | (z < lcl_ewma)))
                lines.append(f"\nPoints hors contrôle : {out_of_control}/{len(data)}")
                self.cc_result_text.setText("\n".join(lines))
                self._plot_control_charts(data, cc_type, subgroup)
                self._refresh_report_panel()
                return
            elif cc_type == "CUSUM":
                k = self.cc_k.value() * std_val
                h = 5 * std_val
                s_pos = np.zeros(len(data))
                s_neg = np.zeros(len(data))
                for i in range(1, len(data)):
                    s_pos[i] = max(0, s_pos[i-1] + data[i] - mean_val - k)
                    s_neg[i] = max(0, s_neg[i-1] + mean_val - data[i] - k)
                out_of_control = len(np.where(s_pos > h)[0]) + len(np.where(s_neg > h)[0])
                lines.append(f"\nMoyenne = {mean_val:.6f}")
                lines.append(f"k = {k:.6f}, h = {h:.6f}")
                lines.append(f"\nSignaux CUSUM : {out_of_control}/{len(data)}")
                self.cc_result_text.setText("\n".join(lines))
                self._plot_control_charts(data, cc_type, subgroup)
                self._refresh_report_panel()
                return
            else:
                lines.append(f"\nMoyenne = {mean_val:.6f}")

            lines.append(f"UCL = {ucl:.6f}")
            lines.append(f"LCL = {lcl:.6f}")
            out_of_control = int(np.sum((plotted > ucl) | (plotted < lcl)))
            lines.append(f"\nPoints hors contrôle : {out_of_control}/{len(plotted)}")
            rules = self._detect_control_rules(plotted, center, sigma)
            lines.append("\nRègles Western Electric / Nelson :")
            lines.extend(rules if rules else ["  Aucun signal détecté"])
            self.cc_result_text.setText("\n".join(lines))
            self._plot_control_charts(data, cc_type, subgroup)
            self._refresh_report_panel()
        except Exception as e:
            QMessageBox.critical(self, "Erreur", str(e))

    def _plot_xbar_control_pair(self, stats_dict, chart_label):
        """Trace la paire de cartes X̄-R ou X̄-S."""
        self.cc_canvas.fig.clear()
        ax1 = self.cc_canvas.fig.add_subplot(211)
        ax2 = self.cc_canvas.fig.add_subplot(212)
        x = np.arange(1, stats_dict["n_groups"] + 1)

        for ax, item in [(ax1, stats_dict["main"]), (ax2, stats_dict["dispersion"] )]:
            values = np.asarray(item["values"], dtype=float)
            ax.plot(x, values, "bo-", markersize=4, linewidth=1)
            ax.axhline(item["center"], color="green", linestyle="--", label=f"CL={item['center']:.4g}")
            ax.axhline(item["ucl"], color="red", linestyle="-.", label=f"UCL={item['ucl']:.4g}")
            ax.axhline(item["lcl"], color="red", linestyle="-.", label=f"LCL={item['lcl']:.4g}")
            bad = np.where((values > item["ucl"]) | (values < item["lcl"]))[0]
            if len(bad):
                ax.scatter(x[bad], values[bad], color="red", s=45, zorder=5, label="Hors contrôle")
            ax.set_title(item["title"])
            ax.set_xlabel("Sous-groupe")
            ax.set_ylabel(item["ylabel"])
            ax.grid(True, alpha=0.25)
            ax.legend(fontsize=8)

        self.cc_canvas.fig.suptitle(chart_label, fontsize=12, fontweight="bold")
        try:
            self.cc_canvas.fig.tight_layout(rect=[0, 0, 1, 0.96])
        except Exception:
            pass
        self.cc_canvas.draw()

    def _plot_control_charts(self, data, cc_type, subgroup):
        self.cc_canvas.fig.clear()
        data = np.asarray(data, dtype=float)
        data = data[~np.isnan(data)]

        if cc_type in ("X̄-R", "Xbar-R", "X-R"):
            stats_xr = self._cc_xbar_stats(data, subgroup, "XBAR_R")
            self._plot_xbar_control_pair(stats_xr, f"Carte X̄-R — sous-groupe {subgroup}")
            return

        if cc_type in ("X̄-S", "Xbar-S", "X-S"):
            stats_xs = self._cc_xbar_stats(data, subgroup, "XBAR_S")
            self._plot_xbar_control_pair(stats_xs, f"Carte X̄-S — sous-groupe {subgroup}")
            return

        mean_val = np.mean(data)
        std_val = np.std(data, ddof=1) if len(data) > 1 else 0
        if cc_type in ["P", "U", "C"]:
            if cc_type == "P":
                n = self.cc_sample_size.value()
                y = data if np.nanmax(data) <= 1 else data / n
                center = np.mean(y)
                sigma = np.sqrt(max(center * (1 - center) / n, 0))
                ucl = min(1, center + 3 * sigma)
                lcl = max(0, center - 3 * sigma)
                ylabel = "Proportion"
            elif cc_type == "U":
                opp = self.cc_opportunities.value()
                y = data / opp
                center = np.mean(y)
                sigma = np.sqrt(center / opp)
                ucl = center + 3 * sigma
                lcl = max(0, center - 3 * sigma)
                ylabel = "Défauts / unité"
            else:
                y = data
                center = np.mean(y)
                sigma = np.sqrt(center)
                ucl = center + 3 * sigma
                lcl = max(0, center - 3 * sigma)
                ylabel = "Nombre de défauts"
            ax = self.cc_canvas.fig.add_subplot(111)
            ax.plot(range(1, len(y) + 1), y, "bo-", markersize=4)
            ax.axhline(center, color="green", linestyle="--", label="CL")
            ax.axhline(ucl, color="red", linestyle="-.", label="UCL")
            ax.axhline(lcl, color="red", linestyle="-.", label="LCL")
            bad = np.where((y > ucl) | (y < lcl))[0]
            if len(bad):
                ax.scatter(bad + 1, y[bad], color="red", s=40, zorder=5, label="Hors contrôle")
            ax.set_title(f"Carte {cc_type}")
            ax.set_ylabel(ylabel)
            ax.set_xlabel("Observation")
            ax.grid(True, alpha=0.25)
            ax.legend(fontsize=8)
            self.cc_canvas.fig.tight_layout()
            self.cc_canvas.draw()
            return

        ucl = mean_val + 3 * std_val
        lcl = mean_val - 3 * std_val
        if cc_type == "EWMA":
            lam = self.cc_lambda.value()
            z = np.zeros(len(data))
            z[0] = mean_val
            for i in range(1, len(data)):
                z[i] = lam * data[i] + (1 - lam) * z[i-1]
            sigma_z = std_val * np.sqrt(lam / (2 - lam) * (1 - (1 - lam)**(2 * np.arange(1, len(data) + 1))))
            ucl_ewma = mean_val + 3 * sigma_z
            lcl_ewma = mean_val - 3 * sigma_z
            ax1 = self.cc_canvas.fig.add_subplot(211)
            ax2 = self.cc_canvas.fig.add_subplot(212)
            ax1.plot(range(len(data)), data, "bo-", markersize=3, alpha=0.5, label="Données")
            ax1.axhline(mean_val, color="green", linestyle="--", label="CL")
            ax1.set_title("Données individuelles")
            ax1.grid(True, alpha=0.25)
            ax1.legend(fontsize=7)
            ax2.plot(range(1, len(data)+1), z, "ro-", markersize=3, label="EWMA")
            ax2.plot(range(1, len(data)+1), ucl_ewma, "r--", linewidth=1, label="UCL")
            ax2.plot(range(1, len(data)+1), lcl_ewma, "r--", linewidth=1, label="LCL")
            ax2.axhline(mean_val, color="green", linestyle="--", label="CL")
            ax2.set_title("Carte EWMA")
            ax2.grid(True, alpha=0.25)
            ax2.legend(fontsize=7)
        elif cc_type == "CUSUM":
            k = self.cc_k.value() * std_val
            h = 5 * std_val
            s_pos = np.zeros(len(data))
            s_neg = np.zeros(len(data))
            for i in range(1, len(data)):
                s_pos[i] = max(0, s_pos[i-1] + data[i] - mean_val - k)
                s_neg[i] = max(0, s_neg[i-1] + mean_val - data[i] - k)
            ax1 = self.cc_canvas.fig.add_subplot(211)
            ax2 = self.cc_canvas.fig.add_subplot(212)
            ax1.plot(range(1, len(data)+1), s_pos, "b-", label="CUSUM+")
            ax1.axhline(h, color="red", linestyle="--", label="h")
            ax1.set_title("CUSUM positif")
            ax1.grid(True, alpha=0.25)
            ax1.legend(fontsize=7)
            ax2.plot(range(1, len(data)+1), s_neg, "m-", label="CUSUM-")
            ax2.axhline(h, color="red", linestyle="--", label="h")
            ax2.set_title("CUSUM négatif")
            ax2.grid(True, alpha=0.25)
            ax2.legend(fontsize=7)
        else:
            ax = self.cc_canvas.fig.add_subplot(111)
            ax.plot(range(1, len(data) + 1), data, "bo-", markersize=4)
            ax.axhline(mean_val, color="green", linestyle="--", label="CL")
            ax.axhline(ucl, color="red", linestyle="-.", label="UCL")
            ax.axhline(lcl, color="red", linestyle="-.", label="LCL")
            bad = np.where((data > ucl) | (data < lcl))[0]
            if len(bad):
                ax.scatter(bad + 1, data[bad], color="red", s=40, zorder=5, label="Hors contrôle")
            ax.set_title(f"Carte {cc_type}")
            ax.set_xlabel("Observation")
            ax.set_ylabel("Valeur")
            ax.grid(True, alpha=0.25)
            ax.legend(fontsize=8)

        try:
            self.cc_canvas.fig.tight_layout()
        except Exception:
            pass
        self.cc_canvas.draw()

    def _create_probplot_tab(self):
        tab = QWidget()
        self.tabs.addTab(tab, "Graphiques probabilité")
        left, left_layout, right, self.prob_canvas = self._setup_analysis_layout(tab)
        col_row = QHBoxLayout()
        self.prob_col_combo = QComboBox()
        self.prob_col_combo.setMinimumWidth(80)
        col_row.addWidget(QLabel("Colonne :"))
        col_row.addWidget(self.prob_col_combo)
        refresh_btn = QPushButton("↻")
        refresh_btn.setMaximumWidth(30)
        refresh_btn.clicked.connect(self._refresh_all_combos)
        col_row.addWidget(refresh_btn)
        left_layout.addLayout(col_row)
        params_group = QGroupBox("Paramètres")
        params_layout = QFormLayout(params_group)
        self.prob_dist = QComboBox()
        self.prob_dist.addItems([
            "Normale", "Log-Normale", "Weibull (2P)", "Weibull (3P)",
            "Exponentielle", "Gamma", "Logistique", "Gumbel (max)",
            "Gumbel (min)", "Cauchy", "Rayleigh", "Uniforme",
            "Student-t", "Laplace", "Beta", "Pareto",
        ])
        self._tip(self.prob_dist, "Distribution théorique pour le graphique de probabilité")
        params_layout.addRow("Distribution :", self.prob_dist)
        left_layout.addWidget(params_group)
        calc_btn = QPushButton(" Tracer")
        calc_btn.setMinimumHeight(40)
        calc_btn.clicked.connect(self._run_probplot)
        left_layout.addWidget(calc_btn)
        result_group = QGroupBox("Résultats")
        result_layout = QVBoxLayout(result_group)
        self.prob_result_text = QTextEdit()
        self.prob_result_text.setReadOnly(True)
        self.prob_result_text.setFont(QFont("Courier", 10))
        result_layout.addWidget(self.prob_result_text)
        self._add_copy_button(result_layout, self.prob_result_text)
        left_layout.addWidget(result_group)

    def _run_probplot(self):
        data, col_name = self._get_combo_data(self.prob_col_combo)
        if data is None:
            QMessageBox.warning(self, "Attention", "Sélectionnez une colonne")
            return
        try:
            dist_map = {
                "Normale": "normal", "Log-Normale": "lognormal",
                "Weibull (2P)": "weibull_min", "Weibull (3P)": "weibull_min",
                "Exponentielle": "exponential", "Gamma": "gamma",
                "Logistique": "logistic", "Gumbel (max)": "gumbel_max",
                "Gumbel (min)": "gumbel_min", "Cauchy": "cauchy",
                "Rayleigh": "rayleigh", "Uniforme": "uniform",
                "Student-t": "t", "Laplace": "laplace",
                "Beta": "beta", "Pareto": "pareto",
            }
            dist = self.prob_dist.currentText()
            analysis = ProbabilityPlot(data, distribution=dist_map.get(dist, "normal"))
            self.prob_result_text.setText(analysis.get_summary())
            self._plot_probplot(data, dist_map.get(dist, "normal"))
        except Exception as e:
            QMessageBox.critical(self, "Erreur", str(e))

    def _plot_probplot(self, data, dist):
        self.prob_canvas.fig.clear()
        ax = self.prob_canvas.fig.add_subplot(111)

        dist_map = {
            "normal": "norm", "lognormal": "lognorm",
            "weibull_min": "weibull_min", "exponential": "expon",
            "gamma": "gamma", "logistic": "logistic",
            "gumbel_max": "gumbel_r", "gumbel_min": "gumbel_l",
            "cauchy": "cauchy", "rayleigh": "rayleigh",
            "uniform": "uniform", "t": "t",
            "laplace": "laplace", "beta": "beta",
            "pareto": "pareto",
        }
        d = dist_map.get(dist, "norm")

        if dist == "lognormal":
            valid = data[data > 0]
            if len(valid) > 0:
                scipy_stats.probplot(np.log(valid), dist="norm", plot=ax)
            else:
                ax.text(0.5, 0.5, "Données invalides", ha='center', va='center', transform=ax.transAxes)
        elif dist == "beta":
            d_min, d_max = np.min(data), np.max(data)
            rng = d_max - d_min
            if rng == 0:
                ax.text(0.5, 0.5, "Données constantes", ha='center', va='center', transform=ax.transAxes)
            else:
                eps = rng * 0.001
                scaled = (data - d_min + eps) / (rng + 2 * eps)
                a, b = scipy_stats.beta.fit(scaled, floc=0, fscale=1)
                n = len(data)
                sorted_data = np.sort(data)
                percentiles = np.arange(1, n + 1) / (n + 1)
                theoretical_scaled = scipy_stats.beta.ppf(percentiles, a, b, 0, 1)
                theoretical_orig = d_min + theoretical_scaled * rng
                ax.plot(theoretical_orig, sorted_data, 'o', markersize=3, alpha=0.7)
                min_val = min(theoretical_orig.min(), sorted_data.min())
                max_val = max(theoretical_orig.max(), sorted_data.max())
                ax.plot([min_val, max_val], [min_val, max_val], 'r-', linewidth=1.5)
        else:
            dist_obj = getattr(scipy_stats, d)
            try:
                params = dist_obj.fit(data)
                n = len(data)
                sorted_data = np.sort(data)
                percentiles = np.arange(1, n + 1) / (n + 1)
                theoretical = dist_obj.ppf(percentiles, *params)
                ax.plot(theoretical, sorted_data, 'o', markersize=3, alpha=0.7)
                min_val = min(theoretical.min(), sorted_data.min())
                max_val = max(theoretical.max(), sorted_data.max())
                ax.plot([min_val, max_val], [min_val, max_val], 'r-', linewidth=1.5)
            except Exception:
                scipy_stats.probplot(data, dist=d, plot=ax)

        ax.set_title(f"Graphique de probabilité - {dist}")
        ax.set_xlabel("Valeurs théoriques")
        ax.set_ylabel("Valeurs observées")
        ax.grid(True, alpha=0.3)
        self.prob_canvas.fig.tight_layout()
        self.prob_canvas.draw()

    # ===== MSA =====
    def _create_msa_tab(self):
        tab = QWidget()
        self.tabs.addTab(tab, "MSA / Gage R&R")
        left, left_layout, right, self.msa_canvas = self._setup_analysis_layout(tab)
        info = QLabel("Format : 3 colonnes — Pièces, Opérateurs, Mesures")
        info.setStyleSheet("color: #2c3e50; font-size: 11px; font-weight: bold;")
        left_layout.addWidget(info)

        col_row = QHBoxLayout()
        self.msa_part_combo = QComboBox()
        self.msa_part_combo.setMinimumWidth(115)
        col_row.addWidget(QLabel("Pièces :"))
        col_row.addWidget(self.msa_part_combo)
        self.msa_op_combo = QComboBox()
        self.msa_op_combo.setMinimumWidth(115)
        col_row.addWidget(QLabel("Opérateurs :"))
        col_row.addWidget(self.msa_op_combo)
        self.msa_meas_combo = QComboBox()
        self.msa_meas_combo.setMinimumWidth(115)
        col_row.addWidget(QLabel("Mesures :"))
        col_row.addWidget(self.msa_meas_combo)
        refresh_btn = QPushButton("↻")
        refresh_btn.setMaximumWidth(30)
        refresh_btn.clicked.connect(self._refresh_all_combos)
        col_row.addWidget(refresh_btn)
        left_layout.addLayout(col_row)

        params_group = QGroupBox("Paramètres")
        params_layout = QFormLayout(params_group)

        self.msa_n_parts = QSpinBox()
        self.msa_n_parts.setRange(2, 50)
        self.msa_n_parts.setValue(10)
        self._tip(self.msa_n_parts, "Nombre de pièces dans l'étude. Typiquement 10.")
        params_layout.addRow("Nombre de pièces :", self.msa_n_parts)

        self.msa_n_operators = QSpinBox()
        self.msa_n_operators.setRange(2, 10)
        self.msa_n_operators.setValue(3)
        self._tip(self.msa_n_operators, "Nombre d'opérateurs. Typiquement 2-3.")
        params_layout.addRow("Nombre d'opérateurs :", self.msa_n_operators)

        self.msa_n_trials = QSpinBox()
        self.msa_n_trials.setRange(1, 10)
        self.msa_n_trials.setValue(3)
        self._tip(self.msa_n_trials, "Nombre de répétitions par opérateur/pièce. Typiquement 2-3.")
        params_layout.addRow("Nombre de répétitions :", self.msa_n_trials)

        self.msa_tolerance = QDoubleSpinBox()
        self.msa_tolerance.setRange(0.001, 1e9)
        self.msa_tolerance.setValue(1.0)
        self._tip(self.msa_tolerance, "Tolérance du processus (USL - LSL) pour le calcul P/T")
        params_layout.addRow("Tolérance :", self.msa_tolerance)

        left_layout.addWidget(params_group)
        calc_btn = QPushButton(" Calculer Gage R&R")
        calc_btn.setMinimumHeight(40)
        calc_btn.clicked.connect(self._run_msa)
        left_layout.addWidget(calc_btn)
        result_group = QGroupBox("Résultats")
        result_layout = QVBoxLayout(result_group)
        self.msa_result_text = QTextEdit()
        self.msa_result_text.setReadOnly(True)
        self.msa_result_text.setFont(QFont("Courier", 10))
        result_layout.addWidget(self.msa_result_text)
        self._add_copy_button(result_layout, self.msa_result_text)
        left_layout.addWidget(result_group)

    def _run_msa(self):
        parts_col, _ = self._get_combo_data(self.msa_part_combo)
        ops_col, _ = self._get_combo_data(self.msa_op_combo)
        meas_col, _ = self._get_combo_data(self.msa_meas_combo)
        if parts_col is None or ops_col is None or meas_col is None:
            QMessageBox.warning(self, "Attention", "Sélectionnez Pièces, Opérateurs et Mesures")
            return

        try:
            n_parts = self.msa_n_parts.value()
            n_operators = self.msa_n_operators.value()
            n_trials = self.msa_n_trials.value()
            tolerance = self.msa_tolerance.value()

            min_len = min(len(parts_col), len(ops_col), len(meas_col))
            parts_col = np.array(parts_col[:min_len])
            ops_col = np.array(ops_col[:min_len])
            meas_col = np.array(meas_col[:min_len])

            valid = ~np.isnan(parts_col) & ~np.isnan(ops_col) & ~np.isnan(meas_col)
            parts_col = parts_col[valid]
            ops_col = ops_col[valid]
            meas_col = meas_col[valid]

            if len(parts_col) == 0:
                QMessageBox.warning(self, "Attention", "Aucune donnée valide (sans NaN)")
                return

            unique_parts = np.unique(parts_col)
            unique_ops = np.unique(ops_col)

            data_3d = np.zeros((n_operators, n_parts, n_trials))
            for i, op in enumerate(sorted(unique_ops)):
                if i >= n_operators:
                    break
                for j, part in enumerate(sorted(unique_parts)):
                    if j >= n_parts:
                        break
                    mask = (parts_col == part) & (ops_col == op)
                    vals = meas_col[mask][:n_trials]
                    data_3d[i, j, :len(vals)] = vals

            study = GageRRStudy(data_3d, n_operators=n_operators, n_parts=n_parts,
                                n_trials=n_trials, tolerance=tolerance)
            self.msa_result_text.setText(study.get_summary())
            results = study.get_results()
            self._plot_msa(results)
        except Exception as e:
            QMessageBox.critical(self, "Erreur", str(e))

    def _plot_msa(self, results):
        self.msa_canvas.fig.clear()
        ax1 = self.msa_canvas.fig.add_subplot(221)
        ax2 = self.msa_canvas.fig.add_subplot(222)
        ax3 = self.msa_canvas.fig.add_subplot(223)
        ax4 = self.msa_canvas.fig.add_subplot(224)

        labels = ["Répétabilité", "Reproductibilité", "Pièces", "Total"]
        values = [
            results.get('pct_contribution_repeatability', 0),
            results.get('pct_contribution_reproducibility', 0),
            results.get('pct_contribution_part', 0),
            results.get('pct_contribution_gage_rr', 0),
        ]
        colors_bar = ['#3498db', '#e74c3c', '#2ecc71', '#9b59b6']
        ax1.bar(labels, values, color=colors_bar)
        ax1.set_ylabel("% Contribution")
        ax1.set_title("Composantes de variance")
        ax1.tick_params(axis='x', rotation=15)

        op_means = results.get('op_means', [])
        if op_means:
            ax2.bar(range(len(op_means)), op_means, color='#3498db')
            ax2.set_xticks(range(len(op_means)))
            ax2.set_xticklabels([f"Op{i+1}" for i in range(len(op_means))])
            ax2.set_title("Moyennes par opérateur")

        part_means = results.get('part_means', [])
        if part_means:
            ax3.plot(range(len(part_means)), part_means, 'bo-', markersize=4)
            ax3.set_xticks(range(len(part_means)))
            ax3.set_xticklabels([f"P{i+1}" for i in range(len(part_means))], rotation=45, ha='right')
            ax3.set_title("Moyennes par pièce")

        rr_pct = results.get('pct_contribution_gage_rr', 0)
        part_pct = results.get('pct_contribution_part', 0)
        ax4.pie([rr_pct, max(0, 100 - rr_pct)], labels=['Gage R&R', 'Autre'],
                autopct='%1.1f%%', colors=['#e74c3c', '#2ecc71'])
        ax4.set_title("Répartition % contribution")

        self.msa_canvas.fig.tight_layout()
        self.msa_canvas.draw()

    def _export_results(self):
        texts = []
        for title, widget in [("CAPABILITÉ", self.cap_result_text), ("NORMALITÉ", self.norm_result_text),
                               ("VALEURS ABERRANTES", self.out_result_text), ("CARTES DE CONTRÔLE", self.cc_result_text),
                               ("GRAPHIQUES PROBABILITÉ", self.prob_result_text),
                               ("RÉGRESSION", self.reg_result_text), ("TEST T", self.tt_result_text),
                               ("ANOVA", self.anova_result_text), ("CORRÉLATION", self.corr_result_text),
                               ("BOXPLOTS", self.boxplot_result_text),
                               ("MSA / GAGE R&R", self.msa_result_text),
                               ("IDENTIFICATION DISTRIBUTION", self.dist_result_text)]:
            if widget.toPlainText().strip():
                texts.append((title, widget.toPlainText()))
        if not texts:
            QMessageBox.warning(self, "Attention", "Aucun résultat à exporter")
            return

        file_filters = (
            "Text (*.txt);;"
            "PDF (*.pdf);;"
            "ODT (*.odt);;"
            "DOCX (*.docx)"
        )
        filepath, _ = QFileDialog.getSaveFileName(self, "Exporter résultats", "", file_filters)
        if not filepath:
            return

        try:
            ext = os.path.splitext(filepath)[1].lower()

            if ext == '.txt':
                with open(filepath, "w", encoding="utf-8") as f:
                    f.write("RÉSULTATS D'ANALYSE STATISTIQUE\n")
                    f.write("=" * 60 + "\n")
                    f.write(f"Date : {pd.Timestamp.now().strftime('%Y-%m-%d %H:%M')}\n\n")
                    for title, content in texts:
                        f.write(f"{title}\n{content}\n\n")

            elif ext == '.pdf':
                exporter = ReportExporter("Rapport StatPro", filepath)
                exporter.add_section("RÉSULTATS D'ANALYSE STATISTIQUE",
                                     f"Date : {pd.Timestamp.now().strftime('%Y-%m-%d %H:%M')}")
                for title, content in texts:
                    exporter.add_section(title, content)
                exporter.export(filepath)

            elif ext == '.odt':
                from docx import Document as DocxDocument
                from io import BytesIO
                import subprocess

                doc = DocxDocument()
                doc.add_heading("RÉSULTATS D'ANALYSE STATISTIQUE", level=1)
                doc.add_paragraph(f"Date : {pd.Timestamp.now().strftime('%Y-%m-%d %H:%M')}")
                for title, content in texts:
                    doc.add_heading(title, level=2)
                    doc.add_paragraph(content)
                tmp_docx = filepath.replace('.odt', '.docx')
                doc.save(tmp_docx)
                try:
                    subprocess.run(["libreoffice", "--headless", "--convert-to", "odt", tmp_docx],
                                   capture_output=True, timeout=30)
                    if os.path.exists(filepath):
                        os.remove(tmp_docx)
                    elif os.path.exists(tmp_docx):
                        import shutil
                        shutil.move(tmp_docx, filepath)
                except Exception:
                    if os.path.exists(tmp_docx):
                        os.rename(tmp_docx, filepath)

            elif ext == '.docx':
                from docx import Document
                doc = Document()
                doc.add_heading("RÉSULTATS D'ANALYSE STATISTIQUE", level=1)
                doc.add_paragraph(f"Date : {pd.Timestamp.now().strftime('%Y-%m-%d %H:%M')}")
                for title, content in texts:
                    doc.add_heading(title, level=2)
                    doc.add_paragraph(content)
                doc.save(filepath)

            self.status_bar.showMessage(f"Résultats exportés : {filepath}")
        except ImportError:
            QMessageBox.warning(self, "Export", "La bibliothèque 'python-docx' n'est pas installée.\nInstallez-la avec : pip install python-docx")
        except Exception as e:
            QMessageBox.critical(self, "Erreur", f"Impossible d'exporter :\n{e}")


    # ===== MINIQUAL INTÉGRÉ À STATPRO =====
    def _create_miniqual_tab(self):
        tab = QWidget(); self.tabs.addTab(tab, "MiniQual")
        layout = QHBoxLayout(tab); layout.setContentsMargins(5,5,5,5)
        splitter = QSplitter(Qt.Horizontal); splitter.setChildrenCollapsible(False); layout.addWidget(splitter)
        left = QWidget(); left.setMinimumWidth(380); left.setMaximumWidth(540)
        left_layout = QVBoxLayout(left); left_layout.setContentsMargins(0,0,0,0)
        params_group = QGroupBox("Capabilité MiniQual"); params = QFormLayout(params_group)
        self.miniqual_col_combo = QComboBox(); params.addRow("Colonne mesure", self.miniqual_col_combo)
        self.miniqual_distribution = QComboBox(); self.miniqual_distribution.addItems(["Normal","Log-Normale","Weibull (2P)","Exponentielle","Gamma","Logistique","Gumbel (max)","Cauchy","Rayleigh","Uniforme","Student-t","Laplace"]); params.addRow("Loi statistique", self.miniqual_distribution)
        self.miniqual_lsl=QDoubleSpinBox(); self.miniqual_lsl.setRange(-1e9,1e9); self.miniqual_lsl.setSpecialValueText("Aucune"); self.miniqual_lsl.setValue(self.miniqual_lsl.minimum()); params.addRow("LSL", self.miniqual_lsl)
        self.miniqual_usl=QDoubleSpinBox(); self.miniqual_usl.setRange(-1e9,1e9); self.miniqual_usl.setSpecialValueText("Aucune"); self.miniqual_usl.setValue(self.miniqual_usl.minimum()); params.addRow("USL", self.miniqual_usl)
        self.miniqual_target=QDoubleSpinBox(); self.miniqual_target.setRange(-1e9,1e9); self.miniqual_target.setSpecialValueText("Aucune"); self.miniqual_target.setValue(self.miniqual_target.minimum()); params.addRow("Cible", self.miniqual_target)
        self.miniqual_chi2_mode=QComboBox(); self.miniqual_chi2_mode.addItems(["auto","nombre"]); params.addRow("Mode Khi²", self.miniqual_chi2_mode)
        self.miniqual_chi2_bins=QSpinBox(); self.miniqual_chi2_bins.setRange(4,50); self.miniqual_chi2_bins.setValue(10); params.addRow("Nombre de classes Khi²", self.miniqual_chi2_bins)
        self.miniqual_cpk_accept=QDoubleSpinBox(); self.miniqual_cpk_accept.setRange(.1,10); self.miniqual_cpk_accept.setValue(1.33); self.miniqual_cpk_accept.setSingleStep(.01); params.addRow("Seuil Cpk accept.", self.miniqual_cpk_accept)
        self.miniqual_cpk_excellent=QDoubleSpinBox(); self.miniqual_cpk_excellent.setRange(.1,10); self.miniqual_cpk_excellent.setValue(1.67); self.miniqual_cpk_excellent.setSingleStep(.01); params.addRow("Seuil Cpk excellent", self.miniqual_cpk_excellent)
        self.miniqual_boxcox=QCheckBox("Activer transformation Box-Cox"); params.addRow(self.miniqual_boxcox)
        self.miniqual_exclude_outliers=QCheckBox("Exclure valeurs aberrantes et recalculer"); params.addRow(self.miniqual_exclude_outliers)
        left_layout.addWidget(params_group)
        for label, cb in [("Calculer/afficher p-values lois", self._miniqual_show_pvalues), ("Générer rapport DOCX MiniQual", self._miniqual_generate_docx)]:
            b=QPushButton(label); b.clicked.connect(cb); left_layout.addWidget(b)
        sess=QGroupBox("Session / Résultats"); sl=QVBoxLayout(sess); self.miniqual_text=QTextEdit(); self.miniqual_text.setReadOnly(True); self.miniqual_text.setFont(QFont("Courier",10)); sl.addWidget(self.miniqual_text); left_layout.addWidget(sess,1)
        right=QWidget(); rl=QVBoxLayout(right); rl.setContentsMargins(0,0,0,0); self.miniqual_canvas=MplCanvas(right,width=7,height=8); rl.addWidget(self.miniqual_canvas); tb=SafeNavigationToolbar(self.miniqual_canvas,right); tb.setObjectName("plotToolbar"); rl.addWidget(tb)
        splitter.addWidget(left); splitter.addWidget(right); splitter.setSizes([430,900]); self.miniqual_last_out=None; self._refresh_all_combos()

    def _miniqual_log(self,msg=""):
        if hasattr(self,"miniqual_text"): self.miniqual_text.append(str(msg))

    def _miniqual_params(self):
        col=self.miniqual_col_combo.currentText()
        if not col or col=="(aucune donnée)": raise ValueError("Sélectionnez une colonne mesure contenant des données.")
        df=self._sheet_to_dataframe()
        return dict(df=df,column=col,lsl=(None if self.miniqual_lsl.value()==self.miniqual_lsl.minimum() else self.miniqual_lsl.value()),usl=(None if self.miniqual_usl.value()==self.miniqual_usl.minimum() else self.miniqual_usl.value()),target=(None if self.miniqual_target.value()==self.miniqual_target.minimum() else self.miniqual_target.value()),distribution=self.miniqual_distribution.currentText(),chi2_bins=("auto" if self.miniqual_chi2_mode.currentText()=="auto" else self.miniqual_chi2_bins.value()),boxcox=self.miniqual_boxcox.isChecked(),exclude=self.miniqual_exclude_outliers.isChecked(),cpk_accept=self.miniqual_cpk_accept.value(),cpk_excellent=self.miniqual_cpk_excellent.value())

    def _miniqual_dist(self,name):
        return {"Normal":("normal","normale",scipy_stats.norm,"norm"),"Log-Normale":("lognormal","lognormale",scipy_stats.lognorm,"lognorm"),"Weibull (2P)":("weibull","Weibull",scipy_stats.weibull_min,"weibull_min"),"Exponentielle":("exponential","exponentielle",scipy_stats.expon,"expon"),"Gamma":("gamma","gamma",scipy_stats.gamma,"gamma"),"Logistique":("logistic","logistique",scipy_stats.logistic,"logistic"),"Gumbel (max)":("gumbel","Gumbel",scipy_stats.gumbel_r,"gumbel_r"),"Cauchy":("cauchy","Cauchy",scipy_stats.cauchy,"cauchy"),"Rayleigh":("rayleigh","Rayleigh",scipy_stats.rayleigh,"rayleigh"),"Uniforme":("uniform","uniforme",scipy_stats.uniform,"uniform"),"Student-t":("student_t","Student-t",scipy_stats.t,"t"),"Laplace":("laplace","Laplace",scipy_stats.laplace,"laplace")}.get(name,("normal","normale",scipy_stats.norm,"norm"))

    def _miniqual_filter(self,x,key):
        return x[x>0] if key in ("lognormal","weibull","gamma","rayleigh") else (x[x>=0] if key=="exponential" else x)

    def _miniqual_fit(self,s,name):
        key,label,obj,scipy_name=self._miniqual_dist(name); x=pd.Series(s).dropna().astype(float).to_numpy(); xp=self._miniqual_filter(x,key)
        if len(xp)<3: return dict(loi=key,p_value=np.nan,statistique=np.nan,paramètres="Données incompatibles",params=None,label=label,obj=obj)
        try:
            params=scipy_stats.norm.fit(xp) if key=="normal" else (obj.fit(xp,floc=0) if key in ("lognormal","weibull","gamma","exponential","rayleigh") else obj.fit(xp))
            st,pv=scipy_stats.kstest(xp,scipy_name,args=params); return dict(loi=key,p_value=float(pv),statistique=float(st),paramètres=str(tuple(round(float(v),6) for v in params)),params=params,label=label,obj=obj)
        except Exception as e: return dict(loi=key,p_value=np.nan,statistique=np.nan,paramètres=f"Erreur : {e}",params=None,label=label,obj=obj)

    def _miniqual_pvalues_rows(self,s):
        names=["Normal","Log-Normale","Weibull (2P)","Exponentielle","Gamma","Logistique","Gumbel (max)","Cauchy","Rayleigh","Uniforme","Student-t","Laplace"]
        rows=[{k:v for k,v in self._miniqual_fit(s,n).items() if k not in ("params","label","obj")} for n in names]
        return sorted(rows,key=lambda r:-r["p_value"] if pd.notna(r["p_value"]) else 1e9)

    def _miniqual_compute(self):
        p=self._miniqual_params(); s=numeric_series(p["df"],p["column"]); validation=validate_capability(p["df"],p["column"],p["lsl"],p["usl"],p["target"],None,"normal")
        analysis_s=s; lsl=p["lsl"]; usl=p["usl"]; target=p["target"]; info={}
        if p["exclude"]:
            outs=outlier_tests(s); vals=set(outs.get("values_to_exclude",[]))
            if vals:
                before=len(analysis_s); analysis_s=analysis_s[~analysis_s.astype(float).isin(vals)]; info["Valeurs aberrantes exclues"]=before-len(analysis_s)
        if p["boxcox"]:
            analysis_s,lsl,usl,target,bc=boxcox_transform(analysis_s,lsl,usl,target); info.update(bc)
        res=capability(analysis_s,lsl,usl,target); res["loi choisie"]=p["distribution"]; res["Box-Cox activé"]=p["boxcox"]; res.update(info)
        return p,s,analysis_s,res,validation

    def _miniqual_dashboard(self,s,res,out_png=None):
        p=self._miniqual_params(); data=pd.Series(s).dropna().astype(float).to_numpy(); fig=self.miniqual_canvas.fig; fig.clear(); fig.set_size_inches(15,12); ax=fig.subplots(2,2); fig.suptitle("Analyse de capabilité du procédé",fontsize=18,fontweight="bold")
        m=float(np.mean(data)); sd=float(np.std(data,ddof=1)) if len(data)>1 else 0.0; a=ax[0,0]; a.hist(data,bins=min(15,max(5,int(np.sqrt(len(data))))),density=True,alpha=.75,edgecolor="black")
        xs=[data.min(),data.max(),m-4*sd,m+4*sd]+[v for v in [res.get("lsl"),res.get("usl"),res.get("target")] if v is not None]
        xx=np.linspace(min(xs),max(xs),300); fit=self._miniqual_fit(data,p["distribution"])
        if fit.get("params") is not None:
            try: a.plot(xx,fit["obj"].pdf(xx,*fit["params"]),"r-",label=f"Courbe {fit['label']} estimée")
            except Exception: pass
        for key,c,lab,ls in [("lsl","green","LSL","--"),("usl","red","USL","--"),("target","purple","Cible",":")]:
            if res.get(key) is not None: a.axvline(res[key],color=c,ls=ls,label=f"{lab}={res[key]:.3g}")
        a.axvline(m,color="orange",label=f"Moyenne={m:.3g}"); a.legend(fontsize=8); a.set_title("Distribution")
        ax[0,1].boxplot(data); ax[0,1].set_title("Boxplot")
        ax[1,0].set_title(f"Q-Q plot — {p['distribution']}"); key=fit.get("loi"); xp=self._miniqual_filter(data,key)
        if fit.get("params") is not None and len(xp)>=3:
            try:
                probs=(np.arange(1,len(xp)+1)-.5)/len(xp); q=fit["obj"].ppf(probs,*fit["params"]); ordered=np.sort(xp); ax[1,0].scatter(q,ordered,s=12); lo=min(q.min(),ordered.min()); hi=max(q.max(),ordered.max()); ax[1,0].plot([lo,hi],[lo,hi],"r-")
            except Exception: ax[1,0].text(.5,.5,"Q-Q plot indisponible",ha="center",va="center")
        labels=[]; vals=[]
        for lab,key in [("Cp","cp"),("Cpk inf.","cpk_lower"),("Cpk sup.","cpk_upper"),("Cpk global","cpk")]:
            v=res.get(key)
            if v is not None and np.isfinite(v): labels.append(lab); vals.append(float(v))
        if vals: ax[1,1].bar(labels,vals); ax[1,1].axhline(p["cpk_accept"],color="orange",ls="--"); ax[1,1].axhline(p["cpk_excellent"],color="green",ls="--")
        ax[1,1].set_title("Indices"); fig.tight_layout(rect=[0,0,1,.96])
        if out_png: fig.savefig(out_png,dpi=180,bbox_inches="tight")
        self.miniqual_canvas.draw()

    def _miniqual_show_pvalues(self):
        try:
            p,s,analysis_s,res,validation=self._miniqual_compute(); self._miniqual_log("=== P-VALUES PAR LOI STATISTIQUE ===")
            for r in self._miniqual_pvalues_rows(s): self._miniqual_log(f"{r['loi']:<12} p={r['p_value']:.5g} stat={r['statistique']:.5g}")
            self._miniqual_dashboard(analysis_s,res)
        except Exception as e: self._miniqual_log(f"ERREUR p-values : {e}"); QMessageBox.warning(self,"MiniQual",str(e))

    def _miniqual_generate_docx(self):
        out_dir=QFileDialog.getExistingDirectory(self,"Choisir le dossier de sortie MiniQual","")
        if not out_dir: return
        out=Path(out_dir); out.mkdir(parents=True,exist_ok=True)
        try:
            p,s,analysis_s,res,validation=self._miniqual_compute(); distrows=self._miniqual_pvalues_rows(s); outs=outlier_tests(s); norm=normality_tests(s); chi2=chi_square_gof(s,"normal",p["chi2_bins"]); nn=nonnormal_capability(s,p["lsl"],p["usl"])
            summary={"Statut global":"OK" if (res.get("cpk") or 0)>=p["cpk_accept"] and res.get("observed_nc_count",0)==0 else "À SURVEILLER / ACTION À ÉVALUER","Cpk retenu":res.get("cpk"),"Statut validation fichier":status(validation),"Khi² p-value":chi2.get("p-value"),"Colonne mesure":p["column"],"Loi choisie":p["distribution"]}
            png=out/"capability_dashboard.png"; self._miniqual_dashboard(analysis_s,res,png)
            sections=[("Résumé décisionnel",summary),("Validation du fichier d’entrée",validation),("P-values par loi statistique",distrows),("Données et statistiques descriptives",descriptive(s)),("Tests de valeurs aberrantes",outs["table"]),("Tests de normalité",norm["table"]),("Test du Khi²",chi2),("Capabilité non normale / percentile",nn),("Capabilité du procédé",res),("Visualisations",str(png))]
            report_path=out/"capability_report.docx"; write_docx("Rapport de capabilité procédé",sections,report_path); self.miniqual_last_out=out; self._miniqual_log(f"Rapport DOCX MiniQual généré : {report_path}"); self.status_bar.showMessage(f"Rapport DOCX MiniQual généré : {report_path}"); QMessageBox.information(self,"MiniQual",f"Rapport DOCX généré :\n{report_path}")
        except Exception as e:
            import traceback; self._miniqual_log("=== ERREUR MINIQUAL ==="); self._miniqual_log(traceback.format_exc()); QMessageBox.critical(self,"MiniQual",f"Impossible de générer le rapport DOCX :\n{e}")

    def _create_report_tab(self):
        tab = QWidget()
        self.tabs.addTab(tab, "Rapport d'analyse")
        layout = QVBoxLayout(tab)
        info = QLabel("Rapport global généré automatiquement à partir des résultats disponibles dans les onglets.")
        info.setStyleSheet("background-color: #eaf4ff; padding: 5px; border: 1px solid #9cc7ed;")
        layout.addWidget(info)
        self.report_text = QTextEdit()
        self.report_text.setReadOnly(True)
        self.report_text.setFont(QFont("Courier", 10))
        layout.addWidget(self.report_text)
        btns = QHBoxLayout()
        refresh_btn = QPushButton(" Actualiser rapport")
        refresh_btn.clicked.connect(self._refresh_report_panel)
        export_btn = QPushButton(" Exporter PDF/DOCX")
        export_btn.clicked.connect(self._export_full_report)
        excel_btn = QPushButton(" Exporter Excel multi-feuilles")
        excel_btn.clicked.connect(self._export_excel_workbook)
        btns.addWidget(refresh_btn)
        btns.addWidget(export_btn)
        btns.addWidget(excel_btn)
        btns.addStretch()
        layout.addLayout(btns)

    def _sheet_to_dataframe(self):
        data = {}
        max_len = 0
        for c in range(self.sheet.columnCount()):
            col_data = []
            for r in range(self.sheet.rowCount()):
                item = self.sheet.item(r, c)
                col_data.append(item.text() if item and item.text() else "")
            while col_data and col_data[-1] == "":
                col_data.pop()
            if col_data:
                data[col_letter(c)] = col_data
                max_len = max(max_len, len(col_data))
        for k in data:
            data[k] += [""] * (max_len - len(data[k]))
        return pd.DataFrame(data)

    def _load_dataframe_strings(self, df):
        self.sheet.clearContents()
        max_rows = min(len(df), self.sheet.rowCount())
        max_cols = min(len(df.columns), self.sheet.columnCount())
        for c in range(max_cols):
            for r in range(max_rows):
                val = df.iloc[r, c]
                if pd.isna(val) or str(val) == "":
                    continue
                item = QTableWidgetItem(str(val))
                item.setTextAlignment(Qt.AlignRight | Qt.AlignVCenter)
                self.sheet.setItem(r, c, item)
        self._refresh_all_combos()

    def _new_project(self):
        if QMessageBox.question(self, "Nouveau projet", "Effacer les données et résultats actuels ?") != QMessageBox.Yes:
            return
        self.sheet.clearContents()
        for widget_name in ["stats_text", "cap_result_text", "norm_result_text", "out_result_text", "cc_result_text", "prob_result_text", "reg_result_text", "tt_result_text", "anova_result_text", "corr_result_text", "boxplot_result_text", "msa_result_text", "dist_result_text", "miniqual_text"]:
            if hasattr(self, widget_name):
                getattr(self, widget_name).clear()
        self.current_project_path = None
        self.analysis_report_sections = []
        self._refresh_report_panel()
        self.status_bar.showMessage("Nouveau projet")

    def _project_payload(self):
        results = {}
        for name in ["cap", "norm", "out", "cc", "prob", "reg", "tt", "anova", "corr", "boxplot", "msa", "dist"]:
            attr = f"{name}_result_text"
            if hasattr(self, attr):
                results[name] = getattr(self, attr).toPlainText()
        return {
            "version": 1,
            "theme": self.current_theme,
            "data": self._sheet_to_dataframe().to_dict(orient="list"),
            "results": results,
            "report": self.report_text.toPlainText() if hasattr(self, "report_text") else "",
        }

    def _save_project(self):
        if not self.current_project_path:
            return self._save_project_as()
        with open(self.current_project_path, "w", encoding="utf-8") as f:
            json.dump(self._project_payload(), f, ensure_ascii=False, indent=2)
        self.status_bar.showMessage(f"Projet enregistré : {self.current_project_path}")

    def _save_project_as(self):
        filepath, _ = QFileDialog.getSaveFileName(self, "Enregistrer projet", "", "Projet StatPro (*.statpro);;JSON (*.json)")
        if not filepath:
            return
        if not filepath.lower().endswith((".statpro", ".json")):
            filepath += ".statpro"
        self.current_project_path = filepath
        self._save_project()

    def _open_project(self):
        filepath, _ = QFileDialog.getOpenFileName(self, "Ouvrir projet", "", "Projet StatPro (*.statpro *.json);;All files (*)")
        if not filepath:
            return
        try:
            with open(filepath, "r", encoding="utf-8") as f:
                payload = json.load(f)
            df = pd.DataFrame(payload.get("data", {}))
            self._load_dataframe_strings(df)
            results = payload.get("results", {})
            mapping = {"cap":"cap_result_text", "norm":"norm_result_text", "out":"out_result_text", "cc":"cc_result_text", "prob":"prob_result_text", "reg":"reg_result_text", "tt":"tt_result_text", "anova":"anova_result_text", "corr":"corr_result_text", "boxplot":"boxplot_result_text", "msa":"msa_result_text", "dist":"dist_result_text"}
            for key, attr in mapping.items():
                if hasattr(self, attr):
                    getattr(self, attr).setText(results.get(key, ""))
            self.current_project_path = filepath
            self._apply_theme(payload.get("theme", "Clair"))
            self._refresh_report_panel()
            self.status_bar.showMessage(f"Projet ouvert : {filepath}")
        except Exception as e:
            QMessageBox.critical(self, "Erreur", f"Impossible d'ouvrir le projet :\n{e}")

    def _apply_theme(self, theme):
        self.current_theme = theme
        if theme == "Sombre":
            self.setStyleSheet("""
                QWidget { background-color: #232629; color: #f0f0f0; }
                QLabel { color: #f0f0f0; }
                QTextEdit, QTableWidget, QLineEdit, QComboBox, QSpinBox, QDoubleSpinBox { background-color: #2f3337; color: #f0f0f0; border: 1px solid #555; }
                QPushButton, QToolButton { background-color: #3d444b; color: #ffffff; padding: 4px; border: 1px solid #666; border-radius: 3px; }
                QPushButton:hover, QToolButton:hover { background-color: #4f5b66; }
                QToolBar { background-color: #2b3035; border: 1px solid #555; spacing: 3px; }
                QHeaderView::section { background-color: #3d444b; color: #ffffff; }
                QGroupBox { border: 1px solid #666; margin-top: 8px; color: #f0f0f0; }
                QGroupBox::title { color: #ffffff; subcontrol-origin: margin; left: 8px; padding: 0 3px; }
            """)
        elif theme == "Minitab-like":
            self.setStyleSheet("""
                QWidget { background-color: #f4f7fb; color: #102a43; }
                QLabel { color: #102a43; }
                QMenuBar { background-color: #ffffff; color: #102a43; border-bottom: 1px solid #c5d4e3; }
                QMenuBar::item:selected { background-color: #d9eafc; }
                QMenu { background-color: #ffffff; color: #102a43; border: 1px solid #9bb7d4; }
                QMenu::item:selected { background-color: #d9eafc; }

                QTabWidget::pane { border: 1px solid #9bb7d4; background-color: #ffffff; }
                QTabBar::tab { background: #dbe9f6; color: #102a43; padding: 6px 10px; border: 1px solid #9bb7d4; border-bottom: none; }
                QTabBar::tab:selected { background: #ffffff; color: #0b2239; font-weight: bold; }
                QTabBar::tab:hover { background: #eef6ff; }

                QGroupBox { border: 1px solid #7aa7d9; border-radius: 4px; margin-top: 10px; background-color: #ffffff; color: #102a43; font-weight: bold; }
                QGroupBox::title { color: #0b2239; subcontrol-origin: margin; left: 8px; padding: 0 4px; background-color: #ffffff; }

                QPushButton { background-color: #d9eafc; color: #102a43; border: 1px solid #7aa7d9; border-radius: 3px; padding: 5px 8px; }
                QPushButton:hover { background-color: #c8e0fa; border: 1px solid #4d90d9; }
                QPushButton:pressed { background-color: #b5d4f5; }
                QPushButton:disabled { background-color: #eef3f8; color: #6b7c8f; border: 1px solid #c5d4e3; }

                QToolBar#mainToolbar { background-color: #ffffff; border: 1px solid #c5d4e3; spacing: 4px; padding: 2px; }
                QToolBar#mainToolbar QToolButton { background-color: #ffffff; color: #102a43; border: 1px solid transparent; border-radius: 3px; padding: 4px 7px; }
                QToolBar#mainToolbar QToolButton:hover { background-color: #d9eafc; border: 1px solid #7aa7d9; }

                /* Barre Matplotlib sous les graphiques : fond foncé pour rendre les icônes claires visibles. */
                QToolBar#plotToolbar { background-color: #2f3b45; border: 1px solid #1f2a33; spacing: 3px; padding: 3px; }
                QToolBar#plotToolbar QToolButton { background-color: #3f4d59; color: #ffffff; border: 1px solid #5f7080; border-radius: 3px; padding: 4px; margin: 1px; }
                QToolBar#plotToolbar QToolButton:hover { background-color: #536575; border: 1px solid #8fbce8; }
                QToolBar#plotToolbar QToolButton:pressed { background-color: #1f2a33; }
                QToolBar#plotToolbar QToolButton:disabled { background-color: #44505b; color: #cfd8e3; border: 1px solid #5f7080; }

                QTextEdit, QTableWidget, QLineEdit, QComboBox, QSpinBox, QDoubleSpinBox { background-color: #ffffff; color: #102a43; border: 1px solid #7aa7d9; selection-background-color: #b5d4f5; selection-color: #102a43; }
                QTextEdit:focus, QLineEdit:focus, QComboBox:focus, QSpinBox:focus, QDoubleSpinBox:focus { border: 1px solid #2f80c9; }
                QComboBox::drop-down { border-left: 1px solid #7aa7d9; background-color: #d9eafc; width: 18px; }
                QComboBox QAbstractItemView { background-color: #ffffff; color: #102a43; selection-background-color: #d9eafc; border: 1px solid #7aa7d9; }

                QHeaderView::section { background-color: #dbe9f6; color: #102a43; font-weight: bold; border: 1px solid #7aa7d9; }
                QTableCornerButton::section { background-color: #dbe9f6; border: 1px solid #7aa7d9; }
                QSplitter::handle { background-color: #c5d4e3; }
                QStatusBar { background-color: #ffffff; color: #102a43; border-top: 1px solid #c5d4e3; }
            """)
        else:
            self.setStyleSheet("")
        if hasattr(self, "tabs") and hasattr(self.tabs, "set_theme"):
            self.tabs.set_theme(theme)
        if hasattr(self, "status_bar"):
            self.status_bar.showMessage(f"Thème appliqué : {theme}")

    def _anova_posthoc(self, group_data, group_labels, alpha):
        lines = []
        comparisons = []
        for i in range(len(group_data)):
            for j in range(i + 1, len(group_data)):
                t_stat, p_val = scipy_stats.ttest_ind(group_data[i], group_data[j], equal_var=False)
                comparisons.append((group_labels[i], group_labels[j], p_val, np.mean(group_data[i]) - np.mean(group_data[j])))
        m = max(len(comparisons), 1)
        for g1, g2, p_val, diff in comparisons:
            p_adj = min(p_val * m, 1.0)
            sig = "SIGNIFICATIF" if p_adj < alpha else "non significatif"
            lines.append(f"  {g1} vs {g2}: diff moy={diff:.4f}, p={p_val:.6f}, p_adj={p_adj:.6f} → {sig}")
        return lines

    def _detect_control_rules(self, values, center, sigma):
        if sigma is None or sigma == 0 or np.isnan(sigma):
            return []
        x = np.asarray(values, dtype=float)
        signals = []
        idx = np.where(np.abs(x - center) > 3 * sigma)[0]
        if len(idx):
            signals.append("  Règle 1 : point au-delà de 3σ aux positions " + ", ".join(str(i + 1) for i in idx[:20]))
        for i in range(len(x) - 2):
            win = x[i:i+3] - center
            if np.sum(win > 2 * sigma) >= 2 or np.sum(win < -2 * sigma) >= 2:
                signals.append(f"  Règle 2 : 2 points sur 3 au-delà de 2σ autour des positions {i+1}-{i+3}")
                break
        for i in range(len(x) - 4):
            win = x[i:i+5] - center
            if np.sum(win > sigma) >= 4 or np.sum(win < -sigma) >= 4:
                signals.append(f"  Règle 3 : 4 points sur 5 au-delà de 1σ autour des positions {i+1}-{i+5}")
                break
        signs = np.sign(x - center)
        for i in range(len(signs) - 7):
            win = signs[i:i+8]
            if np.all(win > 0) or np.all(win < 0):
                signals.append(f"  Nelson : 8 points consécutifs du même côté autour des positions {i+1}-{i+8}")
                break
        for i in range(len(x) - 5):
            diff = np.diff(x[i:i+6])
            if np.all(diff > 0) or np.all(diff < 0):
                signals.append(f"  Nelson : tendance de 6 points autour des positions {i+1}-{i+6}")
                break
        return signals

    def _collect_report_sections(self):
        sections = []
        mapping = [
            ("Identification distribution", "dist_result_text"), ("Capabilité", "cap_result_text"),
            ("Normalité", "norm_result_text"), ("Valeurs aberrantes", "out_result_text"),
            ("Corrélation", "corr_result_text"), ("Régression", "reg_result_text"),
            ("Test t", "tt_result_text"), ("ANOVA", "anova_result_text"),
            ("Boxplots", "boxplot_result_text"), ("Cartes de contrôle", "cc_result_text"),
            ("Graphiques probabilité", "prob_result_text"), ("MSA / Gage R&R", "msa_result_text"),
        ]
        for title, attr in mapping:
            if hasattr(self, attr):
                txt = getattr(self, attr).toPlainText().strip()
                if txt:
                    sections.append((title, txt))
        return sections

    def _refresh_report_panel(self):
        if not hasattr(self, "report_text"):
            return
        sections = self._collect_report_sections()
        lines = ["RAPPORT D'ANALYSE STATPRO", "=" * 70,
                 f"Date : {pd.Timestamp.now().strftime('%Y-%m-%d %H:%M')}", ""]
        if self.current_project_path:
            lines.append(f"Projet : {self.current_project_path}\n")
        if not sections:
            lines.append("Aucun résultat disponible. Lancez une ou plusieurs analyses.")
        for title, content in sections:
            lines.extend(["", title.upper(), "-" * 70, content])
        self.report_text.setText("\n".join(lines))

    def _canvas_map(self):
        return {
            "Identification distribution": self.dist_canvas, "Capabilité": self.cap_canvas, "Normalité": self.norm_canvas,
            "Valeurs aberrantes": self.out_canvas, "Corrélation": self.corr_canvas, "Régression": self.reg_canvas,
            "Test t": self.tt_canvas, "ANOVA": self.anova_canvas, "Boxplots": self.boxplot_canvas,
            "Cartes de contrôle": self.cc_canvas, "Graphiques probabilité": self.prob_canvas, "MSA / Gage R&R": self.msa_canvas,
        }

    def _export_full_report(self):
        self._refresh_report_panel()
        filepath, _ = QFileDialog.getSaveFileName(self, "Exporter rapport complet", "rapport_statpro.pdf", "PDF (*.pdf);;DOCX (*.docx)")
        if not filepath:
            return
        try:
            ext = os.path.splitext(filepath)[1].lower()
            sections = self._collect_report_sections()
            tmpdir = tempfile.mkdtemp(prefix="statpro_report_")
            images = {}
            for title, canvas in self._canvas_map().items():
                if canvas is not None and canvas.fig is not None and canvas.fig.axes:
                    img = os.path.join(tmpdir, title.replace("/", "_").replace(" ", "_") + ".png")
                    canvas.fig.savefig(img, dpi=150, bbox_inches="tight")
                    images[title] = img
            if ext == ".docx":
                from docx import Document
                doc = Document()
                doc.add_heading("Rapport d'analyse StatPro", 0)
                doc.add_paragraph(f"Date : {pd.Timestamp.now().strftime('%Y-%m-%d %H:%M')}")
                for title, content in sections:
                    doc.add_heading(title, level=1)
                    doc.add_paragraph(content)
                    if title in images:
                        doc.add_picture(images[title])
                doc.save(filepath)
            else:
                from reportlab.lib.pagesizes import A4
                from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image, PageBreak
                from reportlab.lib.styles import getSampleStyleSheet
                styles = getSampleStyleSheet()
                story = [Paragraph("Rapport d'analyse StatPro", styles["Title"]), Paragraph(f"Date : {pd.Timestamp.now().strftime('%Y-%m-%d %H:%M')}", styles["Normal"]), Spacer(1, 12)]
                for title, content in sections:
                    story.append(Paragraph(title, styles["Heading1"]))
                    story.append(Paragraph("<br/>".join(content.replace("&", "&amp;").replace("<", "&lt;").splitlines()), styles["Code"]))
                    if title in images:
                        story.append(Spacer(1, 8))
                        story.append(Image(images[title], width=500, height=300, kind="proportional"))
                    story.append(PageBreak())
                SimpleDocTemplate(filepath, pagesize=A4).build(story)
            shutil.rmtree(tmpdir, ignore_errors=True)
            self.status_bar.showMessage(f"Rapport complet exporté : {filepath}")
        except Exception as e:
            QMessageBox.critical(self, "Erreur", f"Impossible d'exporter le rapport :\n{e}")

    def _export_excel_workbook(self):
        filepath, _ = QFileDialog.getSaveFileName(self, "Exporter Excel multi-feuilles", "statpro_resultats.xlsx", "Excel (*.xlsx)")
        if not filepath:
            return
        if not filepath.lower().endswith(".xlsx"):
            filepath += ".xlsx"
        try:
            with pd.ExcelWriter(filepath, engine="openpyxl") as writer:
                self._sheet_to_dataframe().to_excel(writer, sheet_name="Données", index=False)
                for title, content in self._collect_report_sections():
                    safe = title[:31].replace("/", "-")
                    pd.DataFrame({"Résultats": content.splitlines()}).to_excel(writer, sheet_name=safe, index=False)
                pd.DataFrame({"Rapport": (self.report_text.toPlainText() if hasattr(self, "report_text") else "").splitlines()}).to_excel(writer, sheet_name="Rapport", index=False)
            self.status_bar.showMessage(f"Excel multi-feuilles exporté : {filepath}")
        except Exception as e:
            QMessageBox.critical(self, "Erreur", f"Impossible d'exporter Excel :\n{e}")

    def _show_about(self):
        QMessageBox.information(
            self, "À propos",
            "StatPro - Analyse Statistique\n\n"
            "Alternative à Minitab\n\n"
            "Fonctionnalités :\n"
            "- Tableur avec Copier/Coller et menu contextuel\n"
            "- Capabilité (six-pack), Normalité, Identification distribution\n"
            "- Cartes de contrôle (X̄-R, I-MR, EWMA, CUSUM)\n"
            "- Corrélation, Régression, ANOVA, Test t\n"
            "- Boxplots, MSA/Gage R&R, Valeurs aberrantes\n"
            "- Export TXT/PDF/ODT/DOCX\n\n"
            "© 2026",
        )


def main():
    app = QApplication(sys.argv)
    app.setStyle("Fusion")
    window = StatisticalApp()
    window.show()
    sys.exit(app.exec_())


if __name__ == "__main__":
    main()
